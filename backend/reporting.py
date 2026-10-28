from auth import get_current_user
from database import get_db
from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from datetime import datetime, timezone, timedelta
from typing import Optional
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.units import mm
import io

report_router = APIRouter()
REPORT_ROLES = {'admin', 'gestor', 'rh', 'seguranca_trabalho', 'almoxarifado'}  
REPORT_TIMEZONE = ZoneInfo('America/Sao_Paulo')

def _report_guard(current_user):
    if current_user.get('role') not in REPORT_ROLES:
        raise HTTPException(status_code=403, detail='Sem permissao para consultar relatorios')

def _report_dt(value, end=False):
    if not value: return None
    raw = value.strip()
    if len(raw) == 10:
        parsed = datetime.fromisoformat(raw)
        if end: parsed += timedelta(days=1) - timedelta(microseconds=1)
        return parsed.replace(tzinfo=timezone.utc)
    parsed = datetime.fromisoformat(raw.replace('Z', '+00:00'))
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)

def _facial_status(delivery):
    status, message, score = delivery.get('facial_validation_status'), delivery.get('facial_validation_message'), delivery.get('facial_match_score')
    if status: return f'{status} ({score:.0%})' if isinstance(score, (int, float)) else str(status)
    if message: return str(message)
    return f'Reconhecido ({score:.0%})' if isinstance(score, (int, float)) else 'Nao registrado'

def _delivery_kind(delivery, item):
    dtype = str(delivery.get('delivery_type') or '').lower()
    return 'Kit' if 'kit' in dtype or item.get('kit_id') or item.get('kit_name') else 'Individual'

def _report_local_dt(value):
    if not value:
        return None
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace('Z', '+00:00'))
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(REPORT_TIMEZONE)

async def _report_rows(start_date=None, end_date=None, employee_id=None, company_id=None, department=None, epi_id=None, variation=None, include_returns=True, return_only=False, search=None):
    db = await get_db(); query = {}
    start, end = _report_dt(start_date), _report_dt(end_date, end=True)
    if start or end:
        query['created_at'] = {}
        if start: query['created_at']['$gte'] = start
        if end: query['created_at']['$lte'] = end
    if employee_id: query['employee_id'] = employee_id
    if return_only:
        query['is_return'] = True
    elif include_returns is False:
        query['is_return'] = False
    deliveries = await db.deliveries.find(query).sort('created_at', 1).to_list(100000)
    employee_ids = {str(d.get('employee_id')) for d in deliveries if d.get('employee_id')}
    valid_ids = [ObjectId(v) for v in employee_ids if ObjectId.is_valid(v)]
    employee_docs = await db.employees.find({'_id': {'$in': valid_ids}}).to_list(100000) if valid_ids else []
    employees = {str(e['_id']): e for e in employee_docs}
    company_ids = {str(e.get('company_id')) for e in employees.values() if e.get('company_id')}
    valid_company_ids = [ObjectId(v) for v in company_ids if ObjectId.is_valid(v)]
    company_docs = await db.companies.find({'_id': {'$in': valid_company_ids}}).to_list(100000) if valid_company_ids else []
    companies = {str(c['_id']): c for c in company_docs}; rows = []
    for delivery in deliveries:
        employee = employees.get(str(delivery.get('employee_id')), {})
        company = companies.get(str(employee.get('company_id')), {})
        company_name, dept = company.get('trade_name') or company.get('legal_name') or '', employee.get('department') or ''
        if company_id and str(employee.get('company_id')) != str(company_id): continue
        if department and dept != department: continue
        for item in delivery.get('items') or []:
            item_epi_id = str(item.get('epi_id') or item.get('epi_group_key') or '')
            item_variation = item.get('size') or item.get('variation') or item.get('epi_variation_id') or ''
            if epi_id and item_epi_id != str(epi_id): continue
            if variation and str(item_variation) != str(variation): continue
            rows.append({'delivery_id': str(delivery.get('_id')), 'created_at': delivery.get('created_at'), 'record_date': _report_local_dt(delivery.get('created_at')), 'employee_id': str(delivery.get('employee_id') or ''), 'employee_name': delivery.get('employee_name') or employee.get('full_name') or 'Nao informado', 'registration_number': employee.get('registration_number') or '', 'company': company_name, 'company_id': str(employee.get('company_id') or ''), 'department': dept, 'epi_id': item_epi_id, 'epi_name': item.get('epi_name') or item.get('kit_name') or 'EPI nao informado', 'variation': item_variation, 'quantity': float(item.get('quantity') or 0), 'is_return': bool(delivery.get('is_return')), 'delivery_type': 'Devolucao' if delivery.get('is_return') else _delivery_kind(delivery, item), 'delivered_by': delivery.get('delivered_by_name') or delivery.get('session_username') or delivery.get('terminal_user') or 'Nao registrado', 'attendant': delivery.get('attendant_name') or '', 'notes': delivery.get('notes') or item.get('notes') or '', 'facial_validation': _facial_status(delivery), 'kit_name': item.get('kit_name') or ''})
    if search:
        needle = str(search).strip().lower()
        rows = [row for row in rows if needle in ' '.join(str(row.get(key) or '') for key in ('employee_name', 'registration_number', 'epi_name', 'variation', 'department', 'company', 'notes')).lower()]
    return rows

def _group_rows(rows, view):
    keys = {'epi': ('epi_name', 'variation'), 'company': ('company',), 'department': ('department',), 'employee': ('employee_name', 'registration_number')}.get(view, ('epi_name', 'variation')); groups = {}
    for row in rows:
        key = tuple(row.get(k) or 'Nao informado' for k in keys); group = groups.setdefault(key, {'quantity': 0, 'employees': set()}); group['quantity'] += row['quantity']; group['employees'].add(row['employee_id'])
    result = []
    for key, group in sorted(groups.items(), key=lambda x: tuple(str(v).lower() for v in x[0])):
        record = {k: v for k, v in zip(keys, key)}; record.update({'quantity': group['quantity'], 'collaborators': len(group['employees']), 'average': group['quantity'] / len(group['employees']) if group['employees'] else 0}); result.append(record)
    return result

def _xlsx_response(wb, filename):
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return StreamingResponse(output, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="{filename}"'})

def _xlsx_setup(title, period, filters, headers, data, total_values):
    wb = Workbook(); ws = wb.active; ws.title = 'Relatorio'; ws['A1'] = 'GestorEPI'; ws['A1'].font = Font(bold=True, size=16, color='0F766E'); ws['A2'] = title; ws['A2'].font = Font(bold=True, size=13); ws['A3'] = f'Periodo: {period}'; ws['A4'] = 'Filtros: ' + (', '.join(f'{k}: {v}' for k, v in filters.items() if v not in (None, '', 'Todos')) or 'Todos')
    header_row = 6
    for col, value in enumerate(headers, 1):
        cell = ws.cell(header_row, col, value); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='0F766E'); cell.alignment = Alignment(horizontal='center')
    for r_index, row in enumerate(data, header_row + 1):
        for c_index, value in enumerate(row, 1): ws.cell(r_index, c_index, value)
    total_row = header_row + 1 + len(data)
    for col, value in enumerate(total_values, 1): ws.cell(total_row, col, value)
    for cell in ws[total_row]: cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='D1FAE5')
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{max(total_row, header_row + 1)}'; ws.freeze_panes = f'A{header_row + 1}'
    widths = {}
    for row in ws.iter_rows(min_row=header_row, max_row=total_row):
        for cell in row: widths[cell.column] = min(max(widths.get(cell.column, 0), len(str(cell.value or '')) + 2), 34)
    for col, width in widths.items(): ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=header_row + 1, max_row=total_row):
        for cell in row:
            if 'Data' in str(ws.cell(header_row, cell.column).value): cell.number_format = 'DD/MM/YYYY'
            if 'Hora' in str(ws.cell(header_row, cell.column).value): cell.number_format = 'HH:MM'
    return wb

def _pdf_response(title, period, filters, headers, data, totals, filename):
    buffer = io.BytesIO(); styles = getSampleStyleSheet(); elements = [Paragraph('GestorEPI', ParagraphStyle('ReportTitle', parent=styles['Heading1'], alignment=1, textColor=colors.HexColor('#0F766E'))), Paragraph(title, styles['Heading2']), Paragraph(f'Periodo: {period}', styles['Normal']), Paragraph('Filtros: ' + (', '.join(f'{k}: {v}' for k, v in filters.items() if v not in (None, '', 'Todos')) or 'Todos'), styles['Normal']), Paragraph(f'Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}', styles['Normal']), Spacer(1, 8)]
    table = Table([headers] + [[str(v)[:45] for v in row] for row in data] + [totals], repeatRows=1); table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F766E')), ('TEXTCOLOR', (0,0), (-1,0), colors.white), ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'), ('FONTSIZE', (0,0), (-1,-1), 7), ('GRID', (0,0), (-1,-1), .25, colors.grey), ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#F0FDFA')]), ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#D1FAE5')), ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold')]))
    elements.append(table); doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm, topMargin=12*mm, bottomMargin=14*mm); footer = lambda c,d: c.drawRightString(285*mm, 8*mm, f'Pagina {c.getPageNumber()}'); doc.build(elements, onFirstPage=footer, onLaterPages=footer); buffer.seek(0)
    return StreamingResponse(buffer, media_type='application/pdf', headers={'Content-Disposition': f'attachment; filename="{filename}"'})

def _report_guard(current_user):
    if current_user.get('role') not in REPORT_ROLES: raise HTTPException(status_code=403, detail='Sem permissao para consultar relatorios')

@report_router.get('/reports/consumption')
async def report_consumption(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, view: str = 'detail', current_user: dict = Depends(get_current_user)):
    _report_guard(current_user); rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, False); result = _group_rows(rows, view) if view != 'detail' else rows
    return {'view': view, 'rows': result, 'totals': {'quantity': sum(r['quantity'] for r in rows), 'collaborators': len({r['employee_id'] for r in rows}), 'rows': len(rows)}}

@report_router.get('/reports/daily-deliveries')
async def report_daily_deliveries(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _report_guard(current_user); rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, True)
    return {'rows': rows, 'totals': {'deliveries': len({r['delivery_id'] for r in rows}), 'collaborators': len({r['employee_id'] for r in rows}), 'items': sum(r['quantity'] for r in rows), 'types': len({r['epi_id'] for r in rows})}}

@report_router.get('/reports/returns')
async def report_returns(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, search: Optional[str] = None, page: int = 1, page_size: int = 100, current_user: dict = Depends(get_current_user)):
    _report_guard(current_user)
    rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, return_only=True, search=search)
    page = max(1, page); page_size = min(500, max(1, page_size)); total_rows = len(rows)
    start = (page - 1) * page_size
    return {'rows': rows[start:start + page_size], 'page': page, 'page_size': page_size, 'total_rows': total_rows, 'total_pages': (total_rows + page_size - 1) // page_size, 'totals': {'returns': len({r['delivery_id'] for r in rows}), 'collaborators': len({r['employee_id'] for r in rows}), 'items': sum(r['quantity'] for r in rows), 'quantity': sum(r['quantity'] for r in rows)}}

@report_router.get('/reports/returns/excel')
async def report_returns_excel(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, search: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _report_guard(current_user)
    rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, return_only=True, search=search)
    headers = ['Data', 'Hora', 'Colaborador', 'Matricula', 'Empresa', 'Setor', 'EPI', 'Variacao/Tamanho', 'Quantidade', 'Observacao', 'Responsavel', 'Atendente']
    data = [[r['record_date'].date() if r['record_date'] else '', r['record_date'].time() if r['record_date'] else '', r['employee_name'], r['registration_number'], r['company'], r['department'], r['epi_name'], r['variation'], r['quantity'], r['notes'], r['delivered_by'], r['attendant']] for r in rows]
    total = ['', '', f'{len({r["delivery_id"] for r in rows})} devolucoes', '', '', '', '', '', sum(r['quantity'] for r in rows), '', '', '']
    wb = _xlsx_setup('Relatorio de EPIs Devolvidos', f'{start_date or "inicio"} a {end_date or "hoje"}', {'Empresa': company_id, 'Setor': department, 'Colaborador': employee_id, 'EPI': epi_id, 'Variacao': variation, 'Pesquisa': search}, headers, data, total)
    return _xlsx_response(wb, 'Relatorio_Devolucoes_EPI.xlsx')

@report_router.get('/reports/returns/pdf')
async def report_returns_pdf(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, search: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _report_guard(current_user)
    rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, return_only=True, search=search)
    headers = ['Data', 'Hora', 'Colaborador', 'Empresa', 'Setor', 'EPI', 'Variacao', 'Qtd.', 'Observacao', 'Responsavel', 'Atendente']
    data = [[r['record_date'].strftime('%d/%m/%Y') if r['record_date'] else '-', r['record_date'].strftime('%H:%M') if r['record_date'] else '-', r['employee_name'], r['company'], r['department'], r['epi_name'], r['variation'], r['quantity'], r['notes'], r['delivered_by'], r['attendant']] for r in rows]
    total = ['', '', f'{len({r["delivery_id"] for r in rows})} devolucoes', '', '', '', sum(r['quantity'] for r in rows), '', '', '', '']
    return _pdf_response('Relatorio de EPIs Devolvidos', f'{start_date or "inicio"} a {end_date or "hoje"}', {'Empresa': company_id, 'Setor': department, 'Colaborador': employee_id, 'EPI': epi_id, 'Variacao': variation, 'Pesquisa': search}, headers, data, total, 'relatorio_devolucoes.pdf')

@report_router.get('/reports/consumption/excel')
async def report_consumption_excel(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, view: str = 'detail', current_user: dict = Depends(get_current_user)):
    _report_guard(current_user); rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, False)
    if view == 'detail':
        headers = ['Data', 'EPI', 'Variacao/Tamanho', 'Quantidade consumida', 'Colaborador', 'Matricula', 'Empresa', 'Setor', 'Colaboradores atendidos']
        data = [[r['record_date'].date() if r['record_date'] else '', r['epi_name'], r['variation'], r['quantity'], r['employee_name'], r['registration_number'], r['company'], r['department'], ''] for r in rows]
        total = ['', '', 'TOTAL', sum(r['quantity'] for r in rows), '', '', '', '', len({r['employee_id'] for r in rows})]
    else:
        headers = ['EPI', 'Variacao/Tamanho', 'Quantidade consumida', 'Colaborador', 'Matricula', 'Empresa', 'Setor', 'Colaboradores atendidos']
        data = [[r.get('epi_name') or r.get('company') or r.get('department') or r.get('employee_name'), r.get('variation', ''), r['quantity'], '', r.get('registration_number', ''), r.get('company', ''), r.get('department', ''), r['collaborators']] for r in _group_rows(rows, view)]
        total = ['', 'TOTAL', sum(r['quantity'] for r in rows), '', '', '', '', len({r['employee_id'] for r in rows})]
    wb = _xlsx_setup('Relatorio de Consumo de EPI', f'{start_date or "inicio"} a {end_date or "hoje"}', {'Empresa': company_id, 'Setor': department, 'Colaborador': employee_id, 'EPI': epi_id, 'Variacao': variation, 'Modo': view}, headers, data, total)
    return _xlsx_response(wb, 'Relatorio_Consumo_EPI.xlsx')

@report_router.get('/reports/daily-deliveries/excel')
async def report_daily_deliveries_excel(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _report_guard(current_user); rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, True); headers = ['Data', 'Hora', 'Colaborador', 'Matricula', 'Empresa', 'Setor', 'EPI', 'Variacao/Tamanho', 'Quantidade', 'Individual/Kit', 'Responsavel', 'Atendente', 'Validacao facial', 'Tipo']
    data = [[r['record_date'].date() if r['record_date'] else '', r['record_date'].time() if r['record_date'] else '', r['employee_name'], r['registration_number'], r['company'], r['department'], r['epi_name'], r['variation'], r['quantity'], r['delivery_type'], r['delivered_by'], r['attendant'], r['facial_validation'], 'Devolucao' if r['is_return'] else 'Entrega'] for r in rows]
    total = ['', '', f'{len({r["delivery_id"] for r in rows})} entregas', '', '', '', '', '', sum(r['quantity'] for r in rows), '', '', '', '', '']; wb = _xlsx_setup('Relatorio de Entregas Diarias', f'{start_date or "inicio"} a {end_date or "hoje"}', {'Empresa': company_id, 'Setor': department, 'Colaborador': employee_id, 'EPI': epi_id, 'Variacao': variation}, headers, data, total)
    return _xlsx_response(wb, 'Relatorio_Entregas.xlsx')

@report_router.get('/reports/consumption/pdf')
async def report_consumption_pdf(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, view: str = 'detail', current_user: dict = Depends(get_current_user)):
    _report_guard(current_user); rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, False); grouped = _group_rows(rows, view) if view != 'detail' else rows
    if view == 'detail':
        headers = ['Data', 'EPI', 'Variacao', 'Quantidade', 'Colaborador', 'Empresa', 'Setor']
        data = [[r['record_date'].strftime('%d/%m/%Y') if r['record_date'] else '-', r['epi_name'], r['variation'], r['quantity'], r['employee_name'], r['company'], r['department']] for r in grouped]
        total = ['', 'TOTAL', '', sum(r['quantity'] for r in rows), len({r['employee_id'] for r in rows}), '', '']
    else:
        headers = ['EPI/Agrupamento', 'Variacao', 'Quantidade', 'Colaborador/Atendidos', 'Empresa/Media', 'Setor']
        data = [[r.get('epi_name') or r.get('company') or r.get('department') or r.get('employee_name'), r.get('variation', ''), r['quantity'], r['collaborators'], r['average'], ''] for r in grouped]
        total = ['TOTAL', '', sum(r['quantity'] for r in rows), len({r['employee_id'] for r in rows}), '', '']
    return _pdf_response('Relatorio de Consumo de EPI', f'{start_date or "inicio"} a {end_date or "hoje"}', {'Empresa': company_id, 'Setor': department, 'Colaborador': employee_id, 'EPI': epi_id, 'Variacao': variation, 'Modo': view}, headers, data, total, 'relatorio_consumo.pdf')

@report_router.get('/reports/daily-deliveries/pdf')
async def report_daily_deliveries_pdf(start_date: Optional[str] = None, end_date: Optional[str] = None, company_id: Optional[str] = None, department: Optional[str] = None, employee_id: Optional[str] = None, epi_id: Optional[str] = None, variation: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _report_guard(current_user); rows = await _report_rows(start_date, end_date, employee_id, company_id, department, epi_id, variation, True)
    data = [[r['record_date'].strftime('%d/%m/%Y') if r['record_date'] else '', r['record_date'].strftime('%H:%M') if r['record_date'] else '', r['employee_name'], r['epi_name'], r['variation'], r['quantity'], r['delivery_type'], r['delivered_by'], r['attendant'], r['facial_validation']] for r in rows]
    return _pdf_response('Relatorio de Entregas Diarias', f'{start_date or "inicio"} a {end_date or "hoje"}', {'Empresa': company_id, 'Setor': department, 'Colaborador': employee_id, 'EPI': epi_id, 'Variacao': variation}, ['Data', 'Hora', 'Colaborador', 'EPI', 'Variacao', 'Qtd.', 'Tipo', 'Responsavel', 'Atendente', 'Facial'], data, ['', '', f'{len({r["delivery_id"] for r in rows})} entregas', '', '', sum(r['quantity'] for r in rows), '', '', '', ''], 'relatorio_entregas.pdf')

