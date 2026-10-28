from fastapi import FastAPI, APIRouter, Depends, HTTPException, status, UploadFile, File, Form, Request, Query
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from database import connect_db, close_db, get_db
from telemetry_serialization import to_bson_safe
from schemas import *
from auth import *
from datetime import datetime, timezone, timedelta
from typing import List, Optional
from pathlib import Path
from bson import ObjectId
from pymongo.errors import DuplicateKeyError
import os
import time
import uuid
import shutil
import logging
import base64
import json
import io
import re
import math
import unicodedata

# Para importaÃ§Ã£o de Excel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Para geraÃ§Ã£o de PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from services.face_recognition_service import FaceRecognitionService
from services.phase103_benchmark import (
    Phase103CloseRequest,
    Phase103CompareRequest,
    Phase103ReferenceRequest,
    Phase103SessionRequest,
    add_reference as phase103_add_reference,
    close_session as phase103_close_session,
    compare as phase103_compare,
    resume_session as phase103_resume_session,
    start_session as phase103_start_session,
)
from pydantic import ValidationError

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

ROOT_DIR = Path(__file__).parent
UPLOAD_DIR = ROOT_DIR / 'uploads'
UPLOAD_DIR.mkdir(exist_ok=True)

# Dias para expiraÃ§Ã£o de senha
PASSWORD_EXPIRY_DAYS = 30
face_recognition_service = FaceRecognitionService()

app = FastAPI(title='GestorEPI API')
api_router = APIRouter(prefix='/api')

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=['*'],
    allow_methods=['*'],
    allow_headers=['*'],
    expose_headers=['*'],
)

# Middleware para adicionar headers de cache-control e CORS em arquivos estÃ¡ticos
@app.middleware("http")
async def add_cache_headers(request, call_next):
    response = await call_next(request)
    path = request.url.path
    
    # Para arquivos de upload, adicionar headers de cache e CORS
    if '/uploads/' in path or '/api/uploads/' in path:
        response.headers['Cache-Control'] = 'public, max-age=86400'  # 1 dia
        response.headers['Access-Control-Allow-Origin'] = '*'
        response.headers['Access-Control-Allow-Methods'] = 'GET, OPTIONS'
        response.headers['X-Content-Type-Options'] = 'nosniff'
    elif path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, proxy-revalidate, max-age=0'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
    
    return response

# Servir arquivos de upload - DEVE ter prefixo /api para funcionar com ingress
app.mount('/api/uploads', StaticFiles(directory=str(UPLOAD_DIR)), name='api_uploads')
app.mount('/uploads', StaticFiles(directory=str(UPLOAD_DIR)), name='uploads')

@app.on_event('startup')
async def startup_event():
    await connect_db()
    from seed import seed_database
    await seed_database()
    db = await get_db()
    await face_recognition_service.initialize_cache(db)

@app.on_event('shutdown')
async def shutdown_event():
    await close_db()

def doc_to_response(doc, id_field='id'):
    if doc is None:
        return None

    result = {}
    for k, v in doc.items():
        if k == '_id':
            continue

        if isinstance(v, datetime):
            # Mongo armazena datetime em UTC sem tzinfo.
            # Enviamos com Z para o navegador converter corretamente para America/Sao_Paulo.
            if v.tzinfo is None:
                v = v.replace(tzinfo=timezone.utc)
            result[k] = v.isoformat().replace('+00:00', 'Z')
        else:
            result[k] = v

    result[id_field] = str(doc['_id'])
    return result

def parse_datetime_safe(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace('Z', '+00:00'))
        except Exception:
            return None
    return None

def parse_int_safe(value, default=0):
    if value is None or value == '':
        return default
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default

def get_variation_stock_quantity(raw_variation):
    quantity = raw_variation.get('quantidade_estoque')
    if quantity is None:
        quantity = raw_variation.get('current_stock')
    return parse_int_safe(quantity, 0)

def ensure_aware_datetime(value):
    parsed = parse_datetime_safe(value)
    if not parsed:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed

def format_datetime_safe(value, fmt='%d/%m/%Y %H:%M', fallback='-'):
    parsed = ensure_aware_datetime(value)
    return parsed.strftime(fmt) if parsed else fallback

def format_facial_template_response(doc):
    if doc is None:
        return None
    return {
        "id": str(doc["_id"]),
        "employee_id": str(doc.get("employee_id")) if doc.get("employee_id") else None,
        "descriptor": doc.get("descriptor"),
        "pose_label": doc.get("pose_label"),
        "quality_score": doc.get("quality_score"),
        "detection_score": doc.get("detection_score"),
        "created_by": doc.get("created_by"),
        "created_by_name": doc.get("created_by_name"),
        "created_at": doc.get("created_at")
    }

COUNTER_USER_USERNAME = 'balcao_jac'
COUNTER_ATTENDANT_REQUIRED_MESSAGE = 'ATENDENTE NÃO SELECIONADO. Selecione o atendente responsável pela entrega.'
COUNTER_ATTENDANT_INACTIVE_MESSAGE = 'ATENDENTE INATIVO. Selecione outro atendente.'

def is_counter_user(user: dict) -> bool:
    return (user or {}).get('username') == COUNTER_USER_USERNAME

def is_counter_delivery(delivery: dict) -> bool:
    return (
        delivery.get('delivered_by_name') == COUNTER_USER_USERNAME or
        delivery.get('session_username') == COUNTER_USER_USERNAME or
        delivery.get('terminal_user') == COUNTER_USER_USERNAME
    )

def get_delivery_responsible_name(delivery: dict) -> str:
    if is_counter_delivery(delivery) and delivery.get('attendant_name'):
        return delivery.get('attendant_name')
    return (
        delivery.get('delivered_by_name') or
        delivery.get('created_by_user') or
        delivery.get('created_by_name') or
        delivery.get('username') or
        delivery.get('operator_name') or
        '-'
    )

async def get_counter_attendant_for_delivery(db, attendant_id: Optional[str]) -> dict:
    if not attendant_id or not ObjectId.is_valid(str(attendant_id)):
        raise HTTPException(status_code=422, detail=COUNTER_ATTENDANT_REQUIRED_MESSAGE)

    attendant = await db.counter_attendants.find_one({'_id': ObjectId(str(attendant_id))})
    if not attendant:
        raise HTTPException(status_code=422, detail=COUNTER_ATTENDANT_REQUIRED_MESSAGE)
    if not attendant.get('is_active', True):
        raise HTTPException(status_code=422, detail=COUNTER_ATTENDANT_INACTIVE_MESSAGE)
    return attendant

def biometric_quality_label(quality_score: Optional[float]) -> str:
    score = float(quality_score or 0)
    if score >= 0.85:
        return "Excelente"
    if score >= 0.70:
        return "Boa"
    if score >= 0.50:
        return "Regular"
    return "Ruim"

def has_registered_facial_template(template: dict) -> bool:
    descriptor = template.get("descriptor")
    if descriptor is None:
        return False
    if isinstance(descriptor, str):
        stripped = descriptor.strip()
        if not stripped:
            return False
        try:
            parsed = json.loads(stripped)
            if isinstance(parsed, list):
                return len(parsed) > 0
        except Exception:
            pass
        return True
    if isinstance(descriptor, (list, tuple)):
        return len(descriptor) > 0
    return True

BIOMETRIC_REFERENCE_TARGET = 3

def is_current_model_template(template: dict) -> bool:
    """Return whether a stored descriptor is usable by the current 512D model."""
    try:
        return face_recognition_service.descriptor_dimension(template.get("descriptor")) == 512
    except Exception:
        return False

def build_biometric_summary(employee: dict, templates: List[dict]) -> dict:
    registered_templates = [
        template for template in templates
        if has_registered_facial_template(template)
    ]
    compatible_templates = [
        template for template in templates
        if is_current_model_template(template)
    ]
    count = len(compatible_templates)
    legacy_count = len(registered_templates) - count
    if count >= BIOMETRIC_REFERENCE_TARGET:
        status_key = "registered"
        status_label = "Biometria ativa"
    elif count > 0:
        status_key = "incomplete"
        status_label = "Incompleta"
    elif legacy_count > 0:
        status_key = "legacy"
        status_label = "Biometria legada"
    else:
        status_key = "missing"
        status_label = "Nao cadastrada"

    dates = [
        ensure_aware_datetime(template.get("created_at"))
        for template in compatible_templates
        if ensure_aware_datetime(template.get("created_at"))
    ]
    quality_scores = [float(template.get("quality_score") or 0) for template in compatible_templates]
    avg_quality = round(sum(quality_scores) / len(quality_scores), 4) if quality_scores else 0
    latest_template = sorted(
        compatible_templates,
        key=lambda template: ensure_aware_datetime(template.get("created_at")) or datetime.min.replace(tzinfo=timezone.utc),
        reverse=True
    )[0] if compatible_templates else None

    return {
        "status": status_key,
        "status_label": status_label,
        "templates_count": count,
        "total_templates_count": len(templates),
        "compatible_templates_count": len(compatible_templates),
        "legacy_templates_count": legacy_count,
        "first_enrolled_at": min(dates) if dates else None,
        "last_updated_at": max(dates) if dates else None,
        "created_by": latest_template.get("created_by") if latest_template else None,
        "created_by_name": latest_template.get("created_by_name") if latest_template else None,
        "model": "InsightFace",
        "has_consent": employee.get("facial_consent", False),
        "consent_date": employee.get("facial_consent_date"),
        "quality_score": avg_quality,
        "quality_label": biometric_quality_label(avg_quality),
    }

async def get_biometric_summaries(db, employees: List[dict]) -> dict:
    employee_ids = [str(employee["_id"]) for employee in employees]
    object_ids = [employee["_id"] for employee in employees]
    templates = await db.facial_templates.find({
        "employee_id": {"$in": employee_ids + object_ids}
    }).to_list(10000)
    templates_by_employee = {}
    for template in templates:
        templates_by_employee.setdefault(str(template.get("employee_id")), []).append(template)
    return {
        str(employee["_id"]): build_biometric_summary(
            employee,
            templates_by_employee.get(str(employee["_id"]), [])
        )
        for employee in employees
    }

def check_password_expired(user):
    """Verifica se a senha expirou (mais de 30 dias)"""
    password_changed_at = ensure_aware_datetime(user.get('password_changed_at'))
    if not password_changed_at:
        return False

    expiry_date = password_changed_at + timedelta(days=PASSWORD_EXPIRY_DAYS)
    return datetime.now(timezone.utc) > expiry_date

# PermissÃµes por perfil
ROLE_PERMISSIONS = {
    'admin': ['all'],
    'gestor': ['dashboard', 'historico', 'colaboradores', 'empresas', 'epis', 'fornecedores', 'kits', 'estoque', 'movimentacoes', 'alertas', 'usuarios_read', 'courses'],
    'rh': ['dashboard', 'alertas', 'historico', 'colaboradores', 'colaboradores_full'],
    'seguranca_trabalho': ['dashboard', 'alertas', 'historico', 'colaboradores', 'colaboradores_full', 'empresas', 'fornecedores', 'kits', 'biometria', 'courses'],
    'almoxarifado': ['dashboard', 'alertas', 'entrega', 'historico', 'colaboradores', 'colaboradores_full', 'empresas', 'fornecedores', 'epis', 'estoque', 'movimentacoes', 'courses']
}

def can_view_sensitive_data(role):
    """Verifica se o perfil pode ver dados sensÃ­veis (CPF, RG, etc)"""
    return role in ['admin', 'gestor', 'rh', 'seguranca_trabalho', 'almoxarifado']

def can_manage_users(role):
    """Verifica se pode gerenciar usuÃ¡rios"""
    return role in ['admin']

def can_deliver_epi(role):
    """Verifica se pode fazer entregas de EPI"""
    return role in ['admin', 'almoxarifado']

def can_view_delivery_history(role):
    """Verifica se pode consultar historico de entregas."""
    return role in ['admin', 'gestor', 'rh', 'seguranca_trabalho', 'almoxarifado']

def can_manage_epis(role):
    """Verifica se pode gerenciar EPIs - inclui almoxarifado"""
    return role in ['admin', 'almoxarifado']

def can_manage_employees(role):
    """Verifica se pode cadastrar/editar colaboradores"""
    return role in ['admin', 'rh', 'seguranca_trabalho', 'almoxarifado']

def can_change_employee_status(role):
    return role in ['admin', 'rh', 'seguranca_trabalho', 'almoxarifado']

def can_manage_biometrics(role):
    """Verifica se pode cadastrar/reprocessar/excluir templates faciais."""
    return role in ['admin', 'rh', 'seguranca_trabalho', 'almoxarifado']

def can_view_courses(role):
    """Verifica se pode consultar o modulo de cursos."""
    return role in ['admin', 'gestor', 'seguranca_trabalho', 'almoxarifado']

def can_manage_courses(role):
    """Verifica se pode gerenciar cursos e vinculos com EPI."""
    return role in ['admin', 'seguranca_trabalho']


def can_view_employee_courses(role):
    return role in ['admin', 'gestor', 'rh', 'seguranca_trabalho', 'almoxarifado']


def can_manage_employee_courses(role):
    return role in ['admin', 'seguranca_trabalho', 'almoxarifado']


def can_deactivate_employee_courses(role):
    return role in ['admin', 'seguranca_trabalho']


EMPLOYEE_COURSE_FILE_EXTENSIONS = {'.pdf', '.jpg', '.jpeg', '.png'}
EMPLOYEE_COURSE_CONTENT_TYPES = {'application/pdf', 'image/jpeg', 'image/png'}
EMPLOYEE_COURSE_MAX_FILE_SIZE = 8 * 1024 * 1024

def normalize_employee_status(employee: dict) -> str:
    status_value = str((employee or {}).get('status') or 'active').strip().lower()
    return 'inactive' if status_value in ['inactive', 'inativo', 'disabled', 'desligado', 'false', '0'] else 'active'

def is_employee_active(employee: dict) -> bool:
    return normalize_employee_status(employee) == 'active'

def get_request_ip(request: Optional[Request]) -> Optional[str]:
    if not request:
        return None
    forwarded = request.headers.get('x-forwarded-for')
    if forwarded:
        return forwarded.split(',')[0].strip()
    return request.client.host if request.client else None

async def write_audit_log(db, action: str, entity_type: str, entity_id: str, current_user: dict, reason: str = None, previous_status: str = None, new_status: str = None, request: Optional[Request] = None, metadata: Optional[dict] = None):
    audit_doc = {
        'action': action,
        'entity_type': entity_type,
        'entity_id': entity_id,
        'performed_by': current_user.get('id'),
        'performed_by_username': current_user.get('username'),
        'performed_by_profile': current_user.get('role'),
        'reason': reason,
        'previous_status': previous_status,
        'new_status': new_status,
        'ip': get_request_ip(request),
        'metadata': metadata or {},
        'created_at': datetime.now(timezone.utc)
    }
    await db.audit_logs.insert_one(audit_doc)

async def block_inactive_identified_employee(db, result: dict) -> dict:
    employee_id = result.get('employee_id') if result else None
    if result.get('status') != 'approved' or not employee_id or not ObjectId.is_valid(str(employee_id)):
        return result
    employee = await db.employees.find_one({"_id": ObjectId(str(employee_id))}, {"status": 1, "full_name": 1})
    if employee and not is_employee_active(employee):
        return {
            **result,
            'status': 'blocked',
            'message': 'Colaborador identificado, porem esta inativo e nao pode receber EPI.',
            'employee_name': employee.get('full_name') or result.get('employee_name'),
            'liveness_required': False
        }
    return result

def friendly_facial_message(result: dict) -> str:
    if not result or result.get("status") == "approved":
        return result.get("message") if result else ""
    message = str(result.get("message") or "").lower()
    if "mais de uma" in message or "apenas um colaborador" in message:
        return "Mais de uma pessoa detectada. Fique sozinho diante da câmera."
    if "desfoc" in message or "nitid" in message:
        return "Mantenha o rosto parado por um instante."
    if "escura" in message or "ilumina" in message or "qualidade da imagem" in message:
        return "Melhore a iluminação do rosto e evite luz forte atrás de você."
    if "centralize" in message or "distante" in message or "proximo" in message or "nenhum rosto" in message:
        return "Aproxime-se e centralize o rosto."
    if "nenhum frame" in message or "nenhuma imagem" in message:
        return "Aproxime-se e centralize o rosto."
    if "conflit" in message:
        return "Não foi possível confirmar a mesma pessoa. Tente novamente."
    if "não confirmado" in message or "nao confirmado" in message:
        return "Mantenha o rosto parado por um instante e olhe para a câmera."
    return "Não foi possível confirmar o reconhecimento. Tente novamente."

async def save_facial_attempt_log(
    db,
    result: dict,
    current_user: dict,
    endpoint: str,
    frame_count: Optional[int] = None,
    attempt_id: Optional[str] = None,
    frames_captured: Optional[int] = None,
    frames_sent: Optional[int] = None,
    frame_telemetry: Optional[List[dict]] = None,
    backend_total_ms: Optional[float] = None,
    client_telemetry: Optional[dict] = None
):
    if not result:
        return
    now = datetime.now(timezone.utc)
    doc = {
        'created_at': now,
        'endpoint': endpoint,
        'frame_count': frame_count,
        'user_id': current_user.get('id') if current_user else None,
        'username': current_user.get('username') if current_user else None,
        'user_role': current_user.get('role') if current_user else None,
        'status': result.get('status'),
        'message': result.get('message'),
        'approved': result.get('status') == 'approved',
        'employee_id': result.get('employee_id'),
        'employee_name': result.get('employee_name'),
        'candidate_employee_id': result.get('candidate_employee_id'),
        'candidate_employee_name': result.get('candidate_employee_name'),
        'second_candidate_employee_id': result.get('second_candidate_employee_id'),
        'second_candidate_employee_name': result.get('second_candidate_employee_name'),
        'matched_template_id': result.get('matched_template_id'),
        'similarity_score': result.get('similarity_score'),
        'second_similarity_score': result.get('second_similarity_score'),
        'candidate_margin': result.get('candidate_margin'),
        'min_required_margin': result.get('min_required_margin'),
        'detection_confidence': result.get('detection_confidence'),
        'matched_pose_label': result.get('matched_pose_label'),
        'quality': result.get('quality'),
        'decision_rule': result.get('decision_rule'),
        'rejection_reason': None if result.get('status') == 'approved' else result.get('message'),
        'attempt_id': attempt_id,
        'frames_captured': frames_captured,
        'frames_sent': frames_sent,
        'usable_frame_count': sum(1 for frame in (frame_telemetry or []) if frame.get('quality_pass')),
        'backend_total_ms': backend_total_ms,
        'frame_telemetry': frame_telemetry or [],
        'client_telemetry': sanitize_client_telemetry(client_telemetry),
        'final_decision': result.get('status'),
        'final_reason': result.get('message'),
    }
    try:
        await db.facial_recognition_attempts.insert_one(to_bson_safe(doc))
    except Exception as exc:
        logger.warning(f"Falha ao registrar auditoria facial: {exc}")


def sanitize_client_telemetry(value: Optional[dict]) -> dict:
    """Keep only technical timing fields; never persist images or biometrics."""
    if not isinstance(value, dict):
        return {}

    allowed_numbers = {
        'camera_ready_ms', 'burst_ms', 'frames_captured', 'frames_usable',
        'attempt_sequence', 'response_ms', 'ui_result_ms',
    }
    sanitized = {}
    for key in allowed_numbers:
        raw = value.get(key)
        if isinstance(raw, (int, float)) and math.isfinite(float(raw)):
            sanitized[key] = round(float(raw), 3) if not isinstance(raw, int) else int(raw)

    attempt_id = value.get('attempt_id')
    if isinstance(attempt_id, str) and len(attempt_id) <= 80:
        sanitized['attempt_id'] = attempt_id

    result_status = value.get('result_status')
    if isinstance(result_status, str) and len(result_status) <= 40:
        sanitized['result_status'] = result_status

    frames = value.get('frame_telemetry')
    if isinstance(frames, list):
        sanitized_frames = []
        for frame in frames[:5]:
            if not isinstance(frame, dict):
                continue
            item = {}
            for key in ('frame_index', 'capture_ms', 'decode_ms', 'detection_ms', 'compression_ms', 'total_ms'):
                raw = frame.get(key)
                if isinstance(raw, (int, float)) and math.isfinite(float(raw)):
                    item[key] = round(float(raw), 3)
            quality_pass = frame.get('quality_pass')
            if isinstance(quality_pass, bool):
                item['quality_pass'] = quality_pass
            rejection_reason = frame.get('rejection_reason')
            if isinstance(rejection_reason, str) and len(rejection_reason) <= 80:
                item['rejection_reason'] = rejection_reason
            sanitized_frames.append(item)
        sanitized['frame_telemetry'] = sanitized_frames
    return sanitized


def normalized_telemetry_attempt_id(value: Optional[dict]) -> Optional[str]:
    """Return only a canonical UUID suitable for technical correlation."""
    candidate = sanitize_client_telemetry(value).get('attempt_id')
    if not isinstance(candidate, str):
        return None
    try:
        return str(uuid.UUID(candidate))
    except (ValueError, AttributeError, TypeError):
        return None

# ===================== AUTH =====================

@api_router.get('/')
async def root():
    return {'message': 'GestorEPI API'}

@api_router.get('/health')
async def health_check():
    """Endpoint de verificaÃ§Ã£o de saÃºde da API"""
    db_status = 'unknown'
    try:
        db = await get_db()
        await db.command('ping')
        db_status = 'connected'
    except Exception:
        db_status = 'error'
    
    return {
        'status': 'healthy' if db_status == 'connected' else 'degraded',
        'database': db_status,
        'service': 'GestorEPI',
        'version': '5.1.0',
        'timestamp': datetime.now(timezone.utc).isoformat()
    }

@api_router.post('/auth/login', response_model=TokenResponse)
async def login(request: LoginRequest):
    db = await get_db()
    
    # Verificar se o identificador Ã© email ou username
    identifier = request.username.strip()
    if '@' in identifier:
        # Buscar por email
        user = await db.users.find_one({"email": identifier.lower()})
    else:
        # Buscar por username
        user = await db.users.find_one({"username": identifier})
    
    if not user or not verify_password(request.password, user['hashed_password']):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail='Credenciais invÃ¡lidas')
    
    if not user.get('is_active', True):
        raise HTTPException(status_code=400, detail='UsuÃ¡rio inativo')
    
    license_doc = await db.panel_license.find_one({})
    if license_doc:
        now = datetime.now(timezone.utc)
        expires_at = ensure_aware_datetime(license_doc.get('expires_at'))
        if not expires_at:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail='Licenca invalida')
        if now > expires_at and user['role'] != 'admin':
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail='LicenÃ§a expirada')
    
    password_expired = check_password_expired(user)
    
    access_token = create_access_token(data={'sub': user['username'], 'role': user['role']})
    
    return TokenResponse(
        access_token=access_token,
        token_type='bearer',
        must_change_password=user.get('must_change_password', False),
        password_expired=password_expired,
        role=UserRole(user['role']),
        is_primary_admin=user.get('is_primary_admin', False)
    )

@api_router.post('/auth/change-password')
async def change_password(request: ChangePasswordRequest, current_user: dict = Depends(get_current_user)):
    db = await get_db()
    if not verify_password(request.old_password, current_user['hashed_password']):
        raise HTTPException(status_code=400, detail='Senha antiga incorreta')
    
    await db.users.update_one(
        {"_id": ObjectId(current_user['id'])},
        {"$set": {
            "hashed_password": get_password_hash(request.new_password), 
            "must_change_password": False,
            "password_changed_at": datetime.now(timezone.utc)
        }}
    )
    return {'message': 'Senha alterada com sucesso'}

@api_router.get('/auth/me', response_model=UserResponse)
async def get_me(current_user: dict = Depends(get_current_user)):
    return UserResponse(
        id=current_user['id'],
        username=current_user['username'],
        email=current_user['email'],
        role=UserRole(current_user['role']),
        is_active=current_user.get('is_active', True),
        must_change_password=current_user.get('must_change_password', False),
        password_changed_at=current_user.get('password_changed_at'),
        employee_id=current_user.get('employee_id'),
        created_at=current_user.get('created_at', datetime.now(timezone.utc))
    )

# ===================== USERS =====================

@api_router.get('/users', response_model=List[UserResponse])
async def get_users(current_user: dict = Depends(require_role('admin', 'gestor'))):
    db = await get_db()
    users = await db.users.find({}).to_list(1000)
    return [UserResponse(**doc_to_response(u)) for u in users]

@api_router.post('/users', response_model=UserResponse, status_code=status.HTTP_201_CREATED)
async def create_user(user_data: UserCreate, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    
    # RH nÃ£o pode criar admins
    if current_user['role'] == 'rh' and user_data.role == UserRole.ADMIN:
        raise HTTPException(status_code=403, detail='Sem permissÃ£o para criar administradores')
    
    # Normalizar email para lowercase
    email_normalized = user_data.email.lower().strip() if user_data.email else None
    
    existing = await db.users.find_one({"$or": [
        {"username": user_data.username}, 
        {"email": email_normalized}
    ]})
    if existing:
        if existing.get('username') == user_data.username:
            raise HTTPException(status_code=400, detail='Nome de usuÃ¡rio jÃ¡ existe')
        else:
            raise HTTPException(status_code=400, detail='E-mail jÃ¡ estÃ¡ em uso')
    
    new_user = {
        "username": user_data.username,
        "email": email_normalized,
        "hashed_password": get_password_hash(user_data.password),
        "role": user_data.role.value,
        "employee_id": user_data.employee_id,
        "must_change_password": True,
        "is_active": True,
        "password_changed_at": datetime.now(timezone.utc),
        "created_at": datetime.now(timezone.utc),
        "updated_at": datetime.now(timezone.utc)
    }
    result = await db.users.insert_one(new_user)
    new_user['_id'] = result.inserted_id
    return UserResponse(**doc_to_response(new_user))

@api_router.patch('/users/{user_id}', response_model=UserResponse)
async def update_user(user_id: str, user_data: UserUpdate, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    existing_user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not existing_user:
        raise HTTPException(status_code=404, detail='UsuÃƒÂ¡rio nÃƒÂ£o encontrado')
    if current_user['role'] == 'rh' and existing_user.get('role') == 'admin':
        raise HTTPException(status_code=403, detail='Sem permissÃƒÂ£o para editar administradores')

    update_data = {k: v for k, v in user_data.model_dump(exclude_unset=True).items()}
    if 'role' in update_data:
        update_data['role'] = update_data['role'].value
        if current_user['role'] == 'rh' and update_data['role'] == 'admin':
            raise HTTPException(status_code=403, detail='Sem permissÃƒÂ£o para atribuir perfil administrador')
    update_data['updated_at'] = datetime.now(timezone.utc)
    
    result = await db.users.find_one_and_update(
        {"_id": ObjectId(user_id)}, {"$set": update_data}, return_document=True
    )
    if not result:
        raise HTTPException(status_code=404, detail='UsuÃ¡rio nÃ£o encontrado')
    return UserResponse(**doc_to_response(result))

@api_router.delete('/users/{user_id}')
async def delete_user(user_id: str, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    result = await db.users.delete_one({"_id": ObjectId(user_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='UsuÃ¡rio nÃ£o encontrado')
    return {'message': 'UsuÃ¡rio excluÃ­do'}

class ResetPasswordRequest(BaseModel):
    new_password: str

@api_router.post('/users/{user_id}/reset-password')
async def reset_user_password(user_id: str, request: ResetPasswordRequest, current_user: dict = Depends(require_role('admin'))):
    """Permite ao administrador redefinir a senha de qualquer usuÃ¡rio"""
    db = await get_db()
    
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(status_code=404, detail='UsuÃ¡rio nÃ£o encontrado')
    
    # Validar complexidade da senha
    password = request.new_password
    import re
    if len(password) < 8:
        raise HTTPException(status_code=400, detail='A senha deve ter no mÃ­nimo 8 caracteres')
    if not re.search(r'[A-Z]', password):
        raise HTTPException(status_code=400, detail='A senha deve conter pelo menos uma letra maiÃºscula')
    if not re.search(r'[a-z]', password):
        raise HTTPException(status_code=400, detail='A senha deve conter pelo menos uma letra minÃºscula')
    if not re.search(r'\d', password):
        raise HTTPException(status_code=400, detail='A senha deve conter pelo menos um nÃºmero')
    if not re.search(r'[!@#$%^&*(),.?":{}|<>]', password):
        raise HTTPException(status_code=400, detail='A senha deve conter pelo menos um caractere especial')
    
    await db.users.update_one(
        {"_id": ObjectId(user_id)},
        {"$set": {
            "hashed_password": get_password_hash(password),
            "must_change_password": True,
            "password_changed_at": datetime.now(timezone.utc),
            "updated_at": datetime.now(timezone.utc)
        }}
    )
    
    return {'message': 'Senha redefinida com sucesso'}

# ===================== COMPANIES =====================

@api_router.get('/companies', response_model=List[CompanyResponse])
async def get_companies(current_user: dict = Depends(require_role('admin', 'gestor', 'rh', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    companies = await db.companies.find({}).to_list(1000)
    return [CompanyResponse(**doc_to_response(c)) for c in companies]

@api_router.post('/companies', response_model=CompanyResponse, status_code=status.HTTP_201_CREATED)
async def create_company(company_data: CompanyCreate, current_user: dict = Depends(require_role('admin', 'seguranca_trabalho', 'almoxarifado'))):
    logger.info(f"ðŸ“ Recebendo dados de empresa: {company_data.model_dump()}")
    db = await get_db()
    existing = await db.companies.find_one({"cnpj": company_data.cnpj})
    if existing:
        logger.error(f"âŒ CNPJ jÃ¡ cadastrado: {company_data.cnpj}")
        raise HTTPException(status_code=400, detail='CNPJ jÃ¡ cadastrado')
    
    new_company = {**company_data.model_dump(), "created_at": datetime.now(timezone.utc), "updated_at": datetime.now(timezone.utc)}
    logger.info(f"âœ… Inserindo empresa: {new_company}")
    result = await db.companies.insert_one(new_company)
    new_company['_id'] = result.inserted_id
    logger.info(f"âœ… Empresa criada com ID: {result.inserted_id}")
    return CompanyResponse(**doc_to_response(new_company))

@api_router.get('/companies/{company_id}', response_model=CompanyResponse)
async def get_company(company_id: str, current_user: dict = Depends(require_role('admin', 'gestor', 'rh', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    company = await db.companies.find_one({"_id": ObjectId(company_id)})
    if not company:
        raise HTTPException(status_code=404, detail='Empresa nÃ£o encontrada')
    return CompanyResponse(**doc_to_response(company))

@api_router.patch('/companies/{company_id}', response_model=CompanyResponse)
async def update_company(company_id: str, company_data: CompanyUpdate, current_user: dict = Depends(require_role('admin', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    update_data = {k: v for k, v in company_data.model_dump(exclude_unset=True).items()}
    update_data['updated_at'] = datetime.now(timezone.utc)
    result = await db.companies.find_one_and_update(
        {"_id": ObjectId(company_id)}, {"$set": update_data}, return_document=True
    )
    if not result:
        raise HTTPException(status_code=404, detail='Empresa nÃ£o encontrada')
    return CompanyResponse(**doc_to_response(result))

@api_router.delete('/companies/{company_id}')
async def delete_company(company_id: str, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    result = await db.companies.delete_one({"_id": ObjectId(company_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Empresa nÃ£o encontrada')
    return {'message': 'Empresa excluÃ­da'}

# ===================== EMPLOYEES =====================

@api_router.get('/employees')
async def get_employees(current_user: dict = Depends(get_current_user), search: Optional[str] = None, company_id: Optional[str] = None, status_filter: Optional[str] = None):
    db = await get_db()
    query = {}
    if search:
        query["$or"] = [
            {"full_name": {"$regex": search, "$options": "i"}},
            {"cpf": {"$regex": search, "$options": "i"}},
            {"registration_number": {"$regex": search, "$options": "i"}}
        ]
    if company_id:
        query["company_id"] = company_id
    if status_filter:
        normalized_filter = status_filter.strip().lower()
        if normalized_filter in ['active', 'ativos', 'ativo']:
            query["$and"] = query.get("$and", []) + [{"$or": [{"status": "active"}, {"status": {"$exists": False}}, {"status": None}]}]
        elif normalized_filter in ['inactive', 'inativos', 'inativo']:
            query["status"] = "inactive"
    
    employees = await db.employees.find(query).to_list(1000)
    biometric_summaries = await get_biometric_summaries(db, employees)
    
    # Perfis que nÃ£o podem ver dados sensÃ­veis
    if not can_view_sensitive_data(current_user['role']):
        result = []
        for employee in employees:
            item = EmployeePublicResponse(**doc_to_response(employee)).model_dump()
            summary = biometric_summaries.get(str(employee['_id']))
            item["biometric"] = summary
            item["biometric_status"] = summary["status"] if summary else "missing"
            item["biometric_templates_count"] = summary["templates_count"] if summary else 0
            result.append(item)
        return result
    
    result = []
    for employee in employees:
        item = EmployeeResponse(**doc_to_response(employee)).model_dump()
        summary = biometric_summaries.get(str(employee['_id']))
        item["biometric"] = summary
        item["biometric_status"] = summary["status"] if summary else "missing"
        item["biometric_templates_count"] = summary["templates_count"] if summary else 0
        result.append(item)
    return result

@api_router.post('/employees', response_model=EmployeeResponse, status_code=status.HTTP_201_CREATED)
async def create_employee(employee_data: EmployeeCreate, current_user: dict = Depends(get_current_user)):
    if not can_manage_employees(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o para cadastrar colaboradores')
    
    db = await get_db()
    existing = await db.employees.find_one({"cpf": employee_data.cpf})
    if existing:
        raise HTTPException(status_code=400, detail='CPF jÃ¡ cadastrado')
    
    new_employee = {**employee_data.model_dump(), "created_at": datetime.now(timezone.utc), "updated_at": datetime.now(timezone.utc)}
    result = await db.employees.insert_one(new_employee)
    new_employee['_id'] = result.inserted_id
    return EmployeeResponse(**doc_to_response(new_employee))

@api_router.get('/employees/{employee_id}')
async def get_employee(employee_id: str, current_user: dict = Depends(get_current_user)):
    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')
    templates = await db.facial_templates.find({
        "employee_id": {"$in": [employee_id, ObjectId(employee_id)]}
    }).to_list(100)
    biometric_summary = build_biometric_summary(employee, templates)

    if not can_view_sensitive_data(current_user['role']):
        item = EmployeePublicResponse(**doc_to_response(employee)).model_dump()
        item["biometric"] = biometric_summary
        item["biometric_status"] = biometric_summary["status"]
        item["biometric_templates_count"] = biometric_summary["templates_count"]
        return item

    item = EmployeeResponse(**doc_to_response(employee)).model_dump()
    item["biometric"] = biometric_summary
    item["biometric_status"] = biometric_summary["status"]
    item["biometric_templates_count"] = biometric_summary["templates_count"]
    return item

@api_router.patch('/employees/{employee_id}', response_model=EmployeeResponse)
async def update_employee(employee_id: str, employee_data: EmployeeUpdate, current_user: dict = Depends(get_current_user)):
    if not can_manage_employees(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o para editar colaboradores')
    
    db = await get_db()
    update_data = {k: v for k, v in employee_data.model_dump(exclude_unset=True).items()}
    if 'status' in update_data:
        existing_employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
        if not existing_employee:
            raise HTTPException(status_code=404, detail='Colaborador nÃƒÂ£o encontrado')
        requested_status = str(update_data.get('status')).lower()
        current_status = normalize_employee_status(existing_employee)
        if requested_status != current_status:
            raise HTTPException(status_code=400, detail='Use a acao de ativar/inativar com motivo obrigatorio.')
        update_data.pop('status', None)
    update_data['updated_at'] = datetime.now(timezone.utc)
    result = await db.employees.find_one_and_update(
        {"_id": ObjectId(employee_id)}, {"$set": update_data}, return_document=True
    )
    if not result:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')
    return EmployeeResponse(**doc_to_response(result))

@api_router.post('/employees/{employee_id}/deactivate', response_model=EmployeeResponse)
async def deactivate_employee(employee_id: str, payload: EmployeeStatusChangeRequest, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_change_employee_status(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para inativar colaboradores')

    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nao encontrado')

    previous_status = normalize_employee_status(employee)
    if previous_status == 'inactive':
        return EmployeeResponse(**doc_to_response(employee))

    now = datetime.now(timezone.utc)
    reason = payload.reason.strip()
    result = await db.employees.find_one_and_update(
        {"_id": ObjectId(employee_id)},
        {"$set": {
            "status": "inactive",
            "inactivated_at": now,
            "inactivated_by": current_user.get('id'),
            "inactivation_reason": reason,
            "updated_at": now
        }},
        return_document=True
    )
    await write_audit_log(db, 'COLLABORATOR_DEACTIVATED', 'collaborator', employee_id, current_user, reason=reason, previous_status=previous_status, new_status='inactive', request=request)
    return EmployeeResponse(**doc_to_response(result))

@api_router.post('/employees/{employee_id}/reactivate', response_model=EmployeeResponse)
async def reactivate_employee(employee_id: str, payload: EmployeeStatusChangeRequest, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_change_employee_status(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para reativar colaboradores')

    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nao encontrado')

    previous_status = normalize_employee_status(employee)
    if previous_status == 'active':
        return EmployeeResponse(**doc_to_response(employee))

    now = datetime.now(timezone.utc)
    reason = payload.reason.strip()
    result = await db.employees.find_one_and_update(
        {"_id": ObjectId(employee_id)},
        {"$set": {
            "status": "active",
            "reactivated_at": now,
            "reactivated_by": current_user.get('id'),
            "reactivation_reason": reason,
            "updated_at": now
        }},
        return_document=True
    )
    await write_audit_log(db, 'COLLABORATOR_REACTIVATED', 'collaborator', employee_id, current_user, reason=reason, previous_status=previous_status, new_status='active', request=request)
    return EmployeeResponse(**doc_to_response(result))

@api_router.delete('/employees/{employee_id}')
async def delete_employee(employee_id: str, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    result = await db.employees.delete_one({"_id": ObjectId(employee_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')
    return {'message': 'Colaborador excluÃ­do'}

@api_router.post('/employees/{employee_id}/photo')
async def upload_employee_photo(employee_id: str, file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    if not can_manage_employees(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o')
    
    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')
    
    file_ext = Path(file.filename).suffix
    file_name = f'employee_{employee_id}_{datetime.now(timezone.utc).timestamp()}{file_ext}'
    file_path = UPLOAD_DIR / 'employees' / file_name
    file_path.parent.mkdir(exist_ok=True, parents=True)
    
    with file_path.open('wb') as buffer:
        shutil.copyfileobj(file.file, buffer)
    
    photo_path = f'/uploads/employees/{file_name}'
    await db.employees.update_one({"_id": ObjectId(employee_id)}, {"$set": {"photo_path": photo_path}})
    return {'photo_path': photo_path}

# ===================== IMPORTAÃ‡ÃƒO/EXPORTAÃ‡ÃƒO =====================

@api_router.get('/employees/export/excel')
async def export_employees_excel(current_user: dict = Depends(get_current_user)):
    """Exporta colaboradores para Excel"""
    if not can_view_sensitive_data(current_user['role']):
        raise HTTPException(status_code=403, detail='PermissÃ£o insuficiente')
    
    db = await get_db()
    employees = await db.employees.find().to_list(5000)
    companies = {str(c['_id']): c['legal_name'] async for c in db.companies.find()}
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Colaboradores"
    
    # CabeÃ§alho
    headers = ['Nome Completo', 'CPF', 'RG', 'MatrÃ­cula', 'Empresa', 'Cargo', 'Setor', 'Status']
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # Dados
    for row, emp in enumerate(employees, 2):
        ws.cell(row=row, column=1, value=emp.get('full_name', ''))
        ws.cell(row=row, column=2, value=emp.get('cpf', ''))
        ws.cell(row=row, column=3, value=emp.get('rg', ''))
        ws.cell(row=row, column=4, value=emp.get('registration_number', ''))
        ws.cell(row=row, column=5, value=companies.get(emp.get('company_id', ''), ''))
        ws.cell(row=row, column=6, value=emp.get('position', ''))
        ws.cell(row=row, column=7, value=emp.get('department', ''))
        ws.cell(row=row, column=8, value='Ativo' if emp.get('status') == 'active' else 'Inativo')
    
    # Ajustar largura das colunas
    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=colaboradores_{datetime.now().strftime("%Y%m%d")}.xlsx'}
    )

@api_router.get('/employees/template/excel')
async def download_employees_template(current_user: dict = Depends(get_current_user)):
    """Download do template Excel para importaÃ§Ã£o de colaboradores"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Colaboradores"
    
    # CabeÃ§alho com instruÃ§Ãµes
    headers = ['Nome Completo*', 'CPF*', 'RG', 'MatrÃ­cula*', 'Empresa*', 'Cargo', 'Setor', 'Status']
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # Exemplo de preenchimento
    ws.cell(row=2, column=1, value='JoÃ£o da Silva')
    ws.cell(row=2, column=2, value='123.456.789-00')
    ws.cell(row=2, column=3, value='12.345.678-9')
    ws.cell(row=2, column=4, value='MAT001')
    ws.cell(row=2, column=5, value='Nome da Empresa')
    ws.cell(row=2, column=6, value='Operador')
    ws.cell(row=2, column=7, value='ProduÃ§Ã£o')
    ws.cell(row=2, column=8, value='Ativo')
    
    # InstruÃ§Ãµes
    ws2 = wb.create_sheet(title="InstruÃ§Ãµes")
    ws2.cell(row=1, column=1, value="INSTRUÃ‡Ã•ES DE PREENCHIMENTO").font = Font(bold=True, size=14)
    ws2.cell(row=3, column=1, value="* Campos obrigatÃ³rios")
    ws2.cell(row=4, column=1, value="â€¢ Nome Completo: Nome completo do colaborador")
    ws2.cell(row=5, column=1, value="â€¢ CPF: Formato XXX.XXX.XXX-XX ou apenas nÃºmeros")
    ws2.cell(row=6, column=1, value="â€¢ MatrÃ­cula: CÃ³digo Ãºnico do colaborador na empresa")
    ws2.cell(row=7, column=1, value="â€¢ Empresa: Nome exato da empresa cadastrada no sistema")
    ws2.cell(row=8, column=1, value="â€¢ Status: 'Ativo' ou 'Inativo' (padrÃ£o: Ativo)")
    
    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=template_colaboradores.xlsx'}
    )

@api_router.post('/employees/import/excel')
async def import_employees_excel(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user)
):
    """Importa colaboradores de um arquivo Excel"""
    if not can_manage_employees(current_user['role']):
        raise HTTPException(status_code=403, detail='PermissÃ£o insuficiente')
    
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail='Arquivo deve ser Excel (.xlsx ou .xls)')
    
    db = await get_db()
    
    # Carregar empresas
    companies = {}
    async for company in db.companies.find():
        companies[company['legal_name'].lower().strip()] = str(company['_id'])
    
    try:
        content = await file.read()
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        
        results = {'imported': 0, 'errors': [], 'skipped': 0}
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not row or not row[0]:  # Linha vazia
                continue
            
            full_name = str(row[0]).strip() if row[0] else None
            cpf = str(row[1]).strip() if row[1] else None
            rg = str(row[2]).strip() if row[2] else None
            registration_number = str(row[3]).strip() if row[3] else None
            company_name = str(row[4]).strip().lower() if row[4] else None
            position = str(row[5]).strip() if len(row) > 5 and row[5] else None
            department = str(row[6]).strip() if len(row) > 6 and row[6] else None
            status_str = str(row[7]).strip().lower() if len(row) > 7 and row[7] else 'ativo'
            
            # ValidaÃ§Ãµes
            errors = []
            if not full_name:
                errors.append('Nome Ã© obrigatÃ³rio')
            if not cpf:
                errors.append('CPF Ã© obrigatÃ³rio')
            if not registration_number:
                errors.append('MatrÃ­cula Ã© obrigatÃ³ria')
            if not company_name:
                errors.append('Empresa Ã© obrigatÃ³ria')
            elif company_name not in companies:
                errors.append(f'Empresa "{row[4]}" nÃ£o encontrada no sistema')
            
            if errors:
                results['errors'].append({'row': row_num, 'errors': errors, 'name': full_name})
                continue
            
            # Verificar se jÃ¡ existe
            existing = await db.employees.find_one({
                "$or": [
                    {"cpf": cpf},
                    {"registration_number": registration_number, "company_id": companies.get(company_name)}
                ]
            })
            
            if existing:
                results['skipped'] += 1
                results['errors'].append({
                    'row': row_num, 
                    'errors': ['Colaborador jÃ¡ existe (CPF ou MatrÃ­cula duplicada)'],
                    'name': full_name
                })
                continue
            
            # Inserir colaborador
            employee = {
                'full_name': full_name,
                'cpf': cpf,
                'rg': rg,
                'registration_number': registration_number,
                'company_id': companies.get(company_name),
                'position': position,
                'department': department,
                'status': 'active' if status_str == 'ativo' else 'inactive',
                'facial_consent': False,
                'created_at': datetime.now(timezone.utc),
                'updated_at': datetime.now(timezone.utc)
            }
            
            await db.employees.insert_one(employee)
            results['imported'] += 1
        
        return results
        
    except Exception as e:
        logger.error(f"Erro ao importar Excel: {str(e)}")
        raise HTTPException(status_code=400, detail=f'Erro ao processar arquivo: {str(e)}')

# ===================== IMPRESSÃƒO PDF =====================

@api_router.get('/reports/employees/pdf')
async def generate_employees_pdf(current_user: dict = Depends(get_current_user)):
    """Gera relatÃ³rio PDF de colaboradores"""
    if not can_view_sensitive_data(current_user['role']):
        raise HTTPException(status_code=403, detail='PermissÃ£o insuficiente')
    
    db = await get_db()
    employees = await db.employees.find({"status": "active"}).to_list(5000)
    companies = {str(c['_id']): c['legal_name'] async for c in db.companies.find()}
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=30)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=16, alignment=1, spaceAfter=20)
    
    elements = []
    
    # TÃ­tulo
    elements.append(Paragraph("RelatÃ³rio de Colaboradores", title_style))
    elements.append(Paragraph(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 20))
    
    # Tabela
    data = [['Nome', 'CPF', 'MatrÃ­cula', 'Empresa', 'Cargo', 'Setor']]
    
    for emp in employees:
        data.append([
            emp.get('full_name', '')[:30],
            emp.get('cpf', ''),
            emp.get('registration_number', ''),
            companies.get(emp.get('company_id', ''), '')[:25],
            (emp.get('position', '') or '')[:20],
            (emp.get('department', '') or '')[:15]
        ])
    
    table = Table(data, colWidths=[120, 90, 70, 120, 100, 80])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.063, 0.725, 0.506)),  # Emerald
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.95, 0.95)]),
    ]))
    
    elements.append(table)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph(f"Total de colaboradores ativos: {len(employees)}", styles['Normal']))
    
    doc.build(elements)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type='application/pdf',
        headers={'Content-Disposition': f'attachment; filename=colaboradores_{datetime.now().strftime("%Y%m%d")}.pdf'}
    )

@api_router.get('/reports/deliveries/pdf')
async def generate_deliveries_pdf(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Gera relatório PDF de entregas"""
    if not can_view_delivery_history(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissão para consultar entregas')
    db = await get_db()
    
    query = {}
    if start_date:
        query['created_at'] = {'$gte': datetime.fromisoformat(start_date.replace('Z', '+00:00'))}
    if end_date:
        if 'created_at' not in query:
            query['created_at'] = {}
        query['created_at']['$lte'] = datetime.fromisoformat(end_date.replace('Z', '+00:00'))
    
    deliveries = await db.deliveries.find(query).sort('created_at', -1).to_list(1000)
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=30)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=16, alignment=1, spaceAfter=20)
    
    elements = []
    
    # TÃ­tulo
    elements.append(Paragraph("RelatÃ³rio de Entregas de EPIs", title_style))
    elements.append(Paragraph(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 20))
    
    # Tabela
    data = [['Data', 'Colaborador', 'Itens', 'Tipo', 'VerificaÃ§Ã£o Facial']]
    
    for delivery in deliveries:
        items_str = ', '.join([
            f"{item.get('quantity', 1)}x {item.get('epi_name', item.get('kit_name', 'Item'))}"
            for item in delivery.get('items', [])
        ])[:50]
        
        facial_match = delivery.get('facial_match_score')
        facial_str = f"{int(facial_match * 100)}%" if facial_match else 'N/A'
        
        data.append([
            format_datetime_safe(delivery.get('created_at'), '%d/%m/%Y %H:%M'),
            delivery.get('employee_name', '')[:25],
            items_str,
            'DevoluÃ§Ã£o' if delivery.get('is_return') else 'Entrega',
            facial_str
        ])
    
    table = Table(data, colWidths=[100, 150, 200, 70, 80])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.063, 0.725, 0.506)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.95, 0.95)]),
    ]))
    
    elements.append(table)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph(f"Total de entregas: {len(deliveries)}", styles['Normal']))
    
    doc.build(elements)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type='application/pdf',
        headers={'Content-Disposition': f'attachment; filename=entregas_{datetime.now().strftime("%Y%m%d")}.pdf'}
    )

@api_router.get('/reports/employee/{employee_id}/pdf')
async def generate_employee_history_pdf(employee_id: str, current_user: dict = Depends(get_current_user)):
    """Gera PDF com ficha do colaborador e histÃ³rico de entregas"""
    db = await get_db()
    
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')
    
    company = None
    if employee.get('company_id'):
        company = await db.companies.find_one({"_id": ObjectId(employee['company_id'])})
    
    deliveries = await db.deliveries.find({"employee_id": employee_id}).sort('created_at', -1).to_list(500)
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=40, bottomMargin=40)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18, alignment=1, spaceAfter=20)
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Heading2'], fontSize=14, spaceAfter=10)
    
    elements = []
    
    # CabeÃ§alho
    elements.append(Paragraph("Ficha do Colaborador", title_style))
    elements.append(Spacer(1, 10))
    
    # Dados do colaborador
    info_data = [
        ['Nome:', employee.get('full_name', '')],
        ['CPF:', employee.get('cpf', '')],
        ['RG:', employee.get('rg', '') or '-'],
        ['MatrÃ­cula:', employee.get('registration_number', '')],
        ['Empresa:', company.get('legal_name', '') if company else '-'],
        ['Cargo:', employee.get('position', '') or '-'],
        ['Setor:', employee.get('department', '') or '-'],
        ['Status:', 'Ativo' if employee.get('status') == 'active' else 'Inativo'],
    ]
    
    info_table = Table(info_data, colWidths=[100, 350])
    info_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (0, -1), 'RIGHT'),
        ('ALIGN', (1, 0), (1, -1), 'LEFT'),
    ]))
    
    elements.append(info_table)
    elements.append(Spacer(1, 30))
    
    # HistÃ³rico de entregas
    elements.append(Paragraph("HistÃ³rico de Entregas/DevoluÃ§Ãµes", subtitle_style))
    
    if deliveries:
        # ATUALIZADO: Incluir responsÃ¡vel pela entrega
        history_data = [['Data', 'Tipo', 'Itens', 'ResponsÃ¡vel', 'VerificaÃ§Ã£o']]
        
        for delivery in deliveries:
            items_str = ', '.join([
                f"{item.get('quantity', 1)}x {item.get('epi_name', item.get('kit_name', 'Item'))}"
                for item in delivery.get('items', [])
            ])[:50]
            
            facial_match = delivery.get('facial_match_score')
            facial_str = f"{int(facial_match * 100)}%" if facial_match else '-'
            
            # Nome do responsÃ¡vel pela entrega
            responsavel = get_delivery_responsible_name(delivery)
            
            history_data.append([
                format_datetime_safe(delivery.get('created_at'), '%d/%m/%Y'),
                'DevoluÃ§Ã£o' if delivery.get('is_return') else 'Entrega',
                items_str,
                responsavel[:20],  # Limitar tamanho do nome
                facial_str
            ])
        
        history_table = Table(history_data, colWidths=[65, 60, 180, 100, 55])
        history_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.063, 0.725, 0.506)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.95, 0.95)]),
        ]))
        
        elements.append(history_table)
    else:
        elements.append(Paragraph("Nenhuma entrega registrada.", styles['Normal']))
    
    elements.append(Spacer(1, 40))
    elements.append(Paragraph(f"Documento gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    
    doc.build(elements)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type='application/pdf',
        headers={'Content-Disposition': f'attachment; filename=ficha_{employee.get("registration_number", employee_id)}_{datetime.now().strftime("%Y%m%d")}.pdf'}
    )

# ===================== FACIAL TEMPLATES =====================

@api_router.get('/employees/{employee_id}/facial-templates')
async def get_facial_templates(employee_id: str, current_user: dict = Depends(get_current_user)):
    db = await get_db()
    templates = await db.facial_templates.find({
        "employee_id": {"$in": [employee_id, ObjectId(employee_id)]}
    }).to_list(100)
    return [format_facial_template_response(t) for t in templates]

@api_router.get('/employees/{employee_id}/biometric-summary')
async def get_employee_biometric_summary(employee_id: str, current_user: dict = Depends(get_current_user)):
    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nao encontrado')
    templates = await db.facial_templates.find({
        "employee_id": {"$in": [employee_id, ObjectId(employee_id)]}
    }).to_list(100)
    return build_biometric_summary(employee, templates)

@api_router.delete('/employees/{employee_id}/facial-templates')
async def delete_all_facial_templates(employee_id: str, current_user: dict = Depends(get_current_user)):
    if not can_manage_biometrics(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao')
    db = await get_db()
    result = await db.facial_templates.delete_many({
        "employee_id": {"$in": [employee_id, ObjectId(employee_id)]}
    })
    await face_recognition_service.refresh_employee_cache(db, employee_id)
    return {'message': 'Biometria facial excluida', 'deleted_count': result.deleted_count}

@api_router.post('/employees/{employee_id}/facial-templates/reprocess')
async def reprocess_facial_templates(employee_id: str, current_user: dict = Depends(get_current_user)):
    if not can_manage_biometrics(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao')
    db = await get_db()
    await face_recognition_service.refresh_employee_cache(db, employee_id)
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    templates = await db.facial_templates.find({
        "employee_id": {"$in": [employee_id, ObjectId(employee_id)]}
    }).to_list(100)
    return {
        'message': 'Templates reprocessados no cache facial',
        'summary': build_biometric_summary(employee, templates) if employee else None
    }

@api_router.get('/biometric/dashboard')
async def get_biometric_dashboard(current_user: dict = Depends(get_current_user)):
    db = await get_db()
    employees = await db.employees.find({"status": "active"}).to_list(5000)
    summaries = await get_biometric_summaries(db, employees)
    total = len(employees)
    registered = sum(1 for item in summaries.values() if item["status"] == "registered")
    incomplete = sum(1 for item in summaries.values() if item["status"] == "incomplete")
    legacy = sum(1 for item in summaries.values() if item["status"] == "legacy")
    missing = total - registered - incomplete - legacy
    coverage = round((registered / total) * 100, 1) if total else 0
    return {
        "total_employees": total,
        "registered": registered,
        "missing": missing,
        "incomplete": incomplete,
        "legacy": legacy,
        "coverage_percent": coverage
    }

# Endpoint otimizado para buscar TODOS os templates de uma vez
@api_router.get('/facial-templates/all')
async def get_all_facial_templates(
    page: Optional[int] = Query(None, ge=1),
    page_size: int = Query(100, ge=1, le=500),
    current_user: dict = Depends(require_role('admin', 'almoxarifado')),
):
    """Retorna todos os templates faciais com informaÃ§Ãµes do colaborador - otimizado para reconhecimento"""
    db = await get_db()
    
    total = await db.facial_templates.count_documents({})
    paginated = page is not None
    cursor = db.facial_templates.find({}).sort("_id", 1)
    if paginated:
        cursor = cursor.skip((page - 1) * page_size).limit(page_size)
    templates = await cursor.to_list(page_size if paginated else max(total, 1))
    
    if not templates:
        if paginated:
            return {"items": [], "page": page, "page_size": page_size, "total": total, "total_pages": (total + page_size - 1) // page_size}
        return []
    
    # Buscar todos os colaboradores de uma vez
    employee_ids = []
    for t in templates:
        emp_id = t.get('employee_id')
        if emp_id:
            # Converter para ObjectId se for string
            if isinstance(emp_id, str):
                employee_ids.append(ObjectId(emp_id))
            else:
                employee_ids.append(emp_id)
    
    employee_ids = list(set(employee_ids))
    employees = await db.employees.find({"_id": {"$in": employee_ids}}).to_list(max(len(employee_ids), 1))
    
    # Criar mapa de colaboradores
    emp_map = {str(e['_id']): doc_to_response(e) for e in employees}
    
    # Montar resposta com dados do colaborador
    result = []
    for t in templates:
        emp_id = t.get('employee_id')
        emp_id_str = str(emp_id) if emp_id else None
        if emp_id_str and emp_id_str in emp_map:
            result.append({
                'id': str(t['_id']),
                'employee_id': emp_id_str,
                'descriptor': t.get('descriptor'),
                'pose_label': t.get('pose_label'),
                'quality_score': t.get('quality_score'),
                'detection_score': t.get('detection_score'),
                'employee': emp_map[emp_id_str]
            })
    
    if paginated:
        return {
            "items": result,
            "page": page,
            "page_size": page_size,
            "total": total,
            "total_pages": (total + page_size - 1) // page_size,
        }
    return result

# ===================== VERIFICAÃ‡ÃƒO DE DUPLICIDADE BIOMÃ‰TRICA =====================

BIOMETRIC_DUPLICATE_THRESHOLD = 0.40  # DistÃ¢ncia mÃ¡xima para considerar duplicata

def euclidean_distance(a, b):
    """Calcula distÃ¢ncia euclidiana entre dois vetores"""
    import math
    return math.sqrt(sum((x - y) ** 2 for x, y in zip(a, b)))

@api_router.post('/biometric/check-duplicate')
async def check_biometric_duplicate(
    data: BiometricCheckRequest,
    current_user: dict = Depends(get_current_user)
):
    """Verifica se existe biometria duplicada no sistema"""
    db = await get_db()
    
    try:
        # Parsear o descriptor enviado
        new_descriptor = json.loads(data.descriptor)
        if len(new_descriptor) != 128:
            raise HTTPException(status_code=400, detail="Descriptor deve ter 128 valores")
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="Descriptor invÃ¡lido")
    
    # Buscar todos os templates existentes
    templates = await db.facial_templates.find({}).to_list(1000)
    
    # Buscar colaboradores
    employee_ids = list(set([t.get('employee_id') for t in templates if t.get('employee_id')]))
    employees = await db.employees.find({"_id": {"$in": [ObjectId(str(e)) for e in employee_ids]}}).to_list(1000)
    emp_map = {str(e['_id']): e for e in employees}
    
    # Comparar com cada template
    for template in templates:
        emp_id = str(template.get('employee_id'))
        
        # Ignorar o prÃ³prio colaborador (para ediÃ§Ã£o)
        if data.employee_id and emp_id == data.employee_id:
            continue
        
        try:
            existing_descriptor = json.loads(template.get('descriptor', '[]'))
            if len(existing_descriptor) != 128:
                continue
            
            # Calcular distÃ¢ncia
            distance = euclidean_distance(new_descriptor, existing_descriptor)
            similarity = 1 - distance
            
            # Se a distÃ¢ncia for menor que o threshold, Ã© duplicata
            if distance < BIOMETRIC_DUPLICATE_THRESHOLD:
                emp_data = emp_map.get(emp_id, {})
                return BiometricCheckResponse(
                    is_duplicate=True,
                    duplicate_employee_id=emp_id,
                    duplicate_employee_name=emp_data.get('full_name', 'Desconhecido'),
                    similarity_score=similarity,
                    message=f"Esta foto jÃ¡ corresponde ou estÃ¡ muito semelhante Ã  biometria facial do colaborador '{emp_data.get('full_name', 'Desconhecido')}'. NÃ£o Ã© permitido utilizar a mesma foto para colaboradores diferentes."
                )
        except Exception as e:
            continue
    
    return BiometricCheckResponse(
        is_duplicate=False,
        message="Biometria vÃ¡lida - nenhuma duplicata encontrada"
    )

@api_router.get('/biometric/audit-duplicates')
async def audit_biometric_duplicates(current_user: dict = Depends(get_current_user)):
    """Audita a base de dados para encontrar biometrias duplicadas existentes"""
    if current_user['role'] not in ['admin', 'seguranca_trabalho']:
        raise HTTPException(status_code=403, detail="Apenas administradores podem auditar duplicatas")
    
    db = await get_db()
    
    # Buscar todos os templates
    templates = await db.facial_templates.find({}).to_list(1000)
    
    # Buscar colaboradores
    employee_ids = list(set([t.get('employee_id') for t in templates if t.get('employee_id')]))
    employees = await db.employees.find({"_id": {"$in": [ObjectId(str(e)) for e in employee_ids]}}).to_list(1000)
    emp_map = {str(e['_id']): e for e in employees}
    
    duplicates = []
    checked_pairs = set()
    
    # Comparar todos os pares
    for i, t1 in enumerate(templates):
        emp_id1 = str(t1.get('employee_id'))
        
        try:
            desc1 = json.loads(t1.get('descriptor', '[]'))
            if len(desc1) != 128:
                continue
        except:
            continue
        
        for j, t2 in enumerate(templates):
            if i >= j:
                continue
            
            emp_id2 = str(t2.get('employee_id'))
            
            # NÃ£o comparar mesmo colaborador
            if emp_id1 == emp_id2:
                continue
            
            # Evitar duplicar pares
            pair_key = tuple(sorted([emp_id1, emp_id2]))
            if pair_key in checked_pairs:
                continue
            checked_pairs.add(pair_key)
            
            try:
                desc2 = json.loads(t2.get('descriptor', '[]'))
                if len(desc2) != 128:
                    continue
                
                distance = euclidean_distance(desc1, desc2)
                
                if distance < BIOMETRIC_DUPLICATE_THRESHOLD:
                    emp1 = emp_map.get(emp_id1, {})
                    emp2 = emp_map.get(emp_id2, {})
                    
                    duplicates.append({
                        "employee1_id": emp_id1,
                        "employee1_name": emp1.get('full_name', 'Desconhecido'),
                        "employee2_id": emp_id2,
                        "employee2_name": emp2.get('full_name', 'Desconhecido'),
                        "similarity_score": round(1 - distance, 4),
                        "distance": round(distance, 4)
                    })
            except:
                continue
    
    return {
        "total_templates": len(templates),
        "duplicates_found": len(duplicates),
        "duplicates": duplicates
    }

@api_router.get('/facial/audit-attempts')
async def list_facial_attempts(
    limit: int = 200,
    current_user: dict = Depends(require_role('admin', 'seguranca_trabalho'))
):
    db = await get_db()
    safe_limit = max(1, min(limit, 500))
    attempts = await db.facial_recognition_attempts.find({}).sort("created_at", -1).to_list(safe_limit)
    return [doc_to_response(item) for item in attempts]

@api_router.get('/biometric/legacy-templates')
async def list_legacy_biometric_templates(current_user: dict = Depends(require_role('admin', 'seguranca_trabalho'))):
    db = await get_db()
    templates = await db.facial_templates.find({}).to_list(5000)
    legacy = []
    employee_ids = set()

    for template in templates:
        try:
            descriptor = json.loads(template.get('descriptor') or '[]')
        except Exception:
            descriptor = []
        if len(descriptor) and len(descriptor) != 512:
            employee_id = str(template.get('employee_id')) if template.get('employee_id') else None
            if employee_id:
                employee_ids.add(employee_id)
            legacy.append({
                'id': str(template.get('_id')),
                'employee_id': employee_id,
                'dimensions': len(descriptor),
                'pose_label': template.get('pose_label'),
                'created_at': template.get('created_at'),
                'quality_score': template.get('quality_score'),
                'detection_score': template.get('detection_score'),
            })

    employee_object_ids = [ObjectId(item) for item in employee_ids if ObjectId.is_valid(item)]
    employees = await db.employees.find({"_id": {"$in": employee_object_ids}}, {"full_name": 1, "status": 1}).to_list(5000)
    employee_map = {str(item['_id']): item for item in employees}

    for item in legacy:
        employee = employee_map.get(item.get('employee_id')) or {}
        item['employee_name'] = employee.get('full_name')
        item['employee_status'] = employee.get('status')

    return {
        'total_templates': len(templates),
        'legacy_templates': len(legacy),
        'templates': legacy
    }

@api_router.get('/biometric/audit-ambiguous-templates')
async def audit_ambiguous_biometric_templates(
    min_similarity: float = 0.62,
    current_user: dict = Depends(require_role('admin', 'seguranca_trabalho'))
):
    db = await get_db()
    try:
        import numpy as np
    except Exception:
        raise HTTPException(status_code=500, detail='Biblioteca numpy indisponivel para auditoria')

    templates = await db.facial_templates.find({}).to_list(5000)
    vectors = []
    employee_ids = set()

    for template in templates:
        try:
            descriptor = json.loads(template.get('descriptor') or '[]')
            if len(descriptor) != 512:
                continue
            vector = np.asarray(descriptor, dtype=np.float32)
            norm = float(np.linalg.norm(vector))
            if norm == 0:
                continue
            employee_id = str(template.get('employee_id')) if template.get('employee_id') else None
            if not employee_id:
                continue
            employee_ids.add(employee_id)
            vectors.append({
                'template_id': str(template.get('_id')),
                'employee_id': employee_id,
                'created_at': template.get('created_at'),
                'pose_label': template.get('pose_label'),
                'vector': vector / norm,
            })
        except Exception:
            continue

    employee_object_ids = [ObjectId(item) for item in employee_ids if ObjectId.is_valid(item)]
    employees = await db.employees.find({"_id": {"$in": employee_object_ids}}, {"full_name": 1, "status": 1}).to_list(5000)
    employee_map = {str(item['_id']): item for item in employees}

    matches = []
    for index, left in enumerate(vectors):
        for right in vectors[index + 1:]:
            if left['employee_id'] == right['employee_id']:
                continue
            similarity = float(np.dot(left['vector'], right['vector']))
            if similarity >= min_similarity:
                left_employee = employee_map.get(left['employee_id'], {})
                right_employee = employee_map.get(right['employee_id'], {})
                matches.append({
                    'employee1_id': left['employee_id'],
                    'employee1_name': left_employee.get('full_name'),
                    'employee1_status': left_employee.get('status'),
                    'template1_id': left['template_id'],
                    'employee2_id': right['employee_id'],
                    'employee2_name': right_employee.get('full_name'),
                    'employee2_status': right_employee.get('status'),
                    'template2_id': right['template_id'],
                    'similarity_score': round(similarity, 4),
                })

    matches.sort(key=lambda item: item['similarity_score'], reverse=True)
    return {
        'compatible_templates': len(vectors),
        'min_similarity': min_similarity,
        'ambiguous_pairs': len(matches),
        'matches': matches[:500]
    }

@api_router.post('/employees/{employee_id}/biometric-consent')
async def register_biometric_consent(
    employee_id: str,
    consent_data: BiometricConsentRequest,
    request: Request,
    current_user: dict = Depends(get_current_user)
):
    """Registra o consentimento biomÃ©trico do colaborador (LGPD)"""
    db = await get_db()
    
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail="Colaborador nÃ£o encontrado")
    
    # Obter IP do cliente
    client_ip = consent_data.ip_address or request.client.host
    
    # Atualizar colaborador com dados de consentimento
    update_data = {
        "facial_consent": consent_data.accepted,
        "facial_consent_date": datetime.now(timezone.utc),
        "facial_consent_ip": client_ip
    }
    
    await db.employees.update_one(
        {"_id": ObjectId(employee_id)},
        {"$set": update_data}
    )
    
    # Registrar log de consentimento para auditoria
    consent_log = {
        "employee_id": ObjectId(employee_id),
        "employee_name": employee.get('full_name'),
        "accepted": consent_data.accepted,
        "ip_address": client_ip,
        "user_agent": request.headers.get("user-agent", "unknown"),
        "registered_by": current_user.get("username"),
        "created_at": datetime.now(timezone.utc)
    }
    await db.biometric_consent_logs.insert_one(consent_log)
    
    return {
        "success": True,
        "message": "Consentimento biomÃ©trico registrado com sucesso",
        "consent_date": update_data["facial_consent_date"].isoformat(),
        "ip_address": client_ip
    }

@api_router.get('/employees/{employee_id}/biometric-consent')
async def get_biometric_consent(
    employee_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Verifica o status do consentimento biomÃ©trico do colaborador"""
    db = await get_db()
    
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail="Colaborador nÃ£o encontrado")
    
    return {
        "has_consent": employee.get("facial_consent", False),
        "consent_date": employee.get("facial_consent_date"),
        "consent_ip": employee.get("facial_consent_ip")
    }

@api_router.post('/employees/{employee_id}/facial-templates')
async def create_facial_template(employee_id: str, template_data: FacialTemplateCreate, current_user: dict = Depends(get_current_user)):
    if not can_manage_biometrics(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o')
    
    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')
    
    new_template = {
        "employee_id": ObjectId(employee_id),
        "descriptor": template_data.descriptor,
        "pose_label": template_data.pose_label,
        "quality_score": template_data.quality_score,
        "detection_score": template_data.detection_score,
        "created_by": current_user.get("id"),
        "created_by_name": current_user.get("username"),
        "created_at": datetime.now(timezone.utc)
    }
    try:
        result = await db.facial_templates.insert_one(new_template)
        await face_recognition_service.refresh_employee_cache(db, employee_id)
    except Exception as error:
        logger.exception("Falha ao salvar template facial do colaborador %s: %s", employee_id, error)
        raise HTTPException(status_code=500, detail='Erro ao salvar template facial')
    
    # Retornar resposta formatada corretamente
    return {
        "id": str(result.inserted_id),
        "employee_id": employee_id,
        "descriptor": template_data.descriptor,
        "pose_label": template_data.pose_label,
        "quality_score": template_data.quality_score,
        "detection_score": template_data.detection_score,
        "created_by": new_template.get("created_by"),
        "created_by_name": new_template.get("created_by_name"),
        "created_at": new_template["created_at"].isoformat()
    }

@api_router.delete('/employees/{employee_id}/facial-templates/{template_id}')
async def delete_facial_template(employee_id: str, template_id: str, current_user: dict = Depends(get_current_user)):
    if not can_manage_biometrics(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o')
    
    db = await get_db()
    result = await db.facial_templates.delete_one({
        "_id": ObjectId(template_id),
        "employee_id": {"$in": [employee_id, ObjectId(employee_id)]}
    })
    
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Template nÃ£o encontrado')
    
    await face_recognition_service.refresh_employee_cache(db, employee_id)
    return {'message': 'Template excluÃ­do'}

@api_router.post('/facial/enroll', response_model=FacialEnrollResponse, status_code=status.HTTP_201_CREATED)
async def facial_enroll(
    payload: FacialEnrollRequest,
    current_user: dict = Depends(get_current_user)
):
    if not can_manage_biometrics(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao')

    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(payload.employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nao encontrado')

    if not face_recognition_service.available:
        raise HTTPException(status_code=503, detail='ServiÃ§o facial backend indisponÃ­vel')

    try:
        enroll_result = face_recognition_service.enroll_from_image(
            payload.image_base64,
            pose_label=payload.pose_label or 'frontal'
        )
        if enroll_result.get('status') != 'ok':
            logger.warning("Falha no enroll facial do colaborador %s: %s", payload.employee_id, enroll_result.get('message'))
            raise HTTPException(status_code=400, detail=enroll_result.get('message', 'Falha ao cadastrar biometria facial'))
    except HTTPException:
        raise
    except Exception as error:
        logger.exception("Erro ao gerar embedding facial do colaborador %s: %s", payload.employee_id, error)
        raise HTTPException(status_code=500, detail='Erro ao processar biometria facial')

    descriptor_json = json.dumps(enroll_result['embedding'].tolist())
    new_template = {
        "employee_id": ObjectId(payload.employee_id),
        "descriptor": descriptor_json,
        "pose_label": enroll_result.get('pose_label') or payload.pose_label or 'frontal',
        "quality_score": enroll_result.get('quality_score'),
        "detection_score": enroll_result.get('detection_score'),
        "created_by": current_user.get("id"),
        "created_by_name": current_user.get("username"),
        "created_at": datetime.now(timezone.utc)
    }
    try:
        result = await db.facial_templates.insert_one(new_template)
        await face_recognition_service.refresh_employee_cache(db, payload.employee_id)
    except Exception as error:
        logger.exception("Falha ao gravar enroll facial do colaborador %s: %s", payload.employee_id, error)
        raise HTTPException(status_code=500, detail='Erro ao salvar template facial')

    return FacialEnrollResponse(
        status='ok',
        message=enroll_result.get('message', 'Biometria facial cadastrada com sucesso.'),
        employee_id=payload.employee_id,
        template_id=str(result.inserted_id),
        pose_label=new_template['pose_label'],
        quality_score=new_template['quality_score'],
        detection_score=new_template['detection_score']
    )

@api_router.post('/facial/identify-fast', response_model=FacialIdentifyFastResponse)
async def facial_identify_fast(
    payload: FacialIdentifyFastRequest,
    current_user: dict = Depends(get_current_user)
):
    if not can_deliver_epi(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para reconhecimento facial de entrega')
    if not face_recognition_service.available:
        raise HTTPException(status_code=503, detail='Serviço facial backend indisponível')

    attempt_id = str(uuid.uuid4())
    started_at = time.perf_counter()
    db = await get_db()
    result = face_recognition_service.identify_fast(payload.image_base64)
    frame_telemetry = result.pop("_telemetry", {})
    frame_telemetry.update({
        "attempt_id": attempt_id,
        "frame_index": 0,
    })
    result["message"] = friendly_facial_message(result)
    result = await block_inactive_identified_employee(db, result)
    await save_facial_attempt_log(
        db, result, current_user, '/facial/identify-fast', frame_count=1,
        attempt_id=attempt_id, frames_captured=1, frames_sent=1,
        frame_telemetry=[frame_telemetry],
        backend_total_ms=round((time.perf_counter() - started_at) * 1000, 3),
    )
    return FacialIdentifyFastResponse(**result)


def resolve_facial_burst_consensus(results):
    approved = [
        item for item in results
        if item.get("status") == "approved" and item.get("employee_id")
    ]
    by_employee = {}
    for item in approved:
        by_employee.setdefault(item["employee_id"], []).append(item)
    if len(by_employee) > 1:
        best = max(approved, key=lambda item: item.get("similarity_score", 0))
        return {"status": "conflict", "best": best}
    if len(by_employee) == 1:
        employee_items = next(iter(by_employee.values()))
        if len(employee_items) >= 2:
            return {"status": "approved", "best": max(employee_items, key=lambda item: item.get("similarity_score", 0))}
    return {"status": "retry", "best": max(results, key=lambda item: item.get("similarity_score", 0), default=None)}


@api_router.post('/facial/identify-burst', response_model=FacialIdentifyFastResponse)
async def facial_identify_burst(
    payload: FacialIdentifyBurstRequest,
    current_user: dict = Depends(get_current_user)
):
    if not can_deliver_epi(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para reconhecimento facial de entrega')
    if not face_recognition_service.available:
        raise HTTPException(status_code=503, detail='Serviço facial backend indisponível')

    client_telemetry = sanitize_client_telemetry(payload.client_telemetry)
    attempt_id = normalized_telemetry_attempt_id(client_telemetry) or str(uuid.uuid4())
    started_at = time.perf_counter()
    captured_count = len(payload.images_base64 or [])
    frames = [image for image in payload.images_base64[:5] if image]
    frame_telemetry = []
    db = await get_db()

    async def persist(result: dict):
        result["message"] = friendly_facial_message(result)
        await save_facial_attempt_log(
            db, result, current_user, '/facial/identify-burst',
            frame_count=len(frames),
            attempt_id=attempt_id,
            frames_captured=captured_count,
            frames_sent=len(frames),
            frame_telemetry=frame_telemetry,
            backend_total_ms=round((time.perf_counter() - started_at) * 1000, 3),
            client_telemetry=client_telemetry,
        )

    if not frames:
        result = {
            'status': 'retry',
            'message': 'Nenhuma imagem utilizavel recebida. Nova captura automatica em andamento.',
            'employee_id': None,
            'employee_name': None,
            'similarity_score': 0.0,
            'detection_confidence': 0.0,
            'liveness_required': False
        }
        await persist(result)
        return FacialIdentifyFastResponse(**result)

    results = []
    for frame_index, image in enumerate(frames):
        result = face_recognition_service.identify_fast(image)
        telemetry = result.pop("_telemetry", {})
        telemetry.update({
            "attempt_id": attempt_id,
            "frame_index": frame_index,
        })
        frame_telemetry.append(telemetry)
        results.append(result)

        blocked_multi_face = (
            result.get("status") == "blocked"
            and "rosto" in (result.get("message") or "").lower()
        )
        if blocked_multi_face:
            result["message"] = "Mais de uma pessoa detectada. Fique sozinho diante da câmera."
            result["employee_id"] = None
            result["employee_name"] = None
            result["similarity_score"] = 0.0
            result["liveness_required"] = False
            await persist(result)
            return FacialIdentifyFastResponse(**result)

        consensus = resolve_facial_burst_consensus(results)
        if consensus["status"] == "conflict":
            conflict = consensus["best"]
            conflict.update({
                "status": "retry",
                "employee_id": None,
                "employee_name": None,
                "message": "Não foi possível confirmar a mesma pessoa. Tente novamente.",
                "liveness_required": False,
            })
            await persist(conflict)
            return FacialIdentifyFastResponse(**conflict)
        if consensus["status"] == "approved":
            best = consensus["best"]
            best["message"] = f"{best.get('employee_name')} identificado automaticamente."
            best["liveness_required"] = False
            best = await block_inactive_identified_employee(db, best)
            await persist(best)
            return FacialIdentifyFastResponse(**best)

    usable = [item for item in results if item.get('detection_confidence', 0) > 0]
    best_retry = max(usable, key=lambda item: item.get('similarity_score', 0), default=None)
    if best_retry:
        best_retry.update({
            'status': 'retry',
            'employee_id': None,
            'employee_name': None,
            'message': 'Reconhecimento facial não confirmado. Posicione-se novamente diante da câmera.',
            'liveness_required': False
        })
        await persist(best_retry)
        return FacialIdentifyFastResponse(**best_retry)

    result = {
        'status': 'retry',
        'message': 'Nenhum frame utilizavel. Nova captura automatica em andamento.',
        'employee_id': None,
        'employee_name': None,
        'similarity_score': 0.0,
        'detection_confidence': 0.0,
        'liveness_required': False
    }
    await persist(result)
    return FacialIdentifyFastResponse(**result)


@api_router.post('/admin/facial-benchmark/phase-10-3/session')
async def phase103_start(
    payload: Phase103SessionRequest,
    current_user: dict = Depends(require_role('admin')),
):
    return phase103_start_session(current_user['id'], payload.consent)


@api_router.get('/admin/facial-benchmark/phase-10-3/session')
async def phase103_resume(
    current_user: dict = Depends(require_role('admin')),
):
    return phase103_resume_session(current_user['id'])


@api_router.post('/admin/facial-benchmark/phase-10-3/reference')
async def phase103_reference(
    payload: Phase103ReferenceRequest,
    current_user: dict = Depends(require_role('admin')),
):
    return phase103_add_reference(current_user['id'], payload)


@api_router.post('/admin/facial-benchmark/phase-10-3/compare')
async def phase103_compare_endpoint(
    payload: Phase103CompareRequest,
    current_user: dict = Depends(require_role('admin')),
):
    return phase103_compare(current_user['id'], payload)


@api_router.post('/admin/facial-benchmark/phase-10-3/close')
async def phase103_close(
    payload: Phase103CloseRequest,
    current_user: dict = Depends(require_role('admin')),
):
    return phase103_close_session(current_user['id'], payload)


@api_router.post('/facial/performance')
async def facial_performance_update(
    payload: FacialPerformanceUpdateRequest,
    current_user: dict = Depends(get_current_user)
):
    if not can_deliver_epi(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para telemetria facial de entrega')

    attempt_id = normalized_telemetry_attempt_id({'attempt_id': payload.attempt_id})
    if not attempt_id:
        raise HTTPException(status_code=422, detail='Identificador tecnico de tentativa invalido')

    update = sanitize_client_telemetry(payload.model_dump())
    update.pop('attempt_id', None)
    if not update:
        raise HTTPException(status_code=422, detail='Payload de telemetria sem metricas')

    db = await get_db()
    try:
        result = await db.facial_recognition_attempts.update_one(
            {'attempt_id': attempt_id, 'user_id': current_user.get('id')},
            {'$set': {f'client_telemetry.{key}': value for key, value in update.items()}},
        )
    except Exception:
        logger.exception('Falha tecnica ao atualizar telemetria facial')
        raise HTTPException(status_code=503, detail='Telemetria temporariamente indisponivel')

    if result.matched_count == 0:
        logger.warning('Tentativa de telemetria facial nao encontrada para o usuario autenticado')
        raise HTTPException(status_code=404, detail='Tentativa de telemetria nao encontrada')

    return {'status': 'ok', 'updated': result.modified_count > 0}


@api_router.post('/facial/liveness-check', response_model=FacialLivenessResponse)
async def facial_liveness_check(
    payload: FacialLivenessRequest,
    current_user: dict = Depends(get_current_user)
):
    if not can_deliver_epi(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para prova de vida')

    if not face_recognition_service.available:
        raise HTTPException(status_code=503, detail='ServiÃ§o facial backend indisponÃ­vel')

    result = face_recognition_service.liveness_check(
        payload.image_base64,
        previous_image_base64=payload.previous_image_base64
    )
    return FacialLivenessResponse(**result)

@api_router.get('/facial/migration-status', response_model=FacialMigrationStatusResponse)
async def facial_migration_status(current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    status_data = await face_recognition_service.get_migration_status(db)
    return FacialMigrationStatusResponse(**status_data)

@api_router.post('/facial/reload-cache')
async def facial_reload_cache(current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    await face_recognition_service.reload_cache(db)
    return {
        'message': 'Cache facial recarregado com sucesso',
        'service_available': face_recognition_service.available,
        'cache_size': face_recognition_service.cache_size,
        'employees_loaded': face_recognition_service.cached_employee_count,
    }

# ===================== SUPPLIERS =====================

@api_router.get('/suppliers', response_model=List[SupplierResponse])
async def get_suppliers(current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    suppliers = await db.suppliers.find({}).to_list(1000)
    return [SupplierResponse(**doc_to_response(s)) for s in suppliers]

@api_router.post('/suppliers', response_model=SupplierResponse, status_code=status.HTTP_201_CREATED)
async def create_supplier(supplier_data: SupplierCreate, current_user: dict = Depends(require_role('admin', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    new_supplier = {**supplier_data.model_dump(), "created_at": datetime.now(timezone.utc)}
    result = await db.suppliers.insert_one(new_supplier)
    new_supplier['_id'] = result.inserted_id
    return SupplierResponse(**doc_to_response(new_supplier))

@api_router.patch('/suppliers/{supplier_id}', response_model=SupplierResponse)
async def update_supplier(supplier_id: str, supplier_data: SupplierUpdate, current_user: dict = Depends(require_role('admin', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    update_data = {k: v for k, v in supplier_data.model_dump(exclude_unset=True).items()}
    update_data['updated_at'] = datetime.now(timezone.utc)
    result = await db.suppliers.find_one_and_update(
        {"_id": ObjectId(supplier_id)}, {"$set": update_data}, return_document=True
    )
    if not result:
        raise HTTPException(status_code=404, detail='Fornecedor nÃ£o encontrado')
    return SupplierResponse(**doc_to_response(result))

@api_router.delete('/suppliers/{supplier_id}')
async def delete_supplier(supplier_id: str, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    result = await db.suppliers.delete_one({"_id": ObjectId(supplier_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Fornecedor nÃ£o encontrado')
    return {'message': 'Fornecedor excluÃ­do'}

# ===================== COURSES =====================

@api_router.get('/courses', response_model=List[CourseResponse])
async def get_courses(current_user: dict = Depends(get_current_user)):
    if not can_view_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    courses = await db.courses.find({}).sort('name', 1).to_list(1000)
    result = []
    for course in courses:
        result.append(CourseResponse(**(await build_course_response(db, course))))
    return result


@api_router.get('/courses/{course_id}', response_model=CourseResponse)
async def get_course(course_id: str, current_user: dict = Depends(get_current_user)):
    if not can_view_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    course = await db.courses.find_one({"_id": parse_object_id_or_400(course_id, 'course_id')})
    if not course:
        raise HTTPException(status_code=404, detail='Curso nao encontrado')
    return CourseResponse(**(await build_course_response(db, course)))


@api_router.post('/courses', response_model=CourseResponse, status_code=status.HTTP_201_CREATED)
async def create_course(course_data: CourseCreate, current_user: dict = Depends(get_current_user)):
    if not can_manage_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    normalized_name = normalize_course_name(course_data.name)
    existing = await db.courses.find_one({"name_normalized": normalized_name})
    if existing:
        raise HTTPException(status_code=409, detail='Ja existe um curso com este nome')

    epi_ids = await validate_epi_ids_exist(db, course_data.epi_ids)
    now = datetime.now(timezone.utc)
    new_course = {
        "name": normalize_text(course_data.name),
        "name_normalized": normalized_name,
        "description": normalize_text(course_data.description),
        "active": course_data.active,
        "block_epi_delivery": course_data.block_epi_delivery,
        "alert_days_before_expiration": course_data.alert_days_before_expiration,
        "created_at": now,
        "updated_at": now,
        "created_by": current_user.get('username'),
        "updated_by": current_user.get('username')
    }

    try:
        result = await db.courses.insert_one(new_course)
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail='Ja existe um curso com este nome')

    new_course['_id'] = result.inserted_id
    if epi_ids:
        await replace_course_epi_links(db, str(result.inserted_id), epi_ids, current_user.get('username'))

    course = await db.courses.find_one({"_id": result.inserted_id})
    return CourseResponse(**(await build_course_response(db, course)))


@api_router.patch('/courses/{course_id}', response_model=CourseResponse)
async def update_course(course_id: str, course_data: CourseUpdate, current_user: dict = Depends(get_current_user)):
    if not can_manage_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    course_object_id = parse_object_id_or_400(course_id, 'course_id')
    existing = await db.courses.find_one({"_id": course_object_id})
    if not existing:
        raise HTTPException(status_code=404, detail='Curso nao encontrado')

    update_data = {}
    payload = course_data.model_dump(exclude_unset=True)

    if 'name' in payload:
        normalized_name = normalize_course_name(payload['name'])
        duplicate = await db.courses.find_one({
            "name_normalized": normalized_name,
            "_id": {"$ne": course_object_id}
        })
        if duplicate:
            raise HTTPException(status_code=409, detail='Ja existe um curso com este nome')
        update_data['name'] = normalize_text(payload['name'])
        update_data['name_normalized'] = normalized_name

    if 'description' in payload:
        update_data['description'] = normalize_text(payload['description'])
    if 'active' in payload:
        update_data['active'] = payload['active']
    if 'block_epi_delivery' in payload:
        update_data['block_epi_delivery'] = payload['block_epi_delivery']
    if 'alert_days_before_expiration' in payload:
        update_data['alert_days_before_expiration'] = payload['alert_days_before_expiration']

    update_data['updated_at'] = datetime.now(timezone.utc)
    update_data['updated_by'] = current_user.get('username')

    try:
        await db.courses.update_one({"_id": course_object_id}, {"$set": update_data})
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail='Ja existe um curso com este nome')

    if payload.get('epi_ids') is not None:
        await replace_course_epi_links(db, course_id, payload.get('epi_ids') or [], current_user.get('username'))

    course = await db.courses.find_one({"_id": course_object_id})
    return CourseResponse(**(await build_course_response(db, course)))


@api_router.get('/courses/{course_id}/epis', response_model=List[CourseLinkedEPI])
async def get_course_linked_epis_endpoint(course_id: str, current_user: dict = Depends(get_current_user)):
    if not can_view_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    course = await db.courses.find_one({"_id": parse_object_id_or_400(course_id, 'course_id')})
    if not course:
        raise HTTPException(status_code=404, detail='Curso nao encontrado')
    linked_epis = await get_course_linked_epis(db, course_id)
    return [CourseLinkedEPI(**item) for item in linked_epis]


@api_router.post('/courses/{course_id}/epis', response_model=CourseEpiLinkResponse, status_code=status.HTTP_201_CREATED)
async def create_course_epi_link(course_id: str, link_data: CourseEpiLinkCreate, current_user: dict = Depends(get_current_user)):
    if not can_manage_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    course_object_id = parse_object_id_or_400(course_id, 'course_id')
    epi_object_id = parse_object_id_or_400(link_data.epi_id, 'epi_id')
    course = await db.courses.find_one({"_id": course_object_id})
    if not course:
        raise HTTPException(status_code=404, detail='Curso nao encontrado')
    epi = await db.epis.find_one({"_id": epi_object_id})
    if not epi:
        raise HTTPException(status_code=404, detail='EPI nao encontrado')

    new_link = {
        "course_id": course_id,
        "epi_id": link_data.epi_id,
        "created_at": datetime.now(timezone.utc),
        "created_by": current_user.get('username')
    }
    try:
        result = await db.course_epi_links.insert_one(new_link)
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail='Este vinculo Curso x EPI ja existe')
    new_link['_id'] = result.inserted_id
    return CourseEpiLinkResponse(**doc_to_response(new_link))


@api_router.delete('/courses/{course_id}/epis/{epi_id}')
async def delete_course_epi_link(course_id: str, epi_id: str, current_user: dict = Depends(get_current_user)):
    if not can_manage_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Permissao insuficiente')

    db = await get_db()
    course = await db.courses.find_one({"_id": parse_object_id_or_400(course_id, 'course_id')})
    if not course:
        raise HTTPException(status_code=404, detail='Curso nao encontrado')

    parse_object_id_or_400(epi_id, 'epi_id')
    result = await db.course_epi_links.delete_one({"course_id": course_id, "epi_id": epi_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Vinculo Curso x EPI nao encontrado')
    return {'message': 'Vinculo removido com sucesso'}



@api_router.get('/employees/{employee_id}/courses', response_model=List[EmployeeCourseResponse])
async def get_employee_course_records(employee_id: str, current_user: dict = Depends(get_current_user)):
    if not can_view_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para consultar certificados do colaborador')
    db = await get_db()
    await ensure_employee_exists(db, employee_id)
    records = await build_employee_course_responses(db, employee_id)
    return [EmployeeCourseResponse(**record) for record in records]


@api_router.get('/employees/{employee_id}/courses/{employee_course_id}', response_model=EmployeeCourseResponse)
async def get_employee_course_record(employee_id: str, employee_course_id: str, current_user: dict = Depends(get_current_user)):
    if not can_view_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para consultar certificados do colaborador')
    db = await get_db()
    await ensure_employee_exists(db, employee_id)
    record = await get_employee_course_response_by_id(db, employee_id, employee_course_id)
    return EmployeeCourseResponse(**record)


@api_router.post('/employees/{employee_id}/courses', response_model=EmployeeCourseResponse, status_code=status.HTTP_201_CREATED)
async def create_employee_course_record(employee_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_manage_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para cadastrar certificados do colaborador')
    db = await get_db()
    employee = await ensure_employee_exists(db, employee_id)
    await ensure_employee_course_write_allowed(employee, current_user)
    payload, certificate_file = await parse_employee_course_request(request)
    try:
        employee_course_data = EmployeeCourseCreate(**payload)
    except ValidationError as exc:
        raise HTTPException(status_code=422, detail=[{'loc': list(err.get('loc', [])), 'msg': err.get('msg', 'Valor invalido'), 'type': err.get('type', 'value_error')} for err in exc.errors()])
    course = await ensure_course_exists(db, employee_course_data.course_id, require_active=True)
    now = datetime.now(timezone.utc)
    document = {
        **employee_course_data.model_dump(),
        'employee_id': employee_id,
        'active': True,
        'created_at': now,
        'updated_at': now,
        'created_by': current_user.get('username'),
        'updated_by': current_user.get('username'),
        'renewed_from_id': None,
        'renewed_into_id': None,
        'deactivated_at': None,
        'deactivated_by': None,
        'deactivation_reason': None,
    }
    if certificate_file:
        document.update(await save_employee_course_file(employee_id, employee_course_data.course_id, certificate_file))
    result = await db.employee_courses.insert_one(document)
    await write_audit_log(db, 'EMPLOYEE_COURSE_CREATED', 'employee_course', str(result.inserted_id), current_user, request=request, metadata={
        'employee_id': employee_id,
        'employee_name': employee.get('full_name'),
        'course_id': employee_course_data.course_id,
        'course_name': course.get('name'),
    })
    record = await get_employee_course_response_by_id(db, employee_id, str(result.inserted_id))
    return EmployeeCourseResponse(**record)


@api_router.patch('/employees/{employee_id}/courses/{employee_course_id}', response_model=EmployeeCourseResponse)
async def update_employee_course_record(employee_id: str, employee_course_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_manage_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para editar certificados do colaborador')
    db = await get_db()
    employee = await ensure_employee_exists(db, employee_id)
    await ensure_employee_course_write_allowed(employee, current_user)
    existing = await db.employee_courses.find_one({'_id': parse_object_id_or_400(employee_course_id, 'employee_course_id')})
    if not existing or existing.get('employee_id') != employee_id:
        raise HTTPException(status_code=404, detail='Certificado do colaborador nao encontrado')
    if not existing.get('active', True):
        raise HTTPException(status_code=400, detail='Certificado inativo nao pode ser editado')
    payload, certificate_file = await parse_employee_course_request(request)
    try:
        employee_course_data = EmployeeCourseUpdate(**payload)
    except ValidationError as exc:
        raise HTTPException(status_code=422, detail=[{'loc': list(err.get('loc', [])), 'msg': err.get('msg', 'Valor invalido'), 'type': err.get('type', 'value_error')} for err in exc.errors()])
    update_data = employee_course_data.model_dump(exclude_unset=True)
    effective_course_id = update_data.get('course_id') or existing.get('course_id')
    await ensure_course_exists(db, effective_course_id, require_active=True)
    ensure_employee_course_dates(update_data.get('start_date') or existing.get('start_date'), update_data.get('expiration_date') or existing.get('expiration_date'))
    if certificate_file:
        update_data.update(await save_employee_course_file(employee_id, effective_course_id, certificate_file))
    if not update_data:
        raise HTTPException(status_code=400, detail='Nenhum campo informado para atualizacao')
    update_data['updated_at'] = datetime.now(timezone.utc)
    update_data['updated_by'] = current_user.get('username')
    await db.employee_courses.update_one({'_id': ObjectId(employee_course_id)}, {'$set': update_data})
    await write_audit_log(db, 'EMPLOYEE_COURSE_UPDATED', 'employee_course', employee_course_id, current_user, request=request, metadata={
        'employee_id': employee_id,
        'employee_name': employee.get('full_name'),
        'course_id': effective_course_id,
        'updated_fields': sorted(update_data.keys()),
    })
    record = await get_employee_course_response_by_id(db, employee_id, employee_course_id)
    return EmployeeCourseResponse(**record)


@api_router.post('/employees/{employee_id}/courses/{employee_course_id}/renew', response_model=EmployeeCourseResponse, status_code=status.HTTP_201_CREATED)
async def renew_employee_course_record(employee_id: str, employee_course_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_manage_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para renovar certificados do colaborador')
    db = await get_db()
    employee = await ensure_employee_exists(db, employee_id)
    await ensure_employee_course_write_allowed(employee, current_user)
    existing = await db.employee_courses.find_one({'_id': parse_object_id_or_400(employee_course_id, 'employee_course_id')})
    if not existing or existing.get('employee_id') != employee_id:
        raise HTTPException(status_code=404, detail='Certificado do colaborador nao encontrado')
    if not existing.get('active', True):
        raise HTTPException(status_code=400, detail='Certificado inativo nao pode ser renovado')
    payload, certificate_file = await parse_employee_course_request(request)
    try:
        renew_data = EmployeeCourseRenew(**payload)
    except ValidationError as exc:
        raise HTTPException(status_code=422, detail=[{'loc': list(err.get('loc', [])), 'msg': err.get('msg', 'Valor invalido'), 'type': err.get('type', 'value_error')} for err in exc.errors()])
    course = await ensure_course_exists(db, existing.get('course_id'), require_active=True)
    now = datetime.now(timezone.utc)
    new_document = {
        'employee_id': employee_id,
        'course_id': existing.get('course_id'),
        'start_date': renew_data.start_date,
        'expiration_date': renew_data.expiration_date,
        'certificate_number': renew_data.certificate_number,
        'training_provider': renew_data.training_provider,
        'notes': renew_data.notes,
        'active': True,
        'created_at': now,
        'updated_at': now,
        'created_by': current_user.get('username'),
        'updated_by': current_user.get('username'),
        'renewed_from_id': employee_course_id,
        'renewed_into_id': None,
        'deactivated_at': None,
        'deactivated_by': None,
        'deactivation_reason': None,
    }
    if certificate_file:
        new_document.update(await save_employee_course_file(employee_id, existing.get('course_id'), certificate_file))
    result = await db.employee_courses.insert_one(new_document)
    await db.employee_courses.update_one({'_id': ObjectId(employee_course_id)}, {'$set': {
        'renewed_into_id': str(result.inserted_id),
        'updated_at': now,
        'updated_by': current_user.get('username')
    }})
    await write_audit_log(db, 'EMPLOYEE_COURSE_RENEWED', 'employee_course', str(result.inserted_id), current_user, request=request, metadata={
        'employee_id': employee_id,
        'employee_name': employee.get('full_name'),
        'course_id': existing.get('course_id'),
        'course_name': course.get('name'),
        'previous_employee_course_id': employee_course_id,
    })
    record = await get_employee_course_response_by_id(db, employee_id, str(result.inserted_id))
    return EmployeeCourseResponse(**record)


@api_router.post('/employees/{employee_id}/courses/{employee_course_id}/deactivate', response_model=EmployeeCourseResponse)
async def deactivate_employee_course_record(employee_id: str, employee_course_id: str, payload: EmployeeCourseDeactivateRequest, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_deactivate_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para inativar certificados do colaborador')
    db = await get_db()
    employee = await ensure_employee_exists(db, employee_id)
    existing = await db.employee_courses.find_one({'_id': parse_object_id_or_400(employee_course_id, 'employee_course_id')})
    if not existing or existing.get('employee_id') != employee_id:
        raise HTTPException(status_code=404, detail='Certificado do colaborador nao encontrado')
    if not existing.get('active', True):
        raise HTTPException(status_code=400, detail='Certificado ja esta inativo')
    now = datetime.now(timezone.utc)
    await db.employee_courses.update_one({'_id': ObjectId(employee_course_id)}, {'$set': {
        'active': False,
        'deactivated_at': now,
        'deactivated_by': current_user.get('username'),
        'deactivation_reason': payload.reason,
        'updated_at': now,
        'updated_by': current_user.get('username')
    }})
    await write_audit_log(db, 'EMPLOYEE_COURSE_DEACTIVATED', 'employee_course', employee_course_id, current_user, reason=payload.reason, previous_status='active', new_status='inactive', request=request, metadata={
        'employee_id': employee_id,
        'employee_name': employee.get('full_name'),
        'course_id': existing.get('course_id'),
    })
    record = await get_employee_course_response_by_id(db, employee_id, employee_course_id)
    return EmployeeCourseResponse(**record)


# ===================== EPIS =====================

def normalize_text(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None

def normalize_course_name(value: Optional[str]) -> Optional[str]:
    text = normalize_text(value)
    if not text:
        return None
    collapsed = ' '.join(text.split()).lower()
    normalized = unicodedata.normalize('NFKD', collapsed)
    return ''.join(char for char in normalized if not unicodedata.combining(char))

def parse_object_id_or_400(value: str, field_name: str = 'id') -> ObjectId:
    if not ObjectId.is_valid(value):
        raise HTTPException(status_code=400, detail=f'{field_name} invalido')
    return ObjectId(value)

def normalize_group_text(value: Optional[str]) -> str:
    text = normalize_text(value) or ''
    return ' '.join(text.lower().split())

def normalize_size_value(value: Optional[str]) -> Optional[str]:
    text = normalize_text(value)
    return text.upper() if text else None

def strip_size_from_epi_name(name: Optional[str], size: Optional[str]) -> str:
    base = normalize_text(name) or ''
    normalized_size = normalize_text(size)
    if not base or not normalized_size:
        return base

    candidates = [
        f" tamanho {normalized_size}",
        f" tam {normalized_size}",
        f" tam. {normalized_size}",
        f" size {normalized_size}",
    ]
    lowered = base.lower()
    for suffix in candidates:
        if lowered.endswith(suffix.lower()):
            return base[:len(base) - len(suffix)].strip()

    parts = base.split()
    if parts and parts[-1].upper() == normalized_size.upper():
        return ' '.join(parts[:-1]).strip() or base
    return base

async def build_course_linked_epi(db, epi_id: str) -> Optional[dict]:
    if not ObjectId.is_valid(str(epi_id)):
        return None
    epi = await db.epis.find_one({"_id": ObjectId(epi_id)})
    if not epi:
        return None
    return {
        "epi_id": str(epi["_id"]),
        "name": epi.get('epi_group_name') or epi.get('name') or 'EPI sem nome',
        "category": epi.get('category'),
        "type_category": epi.get('type_category') or epi.get('category'),
        "epi_group_key": epi.get('epi_group_key'),
        "epi_group_name": epi.get('epi_group_name') or epi.get('name')
    }

async def get_course_linked_epis(db, course_id: str) -> List[dict]:
    links = await db.course_epi_links.find({"course_id": course_id}).to_list(1000)
    linked_epis = []
    for link in links:
        linked_epi = await build_course_linked_epi(db, link.get('epi_id'))
        if linked_epi:
            linked_epis.append(linked_epi)
    linked_epis.sort(key=lambda item: (item.get('name') or '').lower())
    return linked_epis

async def build_course_response(db, course: dict) -> dict:
    response = doc_to_response(course)
    linked_epis = await get_course_linked_epis(db, str(course['_id']))
    response['linked_epis'] = linked_epis
    response['linked_epis_count'] = len(linked_epis)
    return response

async def validate_epi_ids_exist(db, epi_ids: List[str]) -> List[str]:
    normalized_ids = []
    seen = set()
    for epi_id in epi_ids:
        if not epi_id or epi_id in seen:
            continue
        seen.add(epi_id)
        parse_object_id_or_400(epi_id, 'epi_id')
        epi = await db.epis.find_one({"_id": ObjectId(epi_id)})
        if not epi:
            raise HTTPException(status_code=404, detail=f'EPI nao encontrado: {epi_id}')
        normalized_ids.append(epi_id)
    return normalized_ids

async def replace_course_epi_links(db, course_id: str, epi_ids: List[str], created_by: Optional[str]) -> None:
    normalized_ids = await validate_epi_ids_exist(db, epi_ids)
    await db.course_epi_links.delete_many({"course_id": course_id})
    if not normalized_ids:
        return
    now = datetime.now(timezone.utc)
    await db.course_epi_links.insert_many([
        {
            "course_id": course_id,
            "epi_id": epi_id,
            "created_at": now,
            "created_by": created_by
        }
        for epi_id in normalized_ids
    ])



def normalize_optional_text(value: Optional[str]) -> Optional[str]:
    normalized = normalize_text(value)
    return normalized if normalized else None


def ensure_employee_course_dates(start_date: Optional[datetime], expiration_date: Optional[datetime]) -> None:
    if start_date and expiration_date and ensure_aware_datetime(expiration_date) < ensure_aware_datetime(start_date):
        raise HTTPException(status_code=422, detail='Data de validade nao pode ser anterior a data de inicio')


def calculate_employee_course_status(employee_course: dict, course: Optional[dict]) -> tuple[str, str, Optional[int]]:
    if not employee_course.get('active', True):
        return 'inactive', 'Inativo', None
    expiration_date = ensure_aware_datetime(employee_course.get('expiration_date'))
    if not expiration_date:
        return 'unknown', 'Sem validade', None
    days_until_expiration = (expiration_date.date() - datetime.now(timezone.utc).date()).days
    alert_days = parse_int_safe((course or {}).get('alert_days_before_expiration'), 0)
    if days_until_expiration < 0:
        return 'expired', 'Vencido', days_until_expiration
    if alert_days > 0 and days_until_expiration <= alert_days:
        return 'expiring', 'A vencer', days_until_expiration
    return 'valid', 'Valido', days_until_expiration


def pick_current_employee_course(records: List[dict]) -> Optional[dict]:
    active_records = [item for item in records if item.get('active', True)]
    if not active_records:
        return None
    ranked = sorted(
        active_records,
        key=lambda item: (
            ensure_aware_datetime(item.get('expiration_date')) or datetime.min.replace(tzinfo=timezone.utc),
            ensure_aware_datetime(item.get('start_date')) or datetime.min.replace(tzinfo=timezone.utc),
            ensure_aware_datetime(item.get('updated_at')) or datetime.min.replace(tzinfo=timezone.utc),
            ensure_aware_datetime(item.get('created_at')) or datetime.min.replace(tzinfo=timezone.utc),
        ),
        reverse=True,
    )
    return ranked[0]


def serialize_date_for_api(value: Optional[datetime]) -> Optional[str]:
    parsed = ensure_aware_datetime(value)
    return parsed.date().isoformat() if parsed else None


def get_course_block_status_from_record(record: Optional[dict]) -> str:
    if not record:
        return 'missing'
    raw_status = str(record.get('status') or '').strip().lower()
    if raw_status in ['cancelled', 'canceled', 'cancelado', 'cancelada']:
        return 'cancelled'
    if raw_status in ['inactive', 'inativo', 'inativa'] or not record.get('active', True):
        return 'inactive'
    expiration_date = ensure_aware_datetime(record.get('expiration_date'))
    if not expiration_date:
        return 'missing'
    if expiration_date.date() < datetime.now(timezone.utc).date():
        return 'expired'
    return 'valid'


def get_course_block_message(status_key: str) -> str:
    if status_key == 'expired':
        return 'O certificado obrigatório para este EPI está vencido.'
    if status_key == 'inactive':
        return 'O certificado obrigatório para este EPI está inativo.'
    if status_key == 'cancelled':
        return 'O certificado obrigatório para este EPI está cancelado.'
    return 'O colaborador não possui certificado válido para retirada deste EPI.'


def pick_latest_employee_course_record(records: List[dict]) -> Optional[dict]:
    if not records:
        return None
    return sorted(
        records,
        key=lambda item: (
            ensure_aware_datetime(item.get('expiration_date')) or datetime.min.replace(tzinfo=timezone.utc),
            ensure_aware_datetime(item.get('updated_at')) or datetime.min.replace(tzinfo=timezone.utc),
            ensure_aware_datetime(item.get('created_at')) or datetime.min.replace(tzinfo=timezone.utc),
        ),
        reverse=True,
    )[0]


def pick_delivery_certificate_record(records: List[dict]) -> tuple[Optional[dict], str]:
    if not records:
        return None, 'missing'

    active_records = [record for record in records if record.get('active', True) and str(record.get('status') or '').strip().lower() not in ['inactive', 'inativo', 'inativa', 'cancelled', 'canceled', 'cancelado', 'cancelada']]
    valid_records = [record for record in active_records if get_course_block_status_from_record(record) == 'valid']
    if valid_records:
        return pick_latest_employee_course_record(valid_records), 'valid'

    if active_records:
        latest_active = pick_latest_employee_course_record(active_records)
        return latest_active, get_course_block_status_from_record(latest_active)

    latest_any = pick_latest_employee_course_record(records)
    return latest_any, get_course_block_status_from_record(latest_any)


async def resolve_delivery_item_epi(db, item: DeliveryItemInput) -> tuple[Optional[str], Optional[dict], Optional[dict]]:
    effective_epi_id = item.epi_id
    variation = None
    if item.epi_variation_id and ObjectId.is_valid(str(item.epi_variation_id)):
        variation = await db.epi_variations.find_one({'_id': ObjectId(item.epi_variation_id)})
        if variation:
            effective_epi_id = str(variation.get('epi_id') or effective_epi_id)
    if not effective_epi_id or not ObjectId.is_valid(str(effective_epi_id)):
        return effective_epi_id, None, variation
    epi = await db.epis.find_one({'_id': ObjectId(effective_epi_id)})
    return effective_epi_id, epi, variation


async def validate_delivery_course_requirements(db, delivery_data: DeliveryCreate, employee: dict, current_user: dict, request: Optional[Request], counter_attendant: Optional[dict] = None) -> None:
    if delivery_data.is_return:
        return

    item_contexts = []
    for item in delivery_data.items:
        effective_epi_id, epi, _ = await resolve_delivery_item_epi(db, item)
        if not effective_epi_id or not epi:
            continue
        item_contexts.append({
            'item': item,
            'epi_id': effective_epi_id,
            'epi_name': epi.get('name') or item.epi_id or effective_epi_id,
            'kit_id': item.kit_id,
            'kit_name': None,
        })

    if not item_contexts:
        return

    kit_ids = sorted({ctx['kit_id'] for ctx in item_contexts if ctx.get('kit_id') and ObjectId.is_valid(str(ctx.get('kit_id')))})
    kits_by_id = {}
    if kit_ids:
        kits = await db.kits.find({'_id': {'$in': [ObjectId(kit_id) for kit_id in kit_ids]}}).to_list(1000)
        kits_by_id = {str(kit['_id']): kit for kit in kits}
        for ctx in item_contexts:
            kit = kits_by_id.get(str(ctx.get('kit_id')))
            if kit:
                ctx['kit_name'] = kit.get('name')

    epi_ids = sorted({ctx['epi_id'] for ctx in item_contexts})
    links = await db.course_epi_links.find({'epi_id': {'$in': epi_ids}}).to_list(1000)
    if not links:
        return

    course_ids = sorted({str(link.get('course_id')) for link in links if link.get('course_id') and ObjectId.is_valid(str(link.get('course_id')))})
    courses_by_id = await get_courses_map(db, course_ids)
    blocking_courses_by_epi = {}
    for link in links:
        epi_id = str(link.get('epi_id'))
        course_id = str(link.get('course_id'))
        course = courses_by_id.get(course_id)
        if not course or not course.get('active', True) or not course.get('block_epi_delivery', False):
            continue
        blocking_courses_by_epi.setdefault(epi_id, []).append(course)

    if not blocking_courses_by_epi:
        return

    required_course_ids = sorted({str(course['_id']) for courses in blocking_courses_by_epi.values() for course in courses})
    records = await db.employee_courses.find({'employee_id': str(employee['_id']), 'course_id': {'$in': required_course_ids}}).to_list(1000)
    records_by_course = {}
    for record in records:
        records_by_course.setdefault(str(record.get('course_id')), []).append(record)

    blocked_items = []
    for ctx in item_contexts:
        for course in blocking_courses_by_epi.get(ctx['epi_id'], []):
            course_id = str(course['_id'])
            certificate, status_key = pick_delivery_certificate_record(records_by_course.get(course_id, []))
            if status_key == 'valid':
                continue
            message = get_course_block_message(status_key)
            blocked_items.append({
                'code': 'COURSE_REQUIRED',
                'status': status_key,
                'course_id': course_id,
                'course_name': course.get('name') or 'Curso obrigatorio',
                'course': course.get('name') or 'Curso obrigatorio',
                'epi_id': ctx['epi_id'],
                'epi_name': ctx['epi_name'],
                'epi': ctx['epi_name'],
                'kit_id': ctx.get('kit_id'),
                'kit_name': ctx.get('kit_name'),
                'expiration_date': serialize_date_for_api(certificate.get('expiration_date')) if certificate else None,
                'message': message,
                'reason': message,
            })

    if not blocked_items:
        return

    has_kit_block = any(item.get('kit_id') for item in blocked_items)
    response_code = 'KIT_COURSE_REQUIRED' if has_kit_block else 'COURSE_REQUIRED'
    first = blocked_items[0]
    detail = {
        'code': response_code,
        'status': first.get('status'),
        'course_id': first.get('course_id'),
        'course_name': first.get('course_name'),
        'course': first.get('course_name'),
        'epi_id': first.get('epi_id'),
        'epi_name': first.get('epi_name'),
        'epi': first.get('epi_name'),
        'kit_id': first.get('kit_id'),
        'kit_name': first.get('kit_name'),
        'expiration_date': first.get('expiration_date'),
        'message': first.get('message'),
        'reason': first.get('reason'),
        'blocked_items': blocked_items if has_kit_block else [],
        'blocked_courses': blocked_items,
    }

    await write_audit_log(
        db,
        'COURSE_DELIVERY_BLOCKED',
        'delivery',
        delivery_data.employee_id,
        current_user,
        reason=first.get('message'),
        request=request,
        metadata={
            'employee_id': delivery_data.employee_id,
            'employee_name': employee.get('full_name'),
            'delivery_type': delivery_data.delivery_type,
            'is_kit': has_kit_block,
            'blocked_items': blocked_items,
            'attendant_id': str(counter_attendant['_id']) if counter_attendant else None,
            'attendant_name': counter_attendant.get('full_name') if counter_attendant else None,
        }
    )
    raise HTTPException(status_code=409, detail=detail)


async def get_courses_map(db, course_ids: List[str]) -> dict:
    object_ids = [ObjectId(course_id) for course_id in course_ids if ObjectId.is_valid(str(course_id))]
    if not object_ids:
        return {}
    courses = await db.courses.find({'_id': {'$in': object_ids}}).to_list(1000)
    return {str(course['_id']): course for course in courses}


async def build_employee_course_responses(db, employee_id: str) -> List[dict]:
    records = await db.employee_courses.find({'employee_id': employee_id}).to_list(1000)
    course_map = await get_courses_map(db, [record.get('course_id') for record in records])
    current_ids = {}
    grouped = {}
    for record in records:
        grouped.setdefault(record.get('course_id'), []).append(record)
    for course_id, items in grouped.items():
        current = pick_current_employee_course(items)
        if current:
            current_ids[str(current['_id'])] = True

    result = []
    for record in records:
        course = course_map.get(record.get('course_id'))
        status_key, status_label, days_until_expiration = calculate_employee_course_status(record, course)
        result.append({
            **doc_to_response(record),
            'course_name': (course or {}).get('name') or 'Curso removido',
            'course_active': (course or {}).get('active', False),
            'alert_days_before_expiration': parse_int_safe((course or {}).get('alert_days_before_expiration'), 0),
            'status': status_key,
            'status_label': status_label,
            'days_until_expiration': days_until_expiration,
            'is_current': bool(current_ids.get(str(record['_id']))),
        })

    result.sort(key=lambda item: (
        (item.get('course_name') or '').lower(),
        0 if item.get('is_current') else 1,
        -(ensure_aware_datetime(item.get('expiration_date')) or datetime.min.replace(tzinfo=timezone.utc)).timestamp(),
    ))
    return result


async def get_employee_course_response_by_id(db, employee_id: str, employee_course_id: str) -> dict:
    record = await db.employee_courses.find_one({'_id': parse_object_id_or_400(employee_course_id, 'employee_course_id')})
    if not record or record.get('employee_id') != employee_id:
        raise HTTPException(status_code=404, detail='Certificado do colaborador nao encontrado')
    records = await build_employee_course_responses(db, employee_id)
    for item in records:
        if item.get('id') == employee_course_id:
            return item
    raise HTTPException(status_code=404, detail='Certificado do colaborador nao encontrado')


async def ensure_employee_exists(db, employee_id: str) -> dict:
    employee = await db.employees.find_one({'_id': parse_object_id_or_400(employee_id, 'employee_id')})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nao encontrado')
    return employee


async def ensure_course_exists(db, course_id: str, require_active: bool = False) -> dict:
    course = await db.courses.find_one({'_id': parse_object_id_or_400(course_id, 'course_id')})
    if not course:
        raise HTTPException(status_code=404, detail='Curso nao encontrado')
    if require_active and not course.get('active', True):
        raise HTTPException(status_code=400, detail='Curso inativo nao pode receber novos certificados')
    return course


async def ensure_employee_course_write_allowed(employee: dict, current_user: dict) -> None:
    if is_employee_active(employee):
        return
    if current_user.get('role') == 'admin':
        return
    raise HTTPException(status_code=400, detail='Colaborador inativo nao pode receber novos lancamentos ou renovacoes')


async def parse_employee_course_request(request: Request) -> tuple[dict, Optional[UploadFile]]:
    content_type = (request.headers.get('content-type') or '').lower()
    fields = ['course_id', 'start_date', 'expiration_date', 'certificate_number', 'training_provider', 'notes']
    if 'multipart/form-data' in content_type or 'application/x-www-form-urlencoded' in content_type:
        form = await request.form()
        payload = {}
        for field in fields:
            value = form.get(field)
            if value is not None and value != '':
                payload[field] = value
        file = form.get('certificate_file')
        if not getattr(file, 'filename', None):
            file = None
        return payload, file
    payload = await request.json()
    return payload or {}, None


async def save_employee_course_file(employee_id: str, course_id: str, file: UploadFile) -> dict:
    original_name = Path(file.filename or 'certificado').name
    extension = Path(original_name).suffix.lower()
    if extension not in EMPLOYEE_COURSE_FILE_EXTENSIONS:
        raise HTTPException(status_code=400, detail='Arquivo do certificado deve ser PDF, JPG, JPEG ou PNG')
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail='Arquivo do certificado esta vazio')
    if len(content) > EMPLOYEE_COURSE_MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail='Arquivo do certificado excede 8MB')
    content_type = (file.content_type or '').lower().strip()
    if content_type and content_type not in EMPLOYEE_COURSE_CONTENT_TYPES:
        raise HTTPException(status_code=400, detail='Tipo de arquivo do certificado nao suportado')
    employee_dir = UPLOAD_DIR / 'employee_courses' / employee_id
    employee_dir.mkdir(exist_ok=True, parents=True)
    file_name = f'employee_course_{employee_id}_{course_id}_{int(datetime.now(timezone.utc).timestamp() * 1000)}{extension}'
    file_path = employee_dir / file_name
    file_path.write_bytes(content)
    return {
        'certificate_file_path': f'/uploads/employee_courses/{employee_id}/{file_name}',
        'certificate_file_name': original_name,
        'certificate_file_content_type': content_type or ('application/pdf' if extension == '.pdf' else 'image/jpeg' if extension in ['.jpg', '.jpeg'] else 'image/png')
    }
def get_epi_group_name(epi: dict) -> str:
    return normalize_text(epi.get('epi_group_name')) or strip_size_from_epi_name(epi.get('name') or epi.get('description'), epi.get('size') or epi.get('tamanho'))

def get_epi_group_key(epi: dict) -> str:
    explicit_key = normalize_text(epi.get('epi_group_key'))
    if explicit_key:
        return explicit_key
    parts = [
        get_epi_group_name(epi),
        infer_epi_category(epi),
        epi.get('model') or epi.get('modelo'),
        epi.get('brand') or epi.get('marca'),
    ]
    return '|'.join(normalize_group_text(part) for part in parts if normalize_text(part))

def slugify_group_key(value: str) -> str:
    normalized = unicodedata.normalize('NFKD', value or '')
    ascii_text = normalized.encode('ascii', 'ignore').decode('ascii')
    slug = re.sub(r'[^A-Za-z0-9]+', '-', ascii_text.upper()).strip('-')
    return slug or 'EPI'

async def generate_unique_epi_group_key(db, group_name: str) -> str:
    base = slugify_group_key(group_name)
    existing_keys = {
        epi.get('epi_group_key')
        for epi in await db.epis.find({"epi_group_key": {"$regex": f"^{base}-", "$options": "i"}}).to_list(1000)
        if epi.get('epi_group_key')
    }
    index = 1
    while True:
        candidate = f"{base}-{index:03d}"
        if candidate not in existing_keys:
            return candidate
        index += 1

async def resolve_epi_group_fields(db, payload: dict, fallback_name: Optional[str] = None) -> tuple[Optional[str], Optional[str]]:
    group_key = normalize_text(payload.get('epi_group_key'))
    group_name = normalize_text(payload.get('epi_group_name'))
    if group_key:
        if not group_name:
            existing = await db.epis.find_one({"epi_group_key": group_key})
            group_name = existing.get('epi_group_name') if existing else fallback_name
        return group_key, group_name
    if group_name:
        return await generate_unique_epi_group_key(db, group_name), group_name
    return None, None

def infer_epi_category(epi: dict) -> str:
    return epi.get('category') or epi.get('type_category') or 'Sem categoria'

def infer_epi_nbr(epi: dict) -> Optional[str]:
    return epi.get('nbr') or epi.get('nbr_number')

def infer_epi_has_size(epi: dict, variations: List[dict]) -> bool:
    if epi.get('possui_variacao_tamanho') is not None:
        return bool(epi.get('possui_variacao_tamanho'))
    if normalize_text(epi.get('size')):
        return True
    return any(normalize_text(v.get('size')) for v in variations)

async def get_supplier_name(db, supplier_id: Optional[str]) -> Optional[str]:
    if not supplier_id:
        return None
    supplier = await db.suppliers.find_one({"_id": ObjectId(supplier_id)})
    return supplier.get('name') if supplier else None

async def create_variation_from_legacy_epi(db, epi: dict) -> Optional[dict]:
    epi_id = str(epi['_id'])
    existing = await db.epi_variations.find_one({"epi_id": epi_id})
    if existing:
        return existing
    has_legacy_data = any([
        normalize_text(epi.get('brand')),
        normalize_text(epi.get('model')),
        normalize_text(epi.get('ca_number')),
        normalize_text(epi.get('supplier_id')),
        normalize_text(epi.get('batch')),
        normalize_text(epi.get('size')),
        epi.get('current_stock') is not None
    ])
    if not has_legacy_data:
        return None

    variation = {
        "epi_id": epi_id,
        "brand": epi.get('brand'),
        "model": epi.get('model'),
        "ca_number": epi.get('ca_number'),
        "supplier_id": epi.get('supplier_id'),
        "supplier_name": await get_supplier_name(db, epi.get('supplier_id')),
        "ca_validity": epi.get('ca_validity'),
        "size": epi.get('size'),
        "color": epi.get('color'),
        "current_stock": epi.get('current_stock', 0) or 0,
        "unit_price": epi.get('unit_price'),
        "purchase_date": epi.get('purchase_date'),
        "batch": epi.get('batch'),
        "status": epi.get('status') or 'ativo',
        "qr_code": epi.get('qr_code'),
        "internal_code": epi.get('internal_code'),
        "invoice_number": epi.get('invoice_number'),
        "validity_date": epi.get('validity_date'),
        "material": epi.get('material'),
        "technical_standard": epi.get('technical_standard'),
        "quantity_purchased": epi.get('quantity_purchased', epi.get('current_stock', 0) or 0),
        "created_at": epi.get('created_at') or datetime.now(timezone.utc),
        "updated_at": datetime.now(timezone.utc)
    }
    result = await db.epi_variations.insert_one(variation)
    variation['_id'] = result.inserted_id
    return variation

async def get_epi_variations(db, epi: dict) -> List[dict]:
    await create_variation_from_legacy_epi(db, epi)
    return await db.epi_variations.find({"epi_id": str(epi['_id'])}).to_list(200)

async def get_epis_by_group_key(db, group_key: Optional[str]) -> List[dict]:
    if not group_key:
        return []
    epis = await db.epis.find({}).to_list(5000)
    return [epi for epi in epis if get_epi_group_key(epi) == group_key]

async def build_epi_group_summary(db, group_key: str, fallback_item: Optional[dict] = None) -> Optional[dict]:
    fallback_item = fallback_item or {}
    group_epis = await get_epis_by_group_key(db, group_key)
    seed = group_epis[0] if group_epis else (fallback_item or {})
    if not seed:
        return None

    available_count = 0
    variations_count = 0
    for epi in group_epis:
        variations = await get_epi_variations(db, epi)
        variations_count += len(variations)
        available_count += len([
            variation for variation in variations
            if (variation.get('current_stock', 0) or 0) > 0 and (variation.get('status') or 'ativo') != 'inativo'
        ])

    return {
        "epi_group_key": group_key,
        "epi_base_key": group_key,
        "name": fallback_item.get('name') or get_epi_group_name(seed),
        "epi_group_name": fallback_item.get('epi_group_name') or get_epi_group_name(seed),
        "description": fallback_item.get('description') or seed.get('description'),
        "category": fallback_item.get('category') or infer_epi_category(seed),
        "type_category": fallback_item.get('type_category') or infer_epi_category(seed),
        "model": fallback_item.get('model') or seed.get('model'),
        "modelo": fallback_item.get('model') or seed.get('model'),
        "brand": fallback_item.get('brand') or seed.get('brand'),
        "marca": fallback_item.get('brand') or seed.get('brand'),
        "variations_count": variations_count,
        "available_variations_count": available_count
    }

def serialize_epi_variation(variation: dict) -> dict:
    return {
        "id": str(variation['_id']),
        "epi_id": variation.get('epi_id'),
        "brand": variation.get('brand'),
        "marca": variation.get('brand'),
        "model": variation.get('model'),
        "modelo": variation.get('model'),
        "ca_number": variation.get('ca_number'),
        "ca": variation.get('ca_number'),
        "supplier_id": variation.get('supplier_id'),
        "supplier_name": variation.get('supplier_name'),
        "fornecedor": variation.get('supplier_name'),
        "ca_validity": variation.get('ca_validity'),
        "validade_ca": variation.get('ca_validity'),
        "size": variation.get('size'),
        "tamanho": variation.get('size'),
        "color": variation.get('color'),
        "cor": variation.get('color'),
        "current_stock": variation.get('current_stock', 0) or 0,
        "quantidade_estoque": variation.get('current_stock', 0) or 0,
        "unit_price": variation.get('unit_price'),
        "valor_unitario": variation.get('unit_price'),
        "purchase_date": variation.get('purchase_date'),
        "data_compra": variation.get('purchase_date'),
        "batch": variation.get('batch'),
        "lote": variation.get('batch'),
        "status": variation.get('status'),
        "qr_code": variation.get('qr_code'),
        "internal_code": variation.get('internal_code'),
        "invoice_number": variation.get('invoice_number'),
        "validity_date": variation.get('validity_date'),
        "material": variation.get('material'),
        "technical_standard": variation.get('technical_standard'),
        "quantity_purchased": variation.get('quantity_purchased', 0) or 0
    }

def choose_primary_variation(variations: List[dict]) -> Optional[dict]:
    if not variations:
        return None
    ranked = sorted(
        variations,
        key=lambda item: (
            0 if (item.get('status') or 'ativo') == 'ativo' else 1,
            -(item.get('current_stock', 0) or 0),
            normalize_text(item.get('size')) is None
        )
    )
    return ranked[0]

async def sync_epi_stock_summary(db, epi_id: str) -> None:
    variations = await db.epi_variations.find({"epi_id": epi_id}).to_list(200)
    total_stock = sum((item.get('current_stock', 0) or 0) for item in variations)
    total_purchased = sum((item.get('quantity_purchased', 0) or 0) for item in variations)
    await db.epis.update_one(
        {"_id": ObjectId(epi_id)},
        {"$set": {
            "current_stock": total_stock,
            "quantity_purchased": total_purchased,
            "updated_at": datetime.now(timezone.utc)
        }}
    )

async def build_epi_response(db, epi: dict) -> dict:
    variations = await get_epi_variations(db, epi)
    for variation in variations:
        if variation.get('supplier_id') and not variation.get('supplier_name'):
            variation['supplier_name'] = await get_supplier_name(db, variation.get('supplier_id'))
    primary = choose_primary_variation(variations)
    serialized_variations = [serialize_epi_variation(item) for item in variations]
    total_stock = sum(item.get('current_stock', 0) or 0 for item in variations)
    total_purchased = sum(item.get('quantity_purchased', 0) or 0 for item in variations)
    validity_candidates = [
        parsed_date
        for item in variations
        for parsed_date in [ensure_aware_datetime(item.get('validity_date') or item.get('ca_validity'))]
        if parsed_date
    ]
    validity_date = min(validity_candidates) if validity_candidates else None
    response = doc_to_response(epi)
    response.update({
        "category": infer_epi_category(epi),
        "type_category": infer_epi_category(epi),
        "description": epi.get('description'),
        "obrigatorio_ca": epi.get('obrigatorio_ca', True),
        "nbr": infer_epi_nbr(epi),
        "nbr_number": infer_epi_nbr(epi),
        "possui_variacao_tamanho": infer_epi_has_size(epi, variations),
        "epi_group_key": epi.get('epi_group_key'),
        "epi_group_name": epi.get('epi_group_name'),
        "variations": serialized_variations,
        "current_stock": total_stock,
        "quantity_purchased": total_purchased,
        "min_stock": epi.get('min_stock', 0) or 0,
        "max_stock": epi.get('max_stock')
    })
    if primary:
        response.update({
            "brand": primary.get('brand'),
            "model": primary.get('model'),
            "color": primary.get('color'),
            "size": primary.get('size'),
            "material": primary.get('material'),
            "ca_number": primary.get('ca_number'),
            "ca_validity": primary.get('ca_validity'),
            "technical_standard": primary.get('technical_standard'),
            "supplier_id": primary.get('supplier_id'),
            "supplier_name": primary.get('supplier_name'),
            "invoice_number": primary.get('invoice_number'),
            "purchase_date": primary.get('purchase_date'),
            "unit_price": primary.get('unit_price'),
            "total_price": (primary.get('unit_price') or 0) * (primary.get('quantity_purchased') or 0),
            "internal_code": primary.get('internal_code'),
            "batch": primary.get('batch'),
            "qr_code": primary.get('qr_code'),
            "validity_date": validity_date or primary.get('validity_date')
        })
    else:
        response["validity_date"] = validity_date
    stock_status, validity_status = calculate_epi_status(response)
    response['stock_status'] = stock_status
    response['validity_status'] = validity_status
    return response

async def build_kit_response(db, kit: dict) -> dict:
    response = doc_to_response(kit)
    items = []
    grouped_items = {}
    for raw_item in kit.get('items', []):
        group_key = raw_item.get('epi_group_key') or raw_item.get('epi_base_key')
        epi_base_id = raw_item.get('epi_base_id') or raw_item.get('epi_id')
        if not group_key and epi_base_id and ObjectId.is_valid(str(epi_base_id)):
            epi = await db.epis.find_one({"_id": ObjectId(epi_base_id)})
            if epi:
                group_key = get_epi_group_key(epi)

        if not group_key:
            continue
        if group_key in grouped_items:
            grouped_items[group_key]["quantity"] = max(grouped_items[group_key]["quantity"], raw_item.get('quantity', 1))
            continue

        summary = await build_epi_group_summary(db, group_key, raw_item)
        if not summary:
            continue
        item = {
            "epi_group_key": group_key,
            "epi_base_key": group_key,
            "epi_base_id": None,
            "epi_id": None,
            "quantity": raw_item.get('quantity', 1),
            **summary
        }
        grouped_items[group_key] = item
        items.append(item)
    response['items'] = items
    return response

def infer_employee_size_field(epi: dict) -> Optional[str]:
    category = (infer_epi_category(epi) or '').lower()
    name = (epi.get('name') or '').lower()
    if any(token in name for token in ['botina', 'bota', 'sapato', 'calÃ§ado', 'calcado']) or 'pÃ©s' in category or 'pes' in category:
        return 'tamanho_calcado'
    if 'luva' in name or 'mÃ£os' in category or 'maos' in category:
        return 'tamanho_luva'
    if any(token in name for token in ['calÃ§a', 'calca', 'bermuda']) or 'pernas' in category:
        return 'tamanho_calca'
    if any(token in name for token in ['camisa', 'camiseta', 'jaqueta', 'jaleco', 'avental']) or 'corpo' in category:
        return 'tamanho_camisa'
    return None

async def resolve_epi_variation(db, epi: dict, employee: Optional[dict] = None, requested_size: Optional[str] = None, requested_variation_id: Optional[str] = None, require_stock: bool = True) -> Optional[dict]:
    variations = await get_epi_variations(db, epi)
    if requested_variation_id:
        for variation in variations:
            if str(variation['_id']) == requested_variation_id:
                if not require_stock or (variation.get('current_stock', 0) or 0) > 0:
                    return variation
                return None
    desired_size = normalize_size_value(requested_size)
    size_field = infer_employee_size_field(epi)
    if not desired_size and employee and size_field:
        desired_size = normalize_size_value(employee.get(size_field))
    active = [item for item in variations if (item.get('status') or 'ativo') != 'inativo']
    candidates = [item for item in active if not require_stock or (item.get('current_stock', 0) or 0) > 0]
    if desired_size:
        exact = [item for item in candidates if normalize_size_value(item.get('size')) == desired_size]
        if exact:
            return choose_primary_variation(exact)
        if require_stock:
            return None
    generic = [item for item in candidates if not normalize_size_value(item.get('size'))]
    if generic:
        return choose_primary_variation(generic)
    return choose_primary_variation(candidates)

async def get_available_variations_for_delivery(db, epi: dict, require_stock: bool = True) -> List[dict]:
    variations = await get_epi_variations(db, epi)
    active = [item for item in variations if (item.get('status') or 'ativo') != 'inativo']
    return [
        item for item in active
        if not require_stock or (item.get('current_stock', 0) or 0) > 0
    ]

async def get_available_group_variations_for_delivery(db, group_key: str, require_stock: bool = True) -> List[dict]:
    variations = []
    for epi in await get_epis_by_group_key(db, group_key):
        variations.extend(await get_available_variations_for_delivery(db, epi, require_stock=require_stock))
    return variations

def choose_suggested_group_variation(variations: List[dict], requested_size: Optional[str] = None) -> Optional[dict]:
    if not variations:
        return None
    desired_size = normalize_size_value(requested_size)
    if desired_size:
        exact = [item for item in variations if normalize_size_value(item.get('size')) == desired_size]
        if exact:
            return choose_primary_variation(exact)
    generic = [item for item in variations if not normalize_size_value(item.get('size'))]
    return choose_primary_variation(generic or variations)

async def build_delivery_suggestions_for_kit(db, employee: dict, kit: dict) -> dict:
    items = []
    processed_groups = set()
    for raw_item in kit.get('items', []):
        group_key = raw_item.get('epi_group_key') or raw_item.get('epi_base_key')
        epi_base_id = raw_item.get('epi_base_id') or raw_item.get('epi_id')
        epi = None
        if not group_key and epi_base_id and ObjectId.is_valid(str(epi_base_id)):
            epi = await db.epis.find_one({"_id": ObjectId(epi_base_id)})
            if epi:
                group_key = get_epi_group_key(epi)
        if not group_key or group_key in processed_groups:
            continue
        processed_groups.add(group_key)

        summary = await build_epi_group_summary(db, group_key, raw_item)
        if not summary:
            continue
        group_epis = await get_epis_by_group_key(db, group_key)
        representative_epi = epi or (group_epis[0] if group_epis else None)
        size_field = infer_employee_size_field(summary)
        requested_size = employee.get(size_field) if size_field else None
        available_variations = await get_available_group_variations_for_delivery(db, group_key, require_stock=True)
        suggested_variation = choose_suggested_group_variation(available_variations, requested_size)
        selected_variation = suggested_variation or (available_variations[0] if len(available_variations) == 1 else None)
        if available_variations:
            items.append({
                "epi_group_key": group_key,
                "epi_base_key": group_key,
                "epi_base_id": None,
                "epi_id": str(selected_variation.get('epi_id')) if selected_variation else (str(representative_epi['_id']) if representative_epi else None),
                "epi_variation_id": str(selected_variation['_id']) if selected_variation and len(available_variations) == 1 else None,
                "name": summary.get('name'),
                "description": summary.get('description'),
                "category": summary.get('category'),
                "type_category": summary.get('type_category'),
                "quantity": raw_item.get('quantity', 1),
                "size_source": size_field,
                "requested_size": requested_size,
                "status": "available",
                "message": "VariaÃ§Ã£o encontrada em estoque",
                "requires_variation_selection": len(available_variations) > 1,
                "suggested_variation_id": str(suggested_variation['_id']) if suggested_variation else None,
                "available_variations": [serialize_epi_variation(variation) for variation in available_variations],
                "variation": serialize_epi_variation(selected_variation) if selected_variation else None
            })
        else:
            items.append({
                "epi_group_key": group_key,
                "epi_base_key": group_key,
                "epi_base_id": None,
                "epi_id": str(representative_epi['_id']) if representative_epi else None,
                "name": summary.get('name'),
                "description": summary.get('description'),
                "category": summary.get('category'),
                "type_category": summary.get('type_category'),
                "quantity": raw_item.get('quantity', 1),
                "size_source": size_field,
                "requested_size": requested_size,
                "status": "unavailable",
                "message": "Sem estoque disponÃ­vel para tamanho solicitado."
            })
    return {
        "employee_id": str(employee['_id']),
        "employee_name": employee.get('full_name', ''),
        "department": employee.get('department'),
        "kit_id": str(kit['_id']),
        "kit_name": kit.get('name'),
        "items": items
    }

def calculate_epi_status(epi):
    """Calcula status de estoque e validade do EPI"""
    stock_status = 'ok'
    validity_status = 'ok'
    
    # Status de estoque
    if epi.get('current_stock', 0) <= 0:
        stock_status = 'out'
    elif epi.get('current_stock', 0) <= epi.get('min_stock', 0):
        stock_status = 'low'
    
    # Status de validade
    validity_date = ensure_aware_datetime(epi.get('validity_date') or epi.get('ca_validity'))
    if validity_date:
        now = datetime.now(timezone.utc)
        days_until_expiry = (validity_date - now).days
        
        if days_until_expiry < 0:
            validity_status = 'expired'
        elif days_until_expiry <= 30:
            validity_status = 'expiring'
    
    return stock_status, validity_status

@api_router.get('/epis', response_model=List[EPIResponse])
async def get_epis(current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    epis = await db.epis.find({}).to_list(1000)
    result = []
    for e in epis:
        result.append(EPIResponse(**(await build_epi_response(db, e))))
    return result

@api_router.post('/epis', response_model=EPIResponse, status_code=status.HTTP_201_CREATED)
async def create_epi(epi_data: EPICreate, current_user: dict = Depends(require_role('admin', 'almoxarifado'))):
    db = await get_db()
    payload = epi_data.model_dump()
    variations_payload = payload.pop('variations', []) or []
    category = payload.get('category') or payload.get('type_category')
    nbr = payload.get('nbr') or payload.get('nbr_number')
    epi_group_key, epi_group_name = await resolve_epi_group_fields(db, payload, payload.get('name'))

    if payload.get('obrigatorio_ca', True):
        if not variations_payload and not payload.get('ca_number'):
            raise HTTPException(status_code=400, detail='Cadastre ao menos uma variaÃ§Ã£o com CA para este EPI.')
        if variations_payload and not any(normalize_text(item.get('ca') or item.get('ca_number')) for item in variations_payload):
            raise HTTPException(status_code=400, detail='Cadastre ao menos uma variaÃ§Ã£o com CA para este EPI.')

    now = datetime.now(timezone.utc)
    new_epi = {
        'name': payload['name'],
        'category': category,
        'type_category': category,
        'description': payload.get('description'),
        'obrigatorio_ca': payload.get('obrigatorio_ca', True),
        'nbr': nbr,
        'nbr_number': nbr,
        'possui_variacao_tamanho': payload.get('possui_variacao_tamanho', False),
        'epi_group_key': epi_group_key,
        'epi_group_name': epi_group_name,
        'min_stock': payload.get('min_stock', 0),
        'max_stock': payload.get('max_stock'),
        'replacement_period': str(payload['replacement_period'].value) if payload.get('replacement_period') else None,
        'replacement_days': payload.get('replacement_days'),
        'created_by': current_user['id'],
        'created_at': now,
        'updated_at': now
    }
    result = await db.epis.insert_one(new_epi)
    new_epi['_id'] = result.inserted_id

    if not variations_payload and any([
        normalize_text(payload.get('brand')),
        normalize_text(payload.get('model')),
        normalize_text(payload.get('ca_number')),
        payload.get('current_stock', 0) > 0
    ]):
        variations_payload = [payload]

    for raw_variation in variations_payload:
        supplier_id = raw_variation.get('supplier_id')
        quantity = get_variation_stock_quantity(raw_variation)
        variation = {
            'epi_id': str(result.inserted_id),
            'brand': raw_variation.get('marca') or raw_variation.get('brand'),
            'model': raw_variation.get('modelo') or raw_variation.get('model'),
            'ca_number': raw_variation.get('ca') or raw_variation.get('ca_number'),
            'supplier_id': supplier_id,
            'supplier_name': raw_variation.get('fornecedor') or await get_supplier_name(db, supplier_id),
            'ca_validity': raw_variation.get('validade_ca') or raw_variation.get('ca_validity'),
            'size': raw_variation.get('tamanho') or raw_variation.get('size'),
            'color': raw_variation.get('cor') or raw_variation.get('color'),
            'current_stock': quantity,
            'unit_price': raw_variation.get('valor_unitario', raw_variation.get('unit_price')),
            'purchase_date': raw_variation.get('data_compra') or raw_variation.get('purchase_date'),
            'batch': raw_variation.get('lote') or raw_variation.get('batch'),
            'status': raw_variation.get('status') or 'ativo',
            'qr_code': raw_variation.get('qr_code'),
            'internal_code': raw_variation.get('internal_code'),
            'invoice_number': raw_variation.get('invoice_number'),
            'validity_date': raw_variation.get('validity_date'),
            'material': raw_variation.get('material'),
            'technical_standard': raw_variation.get('technical_standard'),
            'quantity_purchased': parse_int_safe(raw_variation.get('quantity_purchased'), quantity),
            'created_at': now,
            'updated_at': now
        }
        await db.epi_variations.insert_one(variation)

    await sync_epi_stock_summary(db, str(result.inserted_id))
    refreshed = await db.epis.find_one({'_id': result.inserted_id})
    return EPIResponse(**(await build_epi_response(db, refreshed)))

@api_router.get('/epis/{epi_id}', response_model=EPIResponse)
async def get_epi(epi_id: str, current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    epi = await db.epis.find_one({'_id': ObjectId(epi_id)})
    if not epi:
        raise HTTPException(status_code=404, detail='EPI n?o encontrado')
    return EPIResponse(**(await build_epi_response(db, epi)))

@api_router.patch('/epis/{epi_id}', response_model=EPIResponse)
async def update_epi(epi_id: str, epi_data: EPIUpdate, current_user: dict = Depends(require_role('admin', 'almoxarifado'))):
    db = await get_db()
    existing = await db.epis.find_one({'_id': ObjectId(epi_id)})
    if not existing:
        raise HTTPException(status_code=404, detail='EPI n?o encontrado')

    update_data = {k: v for k, v in epi_data.model_dump(exclude_unset=True).items() if k != 'variations'}
    update_data['updated_at'] = datetime.now(timezone.utc)

    if 'category' in update_data or 'type_category' in update_data:
        category = update_data.get('category') or update_data.get('type_category')
        update_data['category'] = category
        update_data['type_category'] = category
    if 'nbr' in update_data or 'nbr_number' in update_data:
        nbr = update_data.get('nbr') or update_data.get('nbr_number')
        update_data['nbr'] = nbr
        update_data['nbr_number'] = nbr
    if 'epi_group_key' in update_data or 'epi_group_name' in update_data:
        group_key, group_name = await resolve_epi_group_fields(
            db,
            {**existing, **update_data},
            update_data.get('name') or existing.get('name')
        )
        update_data['epi_group_key'] = group_key
        update_data['epi_group_name'] = group_name
    if update_data.get('replacement_period'):
        update_data['replacement_period'] = str(update_data['replacement_period'].value) if hasattr(update_data['replacement_period'], 'value') else str(update_data['replacement_period'])

    await db.epis.update_one({'_id': ObjectId(epi_id)}, {'$set': update_data})

    if epi_data.variations is not None:
        await db.epi_variations.delete_many({'epi_id': epi_id})
        now = datetime.now(timezone.utc)
        for raw_variation in epi_data.variations:
            supplier_id = raw_variation.get('supplier_id')
            quantity = get_variation_stock_quantity(raw_variation)
            variation = {
                'epi_id': epi_id,
                'brand': raw_variation.get('marca') or raw_variation.get('brand'),
                'model': raw_variation.get('modelo') or raw_variation.get('model'),
                'ca_number': raw_variation.get('ca') or raw_variation.get('ca_number'),
                'supplier_id': supplier_id,
                'supplier_name': raw_variation.get('fornecedor') or await get_supplier_name(db, supplier_id),
                'ca_validity': raw_variation.get('validade_ca') or raw_variation.get('ca_validity'),
                'size': raw_variation.get('tamanho') or raw_variation.get('size'),
                'color': raw_variation.get('cor') or raw_variation.get('color'),
                'current_stock': quantity,
                'unit_price': raw_variation.get('valor_unitario', raw_variation.get('unit_price')),
                'purchase_date': raw_variation.get('data_compra') or raw_variation.get('purchase_date'),
                'batch': raw_variation.get('lote') or raw_variation.get('batch'),
                'status': raw_variation.get('status') or 'ativo',
                'qr_code': raw_variation.get('qr_code'),
                'internal_code': raw_variation.get('internal_code'),
                'invoice_number': raw_variation.get('invoice_number'),
                'validity_date': raw_variation.get('validity_date'),
                'material': raw_variation.get('material'),
                'technical_standard': raw_variation.get('technical_standard'),
                'quantity_purchased': parse_int_safe(raw_variation.get('quantity_purchased'), quantity),
                'created_at': now,
                'updated_at': now
            }
            await db.epi_variations.insert_one(variation)

    await sync_epi_stock_summary(db, epi_id)
    refreshed = await db.epis.find_one({'_id': ObjectId(epi_id)})
    return EPIResponse(**(await build_epi_response(db, refreshed)))

@api_router.delete('/epis/{epi_id}')
async def delete_epi(epi_id: str, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    await db.epi_variations.delete_many({"epi_id": epi_id})
    result = await db.epis.delete_one({"_id": ObjectId(epi_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='EPI nÃ£o encontrado')
    return {'message': 'EPI excluÃ­do'}

@api_router.delete('/epis/groups/{group_key}')
async def delete_epi_group(group_key: str, current_user: dict = Depends(require_role('admin', 'almoxarifado'))):
    db = await get_db()
    epis = await db.epis.find({}).to_list(5000)
    group_epis = [epi for epi in epis if get_epi_group_key(epi) == group_key]
    if not group_epis:
        raise HTTPException(status_code=404, detail='Grupo de EPI nÃ£o encontrado')

    epi_ids = [str(epi['_id']) for epi in group_epis]
    object_ids = [epi['_id'] for epi in group_epis]
    await db.epi_variations.delete_many({"epi_id": {"$in": epi_ids}})
    epi_result = await db.epis.delete_many({"_id": {"$in": object_ids}})

    kits_updated = 0
    async for kit in db.kits.find({"items": {"$elemMatch": {"epi_group_key": group_key}}}):
        original_items = kit.get('items', [])
        kept_items = [
            item for item in original_items
            if (item.get('epi_group_key') or item.get('epi_base_key')) != group_key
        ]
        if len(kept_items) != len(original_items):
            await db.kits.update_one(
                {"_id": kit['_id']},
                {"$set": {"items": kept_items, "updated_at": datetime.now(timezone.utc)}}
            )
            kits_updated += 1

    return {
        'message': 'Grupo de EPI excluÃ­do',
        'epis_removed': epi_result.deleted_count,
        'variations_removed_for_epi_ids': epi_ids,
        'kits_updated': kits_updated
    }

# ===================== KITS =====================

async def normalize_kit_items_for_storage(db, kit_items: List[KitItemInput]) -> List[dict]:
    normalized_items = []
    seen_groups = set()
    for item in kit_items:
        group_key = item.epi_group_key
        representative_id = item.epi_base_id or item.epi_id
        if not group_key and representative_id and ObjectId.is_valid(str(representative_id)):
            epi = await db.epis.find_one({"_id": ObjectId(representative_id)})
            if epi:
                group_key = get_epi_group_key(epi)
        if not group_key or group_key in seen_groups:
            continue
        seen_groups.add(group_key)
        normalized_items.append({
            'epi_group_key': group_key,
            'epi_base_key': group_key,
            'name': item.name,
            'description': item.description,
            'category': item.category or item.type_category,
            'type_category': item.type_category or item.category,
            'model': item.model,
            'brand': item.brand,
            'quantity': item.quantity
        })
    return normalized_items

@api_router.get('/kits', response_model=List[KitResponse])
async def get_kits(current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    kits = await db.kits.find({}).to_list(1000)
    return [KitResponse(**(await build_kit_response(db, k))) for k in kits]

@api_router.post('/kits', response_model=KitResponse, status_code=status.HTTP_201_CREATED)
async def create_kit(kit_data: KitCreate, current_user: dict = Depends(require_role('admin', 'seguranca_trabalho'))):
    db = await get_db()
    items = await normalize_kit_items_for_storage(db, kit_data.items)
    new_kit = {
        'name': kit_data.name,
        'description': kit_data.description,
        'sector': kit_data.sector,
        'is_mandatory': kit_data.is_mandatory,
        'items': items,
        'created_at': datetime.now(timezone.utc),
        'updated_at': datetime.now(timezone.utc)
    }
    result = await db.kits.insert_one(new_kit)
    new_kit['_id'] = result.inserted_id
    return KitResponse(**(await build_kit_response(db, new_kit)))

@api_router.get('/kits/{kit_id}', response_model=KitResponse)
async def get_kit(kit_id: str, current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho', 'almoxarifado'))):
    db = await get_db()
    kit = await db.kits.find_one({'_id': ObjectId(kit_id)})
    if not kit:
        raise HTTPException(status_code=404, detail='Kit n?o encontrado')
    return KitResponse(**(await build_kit_response(db, kit)))

@api_router.patch('/kits/{kit_id}', response_model=KitResponse)
async def update_kit(kit_id: str, kit_data: KitUpdate, current_user: dict = Depends(require_role('admin', 'seguranca_trabalho'))):
    db = await get_db()
    update_data = {}
    if kit_data.name is not None:
        update_data['name'] = kit_data.name
    if kit_data.description is not None:
        update_data['description'] = kit_data.description
    if kit_data.sector is not None:
        update_data['sector'] = kit_data.sector
    if kit_data.is_mandatory is not None:
        update_data['is_mandatory'] = kit_data.is_mandatory
    if kit_data.items is not None:
        update_data['items'] = await normalize_kit_items_for_storage(db, kit_data.items)
    update_data['updated_at'] = datetime.now(timezone.utc)
    result = await db.kits.find_one_and_update({'_id': ObjectId(kit_id)}, {'$set': update_data}, return_document=True)
    if not result:
        raise HTTPException(status_code=404, detail='Kit n?o encontrado')
    return KitResponse(**(await build_kit_response(db, result)))

@api_router.delete('/kits/{kit_id}')
async def delete_kit(kit_id: str, current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    result = await db.kits.delete_one({"_id": ObjectId(kit_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Kit nÃ£o encontrado')
    return {'message': 'Kit excluÃ­do'}

# ===================== COUNTER ATTENDANTS =====================

@api_router.get('/counter-attendants/active', response_model=List[CounterAttendantResponse])
async def get_active_counter_attendants(current_user: dict = Depends(require_role('admin', 'almoxarifado'))):
    db = await get_db()
    attendants = await db.counter_attendants.find({'is_active': True}).sort('full_name', 1).to_list(500)
    return [CounterAttendantResponse(**doc_to_response(attendant)) for attendant in attendants]

@api_router.get('/counter-attendants', response_model=List[CounterAttendantResponse])
async def get_counter_attendants(current_user: dict = Depends(require_role('admin'))):
    db = await get_db()
    attendants = await db.counter_attendants.find({}).sort('full_name', 1).to_list(500)
    return [CounterAttendantResponse(**doc_to_response(attendant)) for attendant in attendants]

@api_router.post('/counter-attendants', response_model=CounterAttendantResponse, status_code=status.HTTP_201_CREATED)
async def create_counter_attendant(attendant_data: CounterAttendantCreate, current_user: dict = Depends(require_role('admin'))):
    full_name = attendant_data.full_name.strip()
    if len(full_name) < 3:
        raise HTTPException(status_code=400, detail='Informe o nome completo do atendente.')

    db = await get_db()
    existing = await db.counter_attendants.find_one({
        'full_name': {'$regex': f'^{re.escape(full_name)}$', '$options': 'i'}
    })
    if existing:
        raise HTTPException(status_code=400, detail='Atendente ja cadastrado.')

    now = datetime.now(timezone.utc)
    attendant = {
        'full_name': full_name,
        'is_active': True,
        'created_at': now,
        'created_by': current_user.get('id'),
        'updated_at': now,
        'updated_by': current_user.get('id')
    }
    result = await db.counter_attendants.insert_one(attendant)
    attendant['_id'] = result.inserted_id
    return CounterAttendantResponse(**doc_to_response(attendant))

@api_router.patch('/counter-attendants/{attendant_id}', response_model=CounterAttendantResponse)
async def update_counter_attendant(attendant_id: str, attendant_data: CounterAttendantUpdate, current_user: dict = Depends(require_role('admin'))):
    if not ObjectId.is_valid(attendant_id):
        raise HTTPException(status_code=400, detail='ID invalido')

    db = await get_db()
    update_fields = {}
    if attendant_data.full_name is not None:
        full_name = attendant_data.full_name.strip()
        if len(full_name) < 3:
            raise HTTPException(status_code=400, detail='Informe o nome completo do atendente.')
        existing = await db.counter_attendants.find_one({
            '_id': {'$ne': ObjectId(attendant_id)},
            'full_name': {'$regex': f'^{re.escape(full_name)}$', '$options': 'i'}
        })
        if existing:
            raise HTTPException(status_code=400, detail='Atendente ja cadastrado.')
        update_fields['full_name'] = full_name

    if attendant_data.is_active is not None:
        update_fields['is_active'] = attendant_data.is_active

    if not update_fields:
        attendant = await db.counter_attendants.find_one({'_id': ObjectId(attendant_id)})
        if not attendant:
            raise HTTPException(status_code=404, detail='Atendente nao encontrado')
        return CounterAttendantResponse(**doc_to_response(attendant))

    update_fields['updated_at'] = datetime.now(timezone.utc)
    update_fields['updated_by'] = current_user.get('id')
    result = await db.counter_attendants.update_one({'_id': ObjectId(attendant_id)}, {'$set': update_fields})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail='Atendente nao encontrado')

    attendant = await db.counter_attendants.find_one({'_id': ObjectId(attendant_id)})
    return CounterAttendantResponse(**doc_to_response(attendant))

@api_router.delete('/counter-attendants/{attendant_id}')
async def delete_counter_attendant(attendant_id: str, current_user: dict = Depends(require_role('admin'))):
    if not ObjectId.is_valid(attendant_id):
        raise HTTPException(status_code=400, detail='ID invalido')

    db = await get_db()
    linked_deliveries = await db.deliveries.count_documents({'attendant_id': attendant_id})
    if linked_deliveries > 0:
        raise HTTPException(status_code=400, detail='Atendente possui entregas vinculadas. Inative o atendente para preservar o historico.')

    result = await db.counter_attendants.delete_one({'_id': ObjectId(attendant_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail='Atendente nao encontrado')
    return {'message': 'Atendente excluido'}

# ===================== DELIVERIES =====================

@api_router.post('/deliveries', response_model=DeliveryResponse, status_code=status.HTTP_201_CREATED)
async def create_delivery(delivery_data: DeliveryCreate, request: Request, current_user: dict = Depends(get_current_user)):
    if not can_deliver_epi(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permiss?o para realizar entregas')

    db = await get_db()
    counter_attendant = None
    if is_counter_user(current_user):
        counter_attendant = await get_counter_attendant_for_delivery(db, delivery_data.attendant_id)

    if not ObjectId.is_valid(delivery_data.employee_id):
        raise HTTPException(status_code=400, detail='ID de colaborador invalido')

    employee = await db.employees.find_one({'_id': ObjectId(delivery_data.employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador n?o encontrado')
    if not delivery_data.is_return and not is_employee_active(employee):
        await write_audit_log(
            db,
            'COLLABORATOR_INACTIVE_DELIVERY_BLOCKED',
            'collaborator',
            delivery_data.employee_id,
            current_user,
            reason='Tentativa de entrega para colaborador inativo',
            previous_status=normalize_employee_status(employee),
            new_status=normalize_employee_status(employee),
            request=request,
            metadata={'delivery_type': delivery_data.delivery_type, 'items_count': len(delivery_data.items)}
        )
        raise HTTPException(status_code=409, detail='Este colaborador esta inativo e nao pode receber uma nova entrega de EPI.')
    if not employee.get('photo_path'):
        raise HTTPException(status_code=400, detail='Colaborador n?o possui foto cadastrada. Procure o RH para cadastrar.')

    await validate_delivery_course_requirements(db, delivery_data, employee, current_user, request, counter_attendant)

    items_list = []
    for item in delivery_data.items:
        if item.quantity < 1:
            raise HTTPException(status_code=400, detail='Quantidade deve ser maior ou igual a 1.')

        item_dict = item.model_dump()
        effective_epi_id = item.epi_id
        preselected_variation = None
        if item.epi_variation_id and ObjectId.is_valid(str(item.epi_variation_id)):
            preselected_variation = await db.epi_variations.find_one({'_id': ObjectId(item.epi_variation_id)})
            if preselected_variation:
                effective_epi_id = str(preselected_variation.get('epi_id') or effective_epi_id)

        if effective_epi_id:
            epi = await db.epis.find_one({'_id': ObjectId(effective_epi_id)})
            if not epi:
                raise HTTPException(status_code=404, detail='EPI n?o encontrado para entrega.')

            available_variations = await get_available_variations_for_delivery(
                db,
                epi,
                require_stock=not delivery_data.is_return
            )
            if not item.epi_variation_id and len(available_variations) > 1:
                raise HTTPException(status_code=400, detail='Selecione a variaÃ§Ã£o do EPI antes de concluir a entrega.')

            variation = await resolve_epi_variation(
                db,
                epi,
                employee=employee,
                requested_size=item.size,
                requested_variation_id=item.epi_variation_id,
                require_stock=not delivery_data.is_return
            )
            if not variation:
                raise HTTPException(status_code=400, detail='EPI sem saldo disponÃ­vel.')

            available_quantity = int(variation.get('current_stock', 0) or 0)
            requested_quantity = int(item.quantity or 0)
            if not delivery_data.is_return and available_quantity < requested_quantity:
                raise HTTPException(status_code=400, detail='Estoque insuficiente para realizar esta entrega.')

            stock_change = item.quantity if delivery_data.is_return else -item.quantity
            await db.epi_variations.update_one({'_id': variation['_id']}, {'$inc': {'current_stock': stock_change}})
            await sync_epi_stock_summary(db, effective_epi_id)

            movement = {
                'movement_type': 'return' if delivery_data.is_return else 'delivery',
                'epi_id': effective_epi_id,
                'epi_group_key': get_epi_group_key(epi),
                'epi_variation_id': str(variation['_id']),
                'quantity': item.quantity if delivery_data.is_return else -item.quantity,
                'employee_id': delivery_data.employee_id,
                'created_by': current_user['id'],
                'created_at': datetime.now(timezone.utc)
            }
            await db.stock_movements.insert_one(movement)

            item_dict.update({
                'epi_name': epi.get('name'),
                'epi_id': effective_epi_id,
                'epi_group_key': get_epi_group_key(epi),
                'epi_variation_id': str(variation['_id']),
                'brand': variation.get('brand'),
                'model': variation.get('model'),
                'size': variation.get('size') or item.size,
                'batch': variation.get('batch') or item.batch,
                'qr_code': variation.get('qr_code') or item.qr_code,
                'ca_number': variation.get('ca_number'),
                'ca_validity': variation.get('ca_validity'),
                'validity_date': variation.get('validity_date'),
                'supplier_id': variation.get('supplier_id'),
                'supplier_name': variation.get('supplier_name'),
                'quantity': item.quantity,
                'kit_id': item.kit_id
            })
            if item.kit_id:
                kit_doc = await db.kits.find_one({'_id': ObjectId(item.kit_id)})
                item_dict['kit_name'] = kit_doc.get('name') if kit_doc else None
        else:
            raise HTTPException(status_code=400, detail='Selecione um EPI com saldo disponÃ­vel antes de concluir a entrega.')

        items_list.append(item_dict)

    new_delivery = {
        'employee_id': delivery_data.employee_id,
        'employee_name': employee['full_name'],
        'delivery_type': delivery_data.delivery_type,
        'is_return': delivery_data.is_return,
        'facial_match_score': delivery_data.facial_match_score,
        'facial_photo_path': delivery_data.facial_photo_path,
        'facial_validation_status': delivery_data.facial_validation_status,
        'facial_validation_message': delivery_data.facial_validation_message,
        'facial_liveness_status': delivery_data.facial_liveness_status,
        'facial_second_capture_used': delivery_data.facial_second_capture_used,
        'notes': delivery_data.notes,
        'items': items_list,
        'delivered_by': current_user['id'],
        'delivered_by_name': current_user['username'],
        'session_username': current_user['username'],
        'terminal_user': current_user['username'],
        'attendant_id': str(counter_attendant['_id']) if counter_attendant else None,
        'attendant_name': counter_attendant.get('full_name') if counter_attendant else None,
        'created_at': datetime.now(timezone.utc)
    }
    result = await db.deliveries.insert_one(new_delivery)
    new_delivery['_id'] = result.inserted_id
    return DeliveryResponse(**doc_to_response(new_delivery))

@api_router.post('/deliveries/save-photo')
async def save_delivery_photo(
    employee_id: str = Form(...),
    photo_data: str = Form(...),
    current_user: dict = Depends(get_current_user)
):
    """Salva a foto de confirmaÃ§Ã£o da entrega"""
    if not can_deliver_epi(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o')
    
    try:
        # Decodificar base64
        if ',' in photo_data:
            photo_data = photo_data.split(',')[1]
        
        photo_bytes = base64.b64decode(photo_data)
        
        file_name = f'delivery_{employee_id}_{datetime.now(timezone.utc).timestamp()}.jpg'
        file_path = UPLOAD_DIR / 'deliveries' / file_name
        file_path.parent.mkdir(exist_ok=True, parents=True)
        
        with open(file_path, 'wb') as f:
            f.write(photo_bytes)
        
        return {'photo_path': f'/uploads/deliveries/{file_name}'}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f'Erro ao salvar foto: {str(e)}')

@api_router.get('/deliveries', response_model=List[DeliveryResponse])
async def get_deliveries(
    current_user: dict = Depends(get_current_user), 
    employee_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None
):
    if not can_view_delivery_history(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissÃ£o para consultar entregas')

    db = await get_db()
    query = {}
    if employee_id:
        query['employee_id'] = employee_id
    
    if start_date:
        query['created_at'] = query.get('created_at', {})
        query['created_at']['$gte'] = datetime.fromisoformat(start_date.replace('Z', '+00:00'))
    
    if end_date:
        query['created_at'] = query.get('created_at', {})
        query['created_at']['$lte'] = datetime.fromisoformat(end_date.replace('Z', '+00:00'))
    
    deliveries = await db.deliveries.find(query).sort("created_at", -1).to_list(1000)
    return [DeliveryResponse(**doc_to_response(d)) for d in deliveries]

# ===================== STOCK =====================

ALERT_REPLACEMENT_WARNING_DAYS = 7
ALERT_CA_WARNING_DAYS = 30
STOCK_ALERT_ROLES = {'admin', 'gestor', 'almoxarifado'}

ALERT_DEFINITIONS = {
    'pending_epi': {
        'title': 'EPI obrigatório pendente',
        'severity': 'ATENÇÃO',
        'action': 'Verificar necessidade de entrega.',
    },
    'replacement_expiring': {
        'title': 'Troca de EPI a vencer',
        'severity': 'ATENÇÃO',
        'action': 'Programar substituição.',
    },
    'replacement_due': {
        'title': 'Troca de EPI vencida',
        'severity': 'URGENTE',
        'action': 'Providenciar substituição.',
    },
    'stock_low': {
        'title': 'Estoque baixo',
        'severity': 'ATENÇÃO',
        'action': 'Providenciar reposição.',
    },
    'stock_zero': {
        'title': 'Estoque zerado',
        'severity': 'CRÍTICO',
        'action': 'Reposição urgente.',
    },
    'ca_expiring': {
        'title': 'CA a vencer',
        'severity': 'ATENÇÃO',
        'action': 'Providenciar renovação/verificação.',
    },
    'ca_expired': {
        'title': 'CA vencido',
        'severity': 'CRÍTICO',
        'action': 'Verificar regularização antes de novas utilizações.',
    },
    'course_expiring': {
        'title': 'Curso a vencer',
        'severity': 'ATENÇÃO',
        'action': 'Programar treinamento/renovação.',
    },
    'course_expired': {
        'title': 'Curso vencido',
        'severity': 'URGENTE',
        'action': 'Regularizar treinamento.',
    },
}

ALERT_ORDER = [
    'pending_epi',
    'replacement_expiring',
    'replacement_due',
    'stock_low',
    'stock_zero',
    'ca_expiring',
    'ca_expired',
    'course_expiring',
    'course_expired',
]


def user_can_view_stock_alerts(current_user: dict) -> bool:
    return (current_user or {}).get('role') in STOCK_ALERT_ROLES


def alert_definition(alert_type: str) -> dict:
    return ALERT_DEFINITIONS.get(alert_type, {
        'title': 'Alerta operacional',
        'severity': 'INFO',
        'action': 'Verificar informação.',
    })


def build_alert_id(alert_type: str, *parts) -> str:
    normalized = [str(part or '').strip() for part in parts if str(part or '').strip()]
    return ':'.join([alert_type] + normalized)


def normalize_alert_item(alert_type: str, **payload) -> dict:
    definition = alert_definition(alert_type)
    alert_id = payload.pop('alert_id', None) or build_alert_id(
        alert_type,
        payload.get('employee_id'),
        payload.get('epi_id') or payload.get('epi_group_key'),
        payload.get('variation_id'),
        payload.get('course_id'),
        payload.get('kit_id'),
    )
    return {
        'alert_id': alert_id,
        'type': alert_type,
        'alert_type': alert_type,
        'title': definition['title'],
        'severity': definition['severity'],
        'action': definition['action'],
        **payload,
    }


def empty_alert_summary() -> dict:
    summary = {alert_type: 0 for alert_type in ALERT_ORDER}
    summary.update({
        'pending_epis': 0,
        'replacement_expiring': 0,
        'replacement_due': 0,
        'stock_low': 0,
        'stock_zero': 0,
        'ca_expiring': 0,
        'ca_expired': 0,
        'course_expiring': 0,
        'course_expired': 0,
        'total': 0,
        'total_alerts': 0,
    })
    return summary


def summarize_alerts(alerts_by_type: dict) -> dict:
    summary = empty_alert_summary()
    seen_ids = set()
    for alert_type in ALERT_ORDER:
        items = alerts_by_type.get(alert_type, [])
        summary[alert_type] = len(items)
        for item in items:
            seen_ids.add(item.get('alert_id') or build_alert_id(alert_type, item.get('id')))
    summary['pending_epis'] = summary['pending_epi']
    summary['total'] = len(seen_ids)
    summary['total_alerts'] = len(seen_ids)
    return summary


async def build_epi_maps(db) -> tuple[dict, dict]:
    epis = await db.epis.find({}).to_list(10000)
    epi_by_id = {str(epi['_id']): epi for epi in epis}
    epi_group_by_id = {epi_id: get_epi_group_key(epi) for epi_id, epi in epi_by_id.items()}
    return epi_by_id, epi_group_by_id


async def build_pending_epi_alerts(db) -> List[dict]:
    employees = await db.employees.find(
        {'status': 'active'},
        {'full_name': 1, 'department': 1}
    ).to_list(5000)
    kits = await db.kits.find(
        {'is_mandatory': {'$ne': False}},
        {'name': 1, 'sector': 1, 'items': 1, 'is_mandatory': 1}
    ).to_list(500)
    kits_by_sector = {str(kit.get('sector') or '').lower(): kit for kit in kits if kit.get('sector')}
    if not employees or not kits_by_sector:
        return []

    _, epi_group_by_id = await build_epi_maps(db)
    deliveries = await db.deliveries.find(
        {'is_return': False},
        {'employee_id': 1, 'items.epi_group_key': 1, 'items.epi_id': 1}
    ).to_list(20000)

    delivered_groups_by_employee = {}
    for delivery in deliveries:
        employee_id = str(delivery.get('employee_id') or '')
        groups = delivered_groups_by_employee.setdefault(employee_id, set())
        for item in delivery.get('items', []):
            group_key = item.get('epi_group_key')
            if not group_key and item.get('epi_id'):
                group_key = epi_group_by_id.get(str(item.get('epi_id')))
            if group_key:
                groups.add(group_key)

    summary_cache = {}
    alerts = []
    for employee in employees:
        emp_id = str(employee['_id'])
        department = str(employee.get('department') or '').lower()
        kit = kits_by_sector.get(department)
        if not kit:
            continue

        delivered_group_keys = delivered_groups_by_employee.get(emp_id, set())
        missing_epis = []
        for kit_item in kit.get('items', []):
            group_key = kit_item.get('epi_group_key') or kit_item.get('epi_base_key')
            epi_id = kit_item.get('epi_base_id') or kit_item.get('epi_id')
            if not group_key and epi_id:
                group_key = epi_group_by_id.get(str(epi_id))
            if group_key and group_key not in delivered_group_keys:
                if group_key not in summary_cache:
                    summary_cache[group_key] = await build_epi_group_summary(db, group_key, kit_item)
                summary = summary_cache.get(group_key) or {}
                missing_epis.append({
                    'epi_group_key': group_key,
                    'name': summary.get('name') or kit_item.get('name') or 'EPI',
                    'ca_number': '',
                    'quantity_required': kit_item.get('quantity', 1),
                })

        if missing_epis:
            alerts.append(normalize_alert_item(
                'pending_epi',
                alert_id=build_alert_id('pending_epi', emp_id, str(kit['_id'])),
                employee_id=emp_id,
                employee_name=employee.get('full_name', ''),
                department=employee.get('department', ''),
                kit_name=kit.get('name', ''),
                kit_id=str(kit['_id']),
                missing_epis=missing_epis,
                entity='Colaborador',
                status='Pendente',
                relevant_date=None,
                message=f"{employee.get('full_name')} ({employee.get('department', '')}) não possui {len(missing_epis)} EPI(s) obrigatório(s) do kit",
            ))
    return alerts


def get_replacement_days(epi):
    period = epi.get('replacement_period')
    if period == 'weekly':
        return 7
    if period == 'biweekly':
        return 14
    if period == 'monthly':
        return 30
    if period == 'custom':
        return parse_int_safe(epi.get('replacement_days'), 30)
    return None


async def build_replacement_alerts(db) -> tuple[List[dict], List[dict]]:
    expiring_alerts = []
    due_alerts = []
    epis_with_period = await db.epis.find({'replacement_period': {'$ne': None}}).to_list(1000)
    epi_periods = {str(epi['_id']): epi for epi in epis_with_period if get_replacement_days(epi)}
    if not epi_periods:
        return expiring_alerts, due_alerts

    employees = await db.employees.find({'status': 'active'}, {'full_name': 1, 'department': 1}).to_list(5000)
    emp_map = {str(employee['_id']): employee for employee in employees}
    pipeline = [
        {'$match': {'is_return': False}},
        {'$unwind': '$items'},
        {'$match': {'items.epi_id': {'$in': list(epi_periods.keys())}}},
        {'$sort': {'created_at': -1}},
        {'$group': {
            '_id': {'employee_id': '$employee_id', 'epi_id': '$items.epi_id'},
            'last_delivery': {'$first': '$created_at'},
            'employee_name': {'$first': '$employee_name'},
        }},
    ]
    results = await db.deliveries.aggregate(pipeline).to_list(20000)
    today = datetime.now(timezone.utc).date()
    for result in results:
        emp_id = str(result['_id']['employee_id'])
        if emp_id not in emp_map:
            continue
        epi_id = str(result['_id']['epi_id'])
        epi = epi_periods.get(epi_id)
        last_delivery = ensure_aware_datetime(result.get('last_delivery'))
        if not epi or not last_delivery:
            continue
        replacement_days = get_replacement_days(epi)
        due_date = last_delivery + timedelta(days=replacement_days)
        days_until_due = (due_date.date() - today).days
        base_payload = {
            'employee_id': emp_id,
            'employee_name': result.get('employee_name') or emp_map[emp_id].get('full_name', ''),
            'department': emp_map[emp_id].get('department', ''),
            'epi_id': epi_id,
            'epi_name': epi.get('name', ''),
            'ca_number': epi.get('ca_number', ''),
            'nbr_number': epi.get('nbr_number', ''),
            'last_delivery_date': last_delivery.isoformat(),
            'replacement_due_date': due_date.isoformat(),
            'relevant_date': due_date.isoformat(),
            'replacement_period': epi.get('replacement_period', ''),
            'replacement_days': replacement_days,
            'entity': epi.get('name', ''),
        }
        if days_until_due < 0:
            due_alerts.append(normalize_alert_item(
                'replacement_due',
                alert_id=build_alert_id('replacement_due', emp_id, epi_id),
                days_overdue=abs(days_until_due),
                days_until_due=days_until_due,
                status='Vencida',
                message=f"{base_payload['employee_name']} não realizou retirada de {epi.get('name')} - {abs(days_until_due)} dias de atraso",
                **base_payload,
            ))
        elif days_until_due <= ALERT_REPLACEMENT_WARNING_DAYS:
            expiring_alerts.append(normalize_alert_item(
                'replacement_expiring',
                alert_id=build_alert_id('replacement_expiring', emp_id, epi_id),
                days_until_due=days_until_due,
                status='A vencer',
                message=f"Troca de {epi.get('name')} prevista em {days_until_due} dia(s)",
                **base_payload,
            ))
    return expiring_alerts, due_alerts


async def get_stock_operational_sources(db) -> List[dict]:
    epis = await db.epis.find({}).to_list(10000)
    variations = await db.epi_variations.find({}).to_list(20000)
    variations_by_epi = {}
    for variation in variations:
        variations_by_epi.setdefault(str(variation.get('epi_id')), []).append(variation)

    sources = []
    for epi in epis:
        epi_id = str(epi['_id'])
        epi_variations = variations_by_epi.get(epi_id) or []
        if epi_variations:
            for variation in epi_variations:
                sources.append({'source': 'variation', 'epi': epi, 'variation': variation})
        else:
            sources.append({'source': 'epi', 'epi': epi, 'variation': None})
    return sources


def source_variation_label(variation: Optional[dict]) -> str:
    if not variation:
        return ''
    parts = [
        variation.get('variation_name'),
        variation.get('size'),
        variation.get('model'),
        variation.get('brand'),
        variation.get('color'),
    ]
    return ' / '.join(str(part).strip() for part in parts if str(part or '').strip())


def pick_source_date(source: dict, field_name: str):
    variation = source.get('variation') or {}
    epi = source.get('epi') or {}
    return variation.get(field_name) if variation.get(field_name) is not None else epi.get(field_name)


def pick_source_value(source: dict, field_name: str, default=None):
    variation = source.get('variation') or {}
    epi = source.get('epi') or {}
    return variation.get(field_name) if variation.get(field_name) is not None else epi.get(field_name, default)


async def build_stock_and_ca_alerts(db) -> tuple[List[dict], List[dict], List[dict], List[dict], dict]:
    stock_low = []
    stock_zero = []
    ca_expiring = []
    ca_expired = []
    ca_missing_count = 0
    sources = await get_stock_operational_sources(db)
    today = datetime.now(timezone.utc).date()

    for source in sources:
        epi = source['epi']
        variation = source.get('variation')
        epi_id = str(epi['_id'])
        variation_id = str(variation['_id']) if variation else None
        stock = parse_int_safe(pick_source_value(source, 'current_stock', 0), 0)
        min_stock = parse_int_safe(pick_source_value(source, 'min_stock', 0), 0)
        variation_label = source_variation_label(variation)
        base_payload = {
            'epi_id': epi_id,
            'epi_name': epi.get('name', ''),
            'name': epi.get('name', ''),
            'variation_id': variation_id,
            'variation_name': variation_label,
            'stock_source': source['source'],
            'current_stock': stock,
            'min_stock': min_stock,
            'entity': epi.get('name', ''),
        }
        if stock <= 0:
            stock_zero.append(normalize_alert_item(
                'stock_zero',
                alert_id=build_alert_id('stock_zero', epi_id, variation_id),
                status='Zerado' if stock == 0 else 'Negativo',
                message=f"{epi.get('name', 'EPI')} sem estoque operacional disponível.",
                **base_payload,
            ))
        elif stock <= min_stock:
            stock_low.append(normalize_alert_item(
                'stock_low',
                alert_id=build_alert_id('stock_low', epi_id, variation_id),
                status='Baixo',
                message=f"{epi.get('name', 'EPI')} está com estoque baixo.",
                **base_payload,
            ))

        ca_number = pick_source_value(source, 'ca_number')
        ca_validity = ensure_aware_datetime(pick_source_date(source, 'ca_validity') or pick_source_date(source, 'validity_date'))
        if not ca_number or not ca_validity:
            ca_missing_count += 1
            continue
        days_until_ca = (ca_validity.date() - today).days
        ca_payload = {
            'epi_id': epi_id,
            'epi_name': epi.get('name', ''),
            'variation_id': variation_id,
            'variation_name': variation_label,
            'stock_source': source['source'],
            'ca_number': ca_number,
            'ca_validity': ca_validity.isoformat(),
            'validity_date': ca_validity.isoformat(),
            'relevant_date': ca_validity.isoformat(),
            'days_until_expiration': days_until_ca,
            'entity': epi.get('name', ''),
        }
        if days_until_ca < 0:
            ca_expired.append(normalize_alert_item(
                'ca_expired',
                alert_id=build_alert_id('ca_expired', epi_id, variation_id, ca_number),
                status='Vencido',
                days_overdue=abs(days_until_ca),
                message=f"CA {ca_number} vencido há {abs(days_until_ca)} dia(s).",
                **ca_payload,
            ))
        elif days_until_ca <= ALERT_CA_WARNING_DAYS:
            ca_expiring.append(normalize_alert_item(
                'ca_expiring',
                alert_id=build_alert_id('ca_expiring', epi_id, variation_id, ca_number),
                status='A vencer',
                message=f"CA {ca_number} vence em {days_until_ca} dia(s).",
                **ca_payload,
            ))

    audit = {
        'source_rule': 'variation_when_available_else_epi',
        'ca_warning_days': ALERT_CA_WARNING_DAYS,
        'ca_missing_sources': ca_missing_count,
        'operational_sources': len(sources),
    }
    return stock_low, stock_zero, ca_expiring, ca_expired, audit


async def build_course_operational_alerts(db) -> List[dict]:
    records = await db.employee_courses.find({'active': True, 'expiration_date': {'$ne': None}}).to_list(10000)
    grouped_records = {}
    for record in records:
        employee_id = str(record.get('employee_id') or '')
        course_id = str(record.get('course_id') or '')
        if employee_id and course_id:
            grouped_records.setdefault((employee_id, course_id), []).append(record)

    current_records = [current for items in grouped_records.values() if (current := pick_current_employee_course(items))]
    if not current_records:
        return []

    course_map = await get_courses_map(db, [record.get('course_id') for record in current_records])
    employee_object_ids = [ObjectId(record.get('employee_id')) for record in current_records if ObjectId.is_valid(str(record.get('employee_id')))]
    employees = await db.employees.find({'_id': {'$in': employee_object_ids}, 'status': 'active'}).to_list(5000)
    employee_map = {str(employee['_id']): employee for employee in employees}
    company_object_ids = [ObjectId(employee.get('company_id')) for employee in employees if ObjectId.is_valid(str(employee.get('company_id')))]
    companies = await db.companies.find({'_id': {'$in': company_object_ids}}).to_list(1000)
    company_map = {str(company['_id']): company for company in companies}

    items = []
    for record in current_records:
        employee = employee_map.get(str(record.get('employee_id')))
        course = course_map.get(str(record.get('course_id')))
        if not employee or not course or not course.get('active', True):
            continue
        status_key, status_label, days_until_expiration = calculate_employee_course_status(record, course)
        if status_key not in ('expiring', 'expired'):
            continue
        company = company_map.get(str(employee.get('company_id')))
        alert_type = 'course_expired' if status_key == 'expired' else 'course_expiring'
        items.append(normalize_alert_item(
            alert_type,
            alert_id=build_alert_id(alert_type, str(employee['_id']), str(course['_id'])),
            id=str(record['_id']),
            employee_id=str(employee['_id']),
            employee_name=employee.get('full_name') or 'Colaborador sem nome',
            company_name=(company or {}).get('legal_name') or (company or {}).get('trade_name') or (company or {}).get('name') or 'Sem empresa',
            department=employee.get('department') or '',
            course_id=str(course['_id']),
            course_name=course.get('name') or 'Curso sem nome',
            expiration_date=ensure_aware_datetime(record.get('expiration_date')),
            relevant_date=ensure_aware_datetime(record.get('expiration_date')).isoformat() if ensure_aware_datetime(record.get('expiration_date')) else None,
            status=status_key,
            status_label=status_label,
            days_until_expiration=days_until_expiration,
            entity=course.get('name') or 'Curso sem nome',
            message=f"{course.get('name') or 'Curso'} de {employee.get('full_name') or 'colaborador'} {status_label.lower()}.",
        ))

    items.sort(key=lambda item: (
        0 if item['type'] == 'course_expired' else 1,
        item['days_until_expiration'],
        (item.get('employee_name') or '').lower(),
        (item.get('course_name') or '').lower(),
    ))
    return items


async def build_alert_center(db, current_user: dict, include_stock: Optional[bool] = None) -> dict:
    allow_stock = user_can_view_stock_alerts(current_user) if include_stock is None else bool(include_stock)
    pending_epis = await build_pending_epi_alerts(db)
    replacement_expiring, replacement_due = await build_replacement_alerts(db)
    course_items = await build_course_operational_alerts(db)
    course_expiring = [item for item in course_items if item['type'] == 'course_expiring']
    course_expired = [item for item in course_items if item['type'] == 'course_expired']

    stock_low = []
    stock_zero = []
    ca_expiring = []
    ca_expired = []
    stock_audit = {'source_rule': 'restricted_by_permission'}
    if allow_stock:
        stock_low, stock_zero, ca_expiring, ca_expired, stock_audit = await build_stock_and_ca_alerts(db)

    alerts_by_type = {
        'pending_epi': pending_epis,
        'replacement_expiring': replacement_expiring,
        'replacement_due': replacement_due,
        'stock_low': stock_low,
        'stock_zero': stock_zero,
        'ca_expiring': ca_expiring,
        'ca_expired': ca_expired,
        'course_expiring': course_expiring,
        'course_expired': course_expired,
    }
    all_items = []
    seen = set()
    for alert_type in ALERT_ORDER:
        for item in alerts_by_type.get(alert_type, []):
            alert_id = item.get('alert_id')
            if alert_id in seen:
                continue
            seen.add(alert_id)
            all_items.append(item)

    summary = summarize_alerts(alerts_by_type)
    return {
        'summary': summary,
        'alerts': all_items,
        'alerts_by_type': alerts_by_type,
        'pending_epis': pending_epis,
        'replacement_expiring': replacement_expiring,
        'replacement_due': replacement_due,
        'stock_low': stock_low,
        'stock_zero': stock_zero,
        'ca_expiring': ca_expiring,
        'ca_expired': ca_expired,
        'course_expiring': course_expiring,
        'course_expired': course_expired,
        'courses': {'summary': {'expiring': len(course_expiring), 'expired': len(course_expired)}, 'items': course_items},
        'total_pending_epis': len(pending_epis),
        'total_replacement_expiring': len(replacement_expiring),
        'total_replacement_due': len(replacement_due),
        'total_alerts': summary['total_alerts'],
        'metadata': {
            'stock_visible': allow_stock,
            'stock': stock_audit,
            'replacement_warning_days': ALERT_REPLACEMENT_WARNING_DAYS,
        }
    }


@api_router.get('/stock/alerts')
async def get_stock_alerts(current_user: dict = Depends(require_role('admin', 'gestor', 'almoxarifado'))):
    db = await get_db()
    stock_low, stock_zero, ca_expiring, ca_expired, audit = await build_stock_and_ca_alerts(db)
    return {
        'low_stock': stock_low + stock_zero,
        'expiring_soon': ca_expiring + ca_expired,
        'stock_low': stock_low,
        'stock_zero': stock_zero,
        'ca_expiring': ca_expiring,
        'ca_expired': ca_expired,
        'summary': {
            'stock_low': len(stock_low),
            'stock_zero': len(stock_zero),
            'ca_expiring': len(ca_expiring),
            'ca_expired': len(ca_expired),
        },
        'metadata': audit,
    }


@api_router.get('/alerts/courses')
async def get_course_alerts(current_user: dict = Depends(get_current_user)):
    if not can_view_employee_courses(current_user['role']):
        raise HTTPException(status_code=403, detail='Sem permissao para consultar alertas de cursos')
    db = await get_db()
    items = await build_course_operational_alerts(db)
    return {
        'summary': {
            'expiring': sum(1 for item in items if item['type'] == 'course_expiring'),
            'expired': sum(1 for item in items if item['type'] == 'course_expired'),
        },
        'items': items,
    }


# ===================== ALERTAS DE EPI OBRIGATÓRIO E PERIODICIDADE =====================

@api_router.get('/alerts/pending-epis')
async def get_pending_epi_alerts(current_user: dict = Depends(get_current_user)):
    db = await get_db()
    return await build_pending_epi_alerts(db)


@api_router.get('/alerts/replacement-expiring')
async def get_replacement_expiring_alerts(current_user: dict = Depends(get_current_user)):
    db = await get_db()
    replacement_expiring, _ = await build_replacement_alerts(db)
    return replacement_expiring


@api_router.get('/alerts/replacement-due')
async def get_replacement_due_alerts(current_user: dict = Depends(get_current_user)):
    db = await get_db()
    _, replacement_due = await build_replacement_alerts(db)
    return replacement_due


@api_router.get('/alerts/all')
async def get_all_alerts(current_user: dict = Depends(get_current_user)):
    db = await get_db()
    center = await build_alert_center(db, current_user)
    return {
        'pending_epis': center['pending_epis'],
        'replacement_expiring': center['replacement_expiring'],
        'replacement_due': center['replacement_due'],
        'stock_low': center['stock_low'],
        'stock_zero': center['stock_zero'],
        'ca_expiring': center['ca_expiring'],
        'ca_expired': center['ca_expired'],
        'course_expiring': center['course_expiring'],
        'course_expired': center['course_expired'],
        'alerts': center['alerts'],
        'alerts_by_type': center['alerts_by_type'],
        'summary': center['summary'],
        'total_pending_epis': center['total_pending_epis'],
        'total_replacement_expiring': center['total_replacement_expiring'],
        'total_replacement_due': center['total_replacement_due'],
        'total_stock_low': len(center['stock_low']),
        'total_stock_zero': len(center['stock_zero']),
        'total_ca_expiring': len(center['ca_expiring']),
        'total_ca_expired': len(center['ca_expired']),
        'total_course_expiring': len(center['course_expiring']),
        'total_course_expired': len(center['course_expired']),
        'total_alerts': center['total_alerts'],
        'metadata': center['metadata'],
    }


@api_router.get('/alerts/employee/{employee_id}')
async def get_employee_alerts(employee_id: str, current_user: dict = Depends(get_current_user)):
    db = await get_db()
    employee = await db.employees.find_one({'_id': ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador não encontrado')

    center = await build_alert_center(db, current_user, include_stock=False)
    pending = [item for item in center['pending_epis'] if item.get('employee_id') == employee_id]
    replacement_expiring = [item for item in center['replacement_expiring'] if item.get('employee_id') == employee_id]
    replacement_due = [item for item in center['replacement_due'] if item.get('employee_id') == employee_id]
    course_expiring = [item for item in center['course_expiring'] if item.get('employee_id') == employee_id]
    course_expired = [item for item in center['course_expired'] if item.get('employee_id') == employee_id]
    alerts_by_type = {
        'pending_epi': pending,
        'replacement_expiring': replacement_expiring,
        'replacement_due': replacement_due,
        'stock_low': [],
        'stock_zero': [],
        'ca_expiring': [],
        'ca_expired': [],
        'course_expiring': course_expiring,
        'course_expired': course_expired,
    }
    summary = summarize_alerts(alerts_by_type)
    return {
        'pending_epis': pending,
        'replacement_expiring': replacement_expiring,
        'replacement_due': replacement_due,
        'course_expiring': course_expiring,
        'course_expired': course_expired,
        'alerts_by_type': alerts_by_type,
        'summary': summary,
        'total_alerts': summary['total_alerts'],
    }

# ===================== SETOR-KIT VINCULAÃ‡ÃƒO =====================

@api_router.get('/employees/{employee_id}/delivery-suggestions', response_model=DeliverySuggestionResponse)
async def get_delivery_suggestions(employee_id: str, current_user: dict = Depends(require_role('admin', 'almoxarifado')), kit_id: Optional[str] = None):
    db = await get_db()
    employee = await db.employees.find_one({"_id": ObjectId(employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail='Colaborador nÃ£o encontrado')

    if not is_employee_active(employee):
        raise HTTPException(status_code=409, detail='Este colaborador esta inativo e nao pode receber uma nova entrega de EPI.')

    kit = None
    if kit_id:
        kit = await db.kits.find_one({"_id": ObjectId(kit_id)})
    elif employee.get('department'):
        kit = await db.kits.find_one({
            "sector": {"$regex": f"^{employee.get('department')}$", "$options": "i"},
            "is_mandatory": {"$ne": False}
        })

    if not kit:
        return DeliverySuggestionResponse(
            employee_id=employee_id,
            employee_name=employee.get('full_name', ''),
            department=employee.get('department'),
            kit_id=None,
            kit_name=None,
            items=[]
        )

    return DeliverySuggestionResponse(**(await build_delivery_suggestions_for_kit(db, employee, kit)))

@api_router.get('/kits/by-sector/{sector_name}')
async def get_kit_by_sector(sector_name: str, current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho', 'almoxarifado'))):
    """Retorna o kit obrigatÃ³rio vinculado a um setor"""
    db = await get_db()
    
    kit = await db.kits.find_one({
        "sector": {"$regex": f"^{sector_name}$", "$options": "i"},
        "is_mandatory": {"$ne": False}
    })
    
    if not kit:
        return None

    return KitResponse(**(await build_kit_response(db, kit)))

@api_router.get('/sectors/list')
async def get_sectors_list(current_user: dict = Depends(require_role('admin', 'gestor', 'seguranca_trabalho'))):
    """Retorna lista de setores com seus kits vinculados"""
    db = await get_db()
    
    # Buscar setores dos colaboradores
    pipeline = [
        {"$match": {"department": {"$ne": None, "$ne": ""}}},
        {"$group": {"_id": "$department"}},
        {"$sort": {"_id": 1}}
    ]
    
    sectors_result = await db.employees.aggregate(pipeline).to_list(100)
    sectors = [s['_id'] for s in sectors_result if s['_id']]
    
    # Buscar kits
    kits = await db.kits.find({}).to_list(100)
    kits_by_sector = {(k.get('sector') or '').lower(): k for k in kits}
    
    result = []
    for sector in sectors:
        kit = kits_by_sector.get(sector.lower())
        result.append({
            'sector_name': sector,
            'has_kit': kit is not None,
            'kit_id': str(kit['_id']) if kit else None,
            'kit_name': kit.get('name') if kit else None,
            'kit_items_count': len(kit.get('items', [])) if kit else 0
        })
    
    return result

# ===================== DASHBOARD =====================

@api_router.get('/dashboard/stats')
async def get_dashboard_stats(current_user: dict = Depends(get_current_user)):
    db = await get_db()

    total_employees = await db.employees.count_documents({})
    inactive_employees = await db.employees.count_documents({'status': 'inactive'})
    active_employees = total_employees - inactive_employees
    total_epis = await db.epis.count_documents({})
    active_employee_docs = await db.employees.find({'$or': [{'status': 'active'}, {'status': {'$exists': False}}, {'status': None}]}).to_list(5000)
    biometric_summaries = await get_biometric_summaries(db, active_employee_docs)
    biometric_registered = sum(1 for item in biometric_summaries.values() if item['status'] == 'registered')
    biometric_incomplete = sum(1 for item in biometric_summaries.values() if item['status'] == 'incomplete')
    biometric_missing = active_employees - biometric_registered - biometric_incomplete
    biometric_coverage = round((biometric_registered / active_employees) * 100, 1) if active_employees else 0

    thirty_days_ago = datetime.now(timezone.utc) - timedelta(days=30)
    recent_deliveries = await db.deliveries.count_documents({
        'is_return': False,
        'created_at': {'$gte': thirty_days_ago}
    })

    try:
        alert_center = await build_alert_center(db, current_user)
        summary = alert_center.get('summary', empty_alert_summary())
    except Exception as e:
        logger.error(f"Erro ao buscar alertas: {e}")
        alert_center = {'summary': empty_alert_summary(), 'metadata': {}}
        summary = alert_center['summary']

    return {
        'total_employees': total_employees,
        'active_employees': active_employees,
        'inactive_employees': inactive_employees,
        'colaboradores_total': total_employees,
        'colaboradores_ativos': active_employees,
        'colaboradores_inativos': inactive_employees,
        'total_epis': total_epis,
        'low_stock_count': summary.get('stock_low', 0) + summary.get('stock_zero', 0),
        'recent_deliveries': recent_deliveries,
        'expiring_epis': summary.get('ca_expiring', 0) + summary.get('ca_expired', 0),
        'pending_epi_alerts': summary.get('pending_epi', 0),
        'replacement_expiring_alerts': summary.get('replacement_expiring', 0),
        'replacement_due_alerts': summary.get('replacement_due', 0),
        'stock_low_alerts': summary.get('stock_low', 0),
        'stock_zero_alerts': summary.get('stock_zero', 0),
        'ca_expiring_alerts': summary.get('ca_expiring', 0),
        'ca_expired_alerts': summary.get('ca_expired', 0),
        'course_expiring_alerts': summary.get('course_expiring', 0),
        'course_expired_alerts': summary.get('course_expired', 0),
        'total_alerts': summary.get('total_alerts', 0),
        'alerts_summary': summary,
        'alerts_metadata': alert_center.get('metadata', {}),
        'biometric_total_employees': active_employees,
        'biometric_registered': biometric_registered,
        'biometric_missing': biometric_missing,
        'biometric_incomplete': biometric_incomplete,
        'biometric_coverage_percent': biometric_coverage
    }

# ===================== AUTENTICAÃ‡ÃƒO DE FICHA EPI =====================

import hashlib
import secrets

def generate_auth_code(employee_id: str, timestamp: datetime) -> str:
    """Gera cÃ³digo de autenticaÃ§Ã£o Ãºnico para ficha de EPI"""
    # Criar hash baseado em employee_id, timestamp e secret
    secret_key = os.environ.get('SECRET_KEY', 'CHANGE_ME_IN_ENV')
    data = f"{employee_id}:{timestamp.isoformat()}:{secret_key}:{secrets.token_hex(4)}"
    hash_obj = hashlib.sha256(data.encode())
    hash_hex = hash_obj.hexdigest()[:10].upper()
    return f"AUT-EPI-{hash_hex}"

@api_router.post('/ficha-auth/generate', response_model=FichaAuthenticationResponse)
async def generate_ficha_authentication(
    data: FichaAuthenticationCreate,
    current_user: dict = Depends(require_role('admin', 'rh', 'seguranca_trabalho', 'almoxarifado'))
):
    """Gera cÃ³digo de autenticaÃ§Ã£o para ficha de EPI com validaÃ§Ã£o biomÃ©trica"""
    db = await get_db()
    
    # Buscar colaborador
    employee = await db.employees.find_one({"_id": ObjectId(data.employee_id)})
    if not employee:
        raise HTTPException(status_code=404, detail="Colaborador nÃ£o encontrado")
    
    # Buscar entregas do colaborador
    delivery_query = {"employee_id": data.employee_id}
    if not data.include_all_history and data.delivery_ids:
        delivery_query["_id"] = {"$in": [ObjectId(d) for d in data.delivery_ids]}
    
    deliveries = await db.deliveries.find(delivery_query).to_list(1000)
    delivery_ids = [str(d["_id"]) for d in deliveries]
    
    # Verificar se hÃ¡ validaÃ§Ã£o biomÃ©trica nas entregas
    biometric_validated = False
    max_biometric_score = 0.0
    
    for delivery in deliveries:
        if delivery.get("facial_match_score"):
            biometric_validated = True
            score = delivery.get("facial_match_score", 0)
            if score > max_biometric_score:
                max_biometric_score = score
    
    # Gerar cÃ³digo de autenticaÃ§Ã£o
    now = datetime.now(timezone.utc)
    auth_code = generate_auth_code(data.employee_id, now)
    
    # Dados para QR Code (JSON que pode ser escaneado para verificaÃ§Ã£o)
    qr_data = {
        "type": "EPI_FICHA_AUTH",
        "code": auth_code,
        "employee_id": data.employee_id,
        "employee_name": employee.get("full_name"),
        "date": now.isoformat(),
        "biometric": biometric_validated,
        "verify_url": f"/api/ficha-auth/verify/{auth_code}"
    }
    
    # Salvar no banco
    auth_record = {
        "auth_code": auth_code,
        "employee_id": data.employee_id,
        "employee_name": employee.get("full_name"),
        "delivery_ids": delivery_ids,
        "biometric_validated": biometric_validated,
        "biometric_score": max_biometric_score if biometric_validated else None,
        "created_at": now,
        "created_by": current_user.get("username"),
        "created_by_id": current_user.get("id"),
        "qr_code_data": str(qr_data),
        "status": "active"
    }
    
    result = await db.ficha_authentications.insert_one(auth_record)
    
    return FichaAuthenticationResponse(
        id=str(result.inserted_id),
        auth_code=auth_code,
        employee_id=data.employee_id,
        employee_name=employee.get("full_name"),
        delivery_ids=delivery_ids,
        biometric_validated=biometric_validated,
        biometric_score=max_biometric_score if biometric_validated else None,
        created_at=now,
        created_by=current_user.get("username"),
        qr_code_data=str(qr_data)
    )

@api_router.get('/ficha-auth/verify/{auth_code}', response_model=FichaAuthenticationVerify)
async def verify_ficha_authentication(auth_code: str):
    """Verifica autenticidade de uma ficha de EPI pelo cÃ³digo"""
    db = await get_db()
    
    # Buscar registro de autenticaÃ§Ã£o
    auth_record = await db.ficha_authentications.find_one({"auth_code": auth_code})
    
    if not auth_record:
        return FichaAuthenticationVerify(
            valid=False,
            auth_code=auth_code,
            message="CÃ³digo de autenticaÃ§Ã£o nÃ£o encontrado no sistema"
        )
    
    if auth_record.get("status") == "revoked":
        return FichaAuthenticationVerify(
            valid=False,
            auth_code=auth_code,
            employee_name=auth_record.get("employee_name"),
            validation_date=auth_record.get("created_at"),
            message="Este cÃ³digo foi revogado e nÃ£o Ã© mais vÃ¡lido"
        )
    
    return FichaAuthenticationVerify(
        valid=True,
        auth_code=auth_code,
        employee_name=auth_record.get("employee_name"),
        validation_date=auth_record.get("created_at"),
        biometric_validated=auth_record.get("biometric_validated"),
        message="Documento autÃªntico - validado pelo sistema GestorEPI"
    )

@api_router.get('/ficha-auth/employee/{employee_id}')
async def get_employee_authentications(
    employee_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Lista todas as autenticaÃ§Ãµes de ficha de um colaborador"""
    db = await get_db()
    
    records = await db.ficha_authentications.find(
        {"employee_id": employee_id}
    ).sort("created_at", -1).to_list(100)
    
    return [{
        "id": str(r["_id"]),
        "auth_code": r["auth_code"],
        "created_at": r["created_at"],
        "biometric_validated": r.get("biometric_validated"),
        "status": r.get("status", "active")
    } for r in records]

from reporting import report_router
api_router.include_router(report_router)
app.include_router(api_router)
